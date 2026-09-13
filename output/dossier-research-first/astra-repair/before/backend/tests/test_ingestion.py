from __future__ import annotations

import io

import pytest
from fastapi import UploadFile
from docx import Document
from pypdf import PdfWriter

from app.services import ingestion


@pytest.mark.asyncio
async def test_pdf_upload_always_creates_editable_companion(app_context):
    pdf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=200)
    writer.write(pdf)
    pdf.seek(0)

    upload = UploadFile(file=pdf, filename="blank-source.pdf")
    result = await app_context.ingestion.upload_to_matter("MAT-DEMO-BEACON", upload)

    assert result["extracted_path"]
    companion = app_context.vault.read_markdown(result["extracted_path"])
    assert companion["metadata"]["source_filename"] == "blank-source.pdf"
    assert "No extractable text was found" in companion["content"]


@pytest.mark.asyncio
async def test_upload_rejects_unsupported_file_before_writing(app_context):
    upload = UploadFile(file=io.BytesIO(b"not a supported document"), filename="payload.exe")

    with pytest.raises(ValueError, match="Unsupported file type"):
        await app_context.ingestion.upload_to_matter("MAT-DEMO-BEACON", upload)

    assert not app_context.vault.exists(
        "03_Matters/beacon-instant-onboarding/documents/payload.exe"
    )


@pytest.mark.asyncio
async def test_multi_file_upload_returns_stable_attachment_refs_and_read_only_preview(app_context):
    uploads = [
        UploadFile(file=io.BytesIO(b"Launch plan and customer notice."), filename="plan.md"),
        UploadFile(file=io.BytesIO(b"Support escalation details."), filename="notes.txt"),
    ]

    result = await app_context.ingestion.upload_many_to_matter("MAT-DEMO-BEACON", uploads)

    assert result["presentation"] == "plural"
    assert result["intent_required"] is True
    assert result["apply_required"] is True
    assert len(result["attachments"]) == 2
    assert all(item["source_id"].startswith("SRC-") for item in result["attachments"])
    assert all(len(item["version"]) == 64 for item in result["attachments"])
    assert result["scan"][0]["word_count"] > 0


@pytest.mark.asyncio
async def test_image_upload_is_preserved_as_source_without_extraction(app_context):
    upload = UploadFile(file=io.BytesIO(b"fake-image-content"), filename="screen.png")

    result = await app_context.ingestion.upload_to_matter("MAT-DEMO-BEACON", upload)

    assert result["source_only"] is True
    assert result["extracted_path"] is None
    assert app_context.vault.exists(result["path"])


@pytest.mark.asyncio
async def test_same_content_has_stable_source_and_version(app_context):
    first = await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON", UploadFile(file=io.BytesIO(b"same"), filename="one.txt")
    )
    second = await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON", UploadFile(file=io.BytesIO(b"same"), filename="two.txt")
    )

    assert first["source_id"] == second["source_id"]
    assert first["version"] == second["version"]


@pytest.mark.asyncio
async def test_docx_upload_converts_headings_lists_and_tables_to_markdown(app_context):
    source = io.BytesIO()
    document = Document()
    document.add_heading("Service Terms", level=1)
    document.add_paragraph("Keep records for 30 days.")
    document.add_paragraph("Notify customers", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Notice"
    table.rows[1].cells[1].text = "30 days"
    document.save(source)
    source.seek(0)

    result = await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON", UploadFile(file=source, filename="terms.docx")
    )

    markdown = app_context.vault.read_markdown(result["extracted_path"])["content"]
    assert "# Service Terms" in markdown
    assert "- Notify customers" in markdown
    assert "| Term | Value |" in markdown
    assert "| --- | --- |" in markdown


@pytest.mark.asyncio
async def test_docx_extraction_runs_off_the_event_loop(app_context, monkeypatch):
    source = io.BytesIO()
    document = Document()
    document.add_paragraph("Off event loop")
    document.save(source)
    calls = []

    async def to_thread(function, *args):
        calls.append(function)
        return function(*args)

    monkeypatch.setattr(ingestion.asyncio, "to_thread", to_thread)
    await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON", UploadFile(file=io.BytesIO(source.getvalue()), filename="thread.docx")
    )

    assert calls[0] == app_context.ingestion._extract


@pytest.mark.asyncio
async def test_pdf_extraction_runs_off_the_event_loop(app_context, monkeypatch):
    source = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=200)
    writer.write(source)
    calls = []

    async def to_thread(function, *args):
        calls.append(function)
        return function(*args)

    monkeypatch.setattr(ingestion.asyncio, "to_thread", to_thread)
    await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON", UploadFile(file=io.BytesIO(source.getvalue()), filename="thread.pdf")
    )

    assert calls[0] == app_context.ingestion._extract


@pytest.mark.asyncio
async def test_upload_uses_configured_source_folder_but_batch_records_stay_fixed(app_context):
    app_context.settings_store.write({"matter_files.source_documents_dir": "source-files"})
    result = await app_context.ingestion.upload_many_to_matter(
        "MAT-DEMO-BEACON",
        [UploadFile(file=io.BytesIO(b"Source text"), filename="source.txt")],
    )
    assert "/source-files/source.txt" in result["attachments"][0]["path"]
    assert app_context.vault.exists(
        f"03_Matters/beacon-instant-onboarding/documents/batches/{result['batch_id']}.md"
    )
