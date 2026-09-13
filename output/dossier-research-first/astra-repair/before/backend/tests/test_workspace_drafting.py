from __future__ import annotations

import hashlib
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfReader

from app.models.api import DocumentReviewAction
from app.services.document_review import DocumentReviewService
from app.services.document_export import DocumentExportService
from app.services.workspace import WorkspaceConflict
from app.services.vault import VaultService

MATTER = "MAT-DEMO-BEACON"


def draft(ctx, title="Memo", content="Client means Beacon.\n\nDeliver in ten days.", **kwargs):
    return ctx.work_products.create_draft(MATTER, title=title, content=content, **kwargs)


def revise(ctx, item, content, key="revise:1", **kwargs):
    return ctx.work_products.revise_draft(MATTER, item["vault_path"], title=item["title"], content=content,
        source_action_key=key, document_reviews=ctx.document_reviews,
        expected_revision=item["artifact"]["revision"], **kwargs)


def test_three_named_artifacts_range_revision_and_persistent_retries(app_context):
    ctx = app_context
    memo = draft(ctx, output_type="memo", source_action_key="create:memo")
    clause = draft(ctx, title="Clause", output_type="clause")
    checklist = draft(ctx, title="Checklist", content="- Confirm timing\n- Preserve Client term", output_type="checklist")
    before = {i["vault_path"]: ctx.vault.resolve(i["vault_path"]).read_bytes() for i in (memo, checklist)}
    text = ctx.vault.read_document(clause["vault_path"])["content"]
    start = text.index("ten")
    selection = {"start": start, "end": start + 3, "text": "ten", "prefix": "in ", "suffix": " days."}
    result = revise(ctx, clause, "twenty", selected_range=selection, reason="The reported delivery period changed.")
    assert ctx.vault.read_document(clause["vault_path"])["content"] == text.replace("ten", "twenty")
    assert result["artifact"]["pending_review"]
    assert result["artifact"]["version_changes"][-1]["reason"] == "The reported delivery period changed."
    for path, data in before.items():
        assert ctx.vault.resolve(path).read_bytes() == data
    # The exact old request can reconnect after a newer version exists.
    retry = revise(ctx, clause, "twenty", selected_range=selection, reason="The reported delivery period changed.")
    assert retry["changed_paths"] == []
    with pytest.raises(WorkspaceConflict, match="different draft text"):
        revise(ctx, clause, "forty", selected_range=selection, reason="The reported delivery period changed.")
    older = ctx.vault.read_document(result["artifact"]["version_changes"][-1]["revision_path"])
    assert older["content"] == text
    assert older["metadata"]["immutable"] is True


@pytest.mark.parametrize("conflict", ["content", "review", "local", "range"])
def test_stale_target_preserves_useful_proposal_and_lawyer_text(app_context, conflict):
    ctx = app_context
    item = draft(ctx)
    review = ctx.document_reviews.get(item["vault_path"])
    kwargs = {"expected_review_revision": review["revision"]}
    if conflict == "content":
        ctx.document_reviews.apply(item["vault_path"], DocumentReviewAction(action="save_untracked", content="Lawyer's new text"))
    elif conflict == "review":
        ctx.document_reviews.apply(item["vault_path"], DocumentReviewAction(action="add_comment", quote="Client", body="Keep this term",
            author_name="Counsel", author_id="lawyer", anchor_start=0, anchor_end=6))
    elif conflict == "local":
        kwargs["local_draft_snapshot"] = "Unsaved lawyer text"
    else:
        kwargs["selected_range"] = {"start": 999, "end": 1005, "text": "Client"}
    before = ctx.vault.resolve(item["vault_path"]).read_bytes()
    with pytest.raises(WorkspaceConflict) as error:
        revise(ctx, item, "Useful generated proposal", **kwargs)
    assert ctx.vault.resolve(item["vault_path"]).read_bytes() == before
    assert ctx.vault.read_document(error.value.detail["proposal_path"])["content"] == "Useful generated proposal\n"
    assert ctx.vault.read_document(error.value.detail["proposal_path"])["metadata"]["local_draft_snapshot"] == kwargs.get("local_draft_snapshot")


def test_review_accept_reject_history_metadata_guard_and_comment_retention(app_context):
    ctx = app_context
    item = draft(ctx, content="Client pays ten dollars.")
    before = ctx.document_reviews.get(item["vault_path"])
    result = revise(ctx, item, "Client pays twenty dollars.", expected_review_revision=before["revision"], reason="Updated price")
    review = ctx.document_reviews.get(item["vault_path"])
    changed = review["changes"][0]["change_id"]
    accepted = ctx.document_reviews.apply(item["vault_path"], DocumentReviewAction(action="accept_change", change_id=changed),
        expected_review_revision=review["revision"], reason="Counsel accepted price wording")
    assert accepted["revision"] != review["revision"]
    assert accepted["artifact_revision"] == result["artifact"]["revision"]
    with pytest.raises(WorkspaceConflict):
        ctx.document_reviews.apply(item["vault_path"], DocumentReviewAction(action="set_tracking", enabled=False), expected_review_revision=review["revision"])
    assert accepted["history"][-1]["reason"] == "Counsel accepted price wording"
    assert not ctx.matters.get(MATTER).get("approved_artifact_path")


def test_partial_revision_save_rolls_back_then_retries_once(app_context, monkeypatch):
    ctx = app_context
    item = draft(ctx)
    before = ctx.vault.resolve(item["vault_path"]).read_bytes()
    original = ctx.matters.append_event
    def fail(*args, **kwargs):
        raise OSError("Event write failed")
    monkeypatch.setattr(ctx.matters, "append_event", fail)
    with pytest.raises(OSError):
        revise(ctx, item, "New useful text")
    assert ctx.vault.resolve(item["vault_path"]).read_bytes() == before
    monkeypatch.setattr(ctx.matters, "append_event", original)
    result = revise(ctx, item, "New useful text")
    assert len(result["artifact"]["version_changes"]) == 1
    assert revise(ctx, item, "New useful text")["changed_paths"] == []


def test_template_snapshots_preview_and_custom_outputs(app_context):
    ctx = app_context
    template = {"template_id": "memo-custom", "output_type": "custom-output", "revision": "v1", "content_hash": "abc",
        "revision_path": "00_System/skills/revisions/v1.md", "instructions_snapshot": "Use three short sections.",
        "section_outline_snapshot": "Answer / Conditions / Next", "defaults_snapshot": {"audience": "CEO"},
        "overrides": {"length": "100 words"}, "state": "applied"}
    expected = deepcopy(template)
    item = draft(ctx, template_use=template, output_type="custom-output", preview=True)
    template["instructions_snapshot"] = "Later template change"
    assert all(item["artifact"]["template_use"][k] == v for k, v in expected.items())
    assert item["artifact"]["preview"]
    kept = ctx.work_products.keep_preview(MATTER, item["vault_path"], expected_revision=item["artifact"]["revision"])
    assert kept["preview"] is False
    failed = draft(ctx, title="Decision record draft", template_use={"template_id": "broken"}, output_type="decision-record")
    assert failed["artifact"]["template_use"]["state"] == "failed"
    assert ctx.vault.read_document(failed["vault_path"])["content"].startswith("Client")
    assert not ctx.matters.get(MATTER).get("approved_artifact_path")


def test_fact_reassessment_offer_decline_and_explicit_revision(app_context):
    ctx = app_context
    item = draft(ctx)
    path = item["vault_path"]
    before = ctx.vault.resolve(path).read_bytes()
    with pytest.raises(ValueError, match="cannot change documents"):
        revise(ctx, item, "New analysis", workspace_action="reassess_changed_facts")
    workspace_path = f"{ctx.matters.matter_path(MATTER)}/workspace.md"
    offer = {"offer_id": "offer-1", "artifact_path": path, "base_revision": item["artifact"]["revision"],
        "reason": "Reported timing changed", "state": "declined"}
    ctx.vault.write_markdown(workspace_path, "", {"update_offers": [offer], "unknown": "preserve"})
    with pytest.raises(WorkspaceConflict):
        revise(ctx, item, "New analysis", update_offer_id="offer-1")
    assert ctx.vault.resolve(path).read_bytes() == before
    offer["state"] = "offered"
    offer["offer_id"] = "offer-2"
    ctx.vault.update_markdown(workspace_path, metadata_updates={"update_offers": [offer]})
    revise(ctx, item, "New analysis", key="explicit-update:2", update_offer_id="offer-2", reason=offer["reason"])
    workspace = ctx.vault.read_document(workspace_path)
    assert workspace["metadata"]["update_offers"][0]["state"] == "accepted"
    assert workspace["metadata"]["unknown"] == "preserve"


def test_export_identity_saved_revision_modes_and_native_objects(tmp_path):
    vault = VaultService(tmp_path)
    exports, reviews = DocumentExportService(vault), DocumentReviewService(vault)
    for path, text in (("matter-a/memo.md", "Client pays ten dollars."), ("matter-b/memo.md", "Other client pays five dollars.")):
        vault.write_markdown(path, text)
    first, _ = exports.export("matter-a/memo.md", "docx")
    first_bytes = vault.resolve(first).read_bytes()
    second, _ = exports.export("matter-b/memo.md", "docx")
    assert first != second and vault.resolve(first).read_bytes() == first_bytes
    reviews.get("matter-a/memo.md")
    reviews.propose_agent_revision("matter-a/memo.md", "Client pays twenty dollars.")
    state = reviews.get("matter-a/memo.md")
    reviews.apply("matter-a/memo.md", DocumentReviewAction(action="add_comment", quote="Client", body="Retain defined term",
        author_name="Counsel", author_id="lawyer", anchor_start=0, anchor_end=6))
    marked, _ = exports.export("matter-a/memo.md", "docx", mode="markup")
    with ZipFile(vault.resolve(marked)) as archive:
        xml = archive.read("word/document.xml")
        assert b"<w:ins " in xml and b"<w:del " in xml
        assert b"Retain defined term" in archive.read("word/comments.xml")
    clean, _ = exports.export("matter-a/memo.md", "docx", mode="accepted_text")
    clean_text = "\n".join(p.text for p in Document(vault.resolve(clean)).paragraphs)
    assert "ten" in clean_text and "twenty" not in clean_text
    pdf, _ = exports.export("matter-a/memo.md", "pdf", mode="markup")
    assert any(page.get("/Annots") for page in PdfReader(vault.resolve(pdf)).pages)
    historical = state["history"][-1]["revision_path"]
    old, _ = exports.export("matter-a/memo.md", "docx", revision_path=historical,
        expected_revision=state["history"][-1]["before_revision"])
    assert "ten" in " ".join(p.text for p in Document(vault.resolve(old)).paragraphs)
    with pytest.raises(ValueError, match="does not belong"):
        exports.export("matter-b/memo.md", "docx", revision_path=historical)


def test_export_failure_leaves_draft_and_prior_export_intact(tmp_path, monkeypatch):
    vault = VaultService(tmp_path)
    service = DocumentExportService(vault)
    vault.write_markdown("memo.md", "See [the source](https://example.org/rule).")
    exported, _ = service.export("memo.md", "docx")
    text = " ".join(p.text for p in Document(vault.resolve(exported)).paragraphs)
    assert "https://example.org/rule" in text
    before = vault.resolve("memo.md").read_bytes()
    export_before = vault.resolve(exported).read_bytes()
    monkeypatch.setattr(service, "_docx", lambda *args: (_ for _ in ()).throw(OSError("export failed")))
    with pytest.raises(OSError):
        service.export("memo.md", "docx")
    assert vault.resolve("memo.md").read_bytes() == before
    assert vault.resolve(exported).read_bytes() == export_before


def test_outside_counsel_packet_retries_and_only_selected_original_exports(app_context):
    ctx = app_context
    root = ctx.matters.matter_path(MATTER)
    original = f"{root}/source-documents/agreement.pdf"
    unused = f"{root}/source-documents/internal.txt"
    ctx.vault.write_bytes(original, b"Original supplied PDF bytes")
    ctx.vault.write_bytes(unused, b"Internal notes")
    rev = hashlib.sha256(ctx.vault.resolve(original).read_bytes()).hexdigest()
    attachments = [{"path": original, "title": "Agreement", "revision": rev, "relevance": "Governs custody",
        "selected": True, "reviewed_revision": rev}, {"path": unused, "title": "Internal notes", "revision": "unused",
        "relevance": "Internal drafting only", "selected": False}]
    kwargs = dict(title="Counsel brief", brief_content="Question: custody. Reported date remains uncertain.",
        cover_email_content="Please advise on custody under the attached agreement.", attachments=attachments, source_action_key="packet:1",
        draft_context={"source_run_id": "RUN-outside-counsel"})
    packet = ctx.work_products.prepare_outside_counsel_packet(MATTER, **kwargs)
    assert ctx.work_products.prepare_outside_counsel_packet(MATTER, **kwargs) == packet
    for artifact in (packet["brief"], packet["cover_email"]):
        assert artifact["source_run_id"] == "RUN-outside-counsel"
        assert ctx.vault.read_document(artifact["path"])["metadata"]["source_run_id"] == "RUN-outside-counsel"
    result = ctx.work_products.export_outside_counsel_packet(MATTER, packet["brief"]["path"],
        reviewed_brief_revision=ctx.document_reviews.revision(packet["brief"]["path"]),
        reviewed_cover_revision=ctx.document_reviews.revision(packet["cover_email"]["path"]),
        attachments=attachments, output_format="docx", mode="markup", document_exports=ctx.document_exports)
    assert all(item["export_state"] == "exported" for item in result["export_results"])
    assert result["attachments"][1]["export_state"] == "not_exported"
    assert ctx.vault.resolve(result["attachments"][0]["output_path"]).read_bytes() == ctx.vault.resolve(original).read_bytes()


def test_action_keys_cannot_redirect_between_artifacts(app_context):
    ctx = app_context
    first = draft(ctx, title="First")
    second = draft(ctx, title="Second")
    revise(ctx, first, "First revised", key="one-action")
    untouched = ctx.vault.resolve(second["vault_path"]).read_bytes()
    with pytest.raises(WorkspaceConflict) as error:
        revise(ctx, second, "Second revised", key="one-action")
    assert error.value.detail["code"] == "action_key_conflict"
    assert ctx.vault.resolve(second["vault_path"]).read_bytes() == untouched


def test_template_preview_preserves_existing_work_state(app_context):
    ctx = app_context
    item = draft(ctx)
    matter_path = f"{ctx.matters.matter_path(MATTER)}/matter.md"
    before = ctx.vault.resolve(matter_path).read_bytes()
    draft_before = ctx.vault.resolve(item["vault_path"]).read_bytes()
    preview = draft(ctx, title="Preview", preview=True, recommendation_content="This must stay out of current recommendations")
    assert ctx.vault.resolve(matter_path).read_bytes() == before
    assert ctx.vault.resolve(item["vault_path"]).read_bytes() == draft_before
    assert preview["artifact"]["preview"] is True


def test_packet_partial_creation_retry_and_partial_export(app_context, monkeypatch):
    ctx = app_context
    original_create = ctx.work_products.create_draft
    def fail_cover(*args, **kwargs):
        if kwargs["source_action_key"].endswith(":cover"):
            raise OSError("Cover save failed")
        return original_create(*args, **kwargs)
    kwargs = dict(title="Brief", brief_content="Necessary context", cover_email_content="Please advise.",
        attachments=[], source_action_key="partial:packet")
    monkeypatch.setattr(ctx.work_products, "create_draft", fail_cover)
    with pytest.raises(OSError):
        ctx.work_products.prepare_outside_counsel_packet(MATTER, **kwargs)
    monkeypatch.setattr(ctx.work_products, "create_draft", original_create)
    packet = ctx.work_products.prepare_outside_counsel_packet(MATTER, **kwargs)
    products = [ctx.vault.read_markdown(ctx.vault.relative(p)) for p in ctx.vault.iter_files(ctx.matters.matter_path(MATTER), {".md"})]
    assert sum(d["metadata"].get("source_action_key") == "partial:packet:brief" for d in products) == 1
    original_export = ctx.document_exports.export
    def fail_one(path, *args, **kwargs):
        if path == packet["cover_email"]["path"]:
            raise OSError("Cover export failed")
        return original_export(path, *args, **kwargs)
    monkeypatch.setattr(ctx.document_exports, "export", fail_one)
    result = ctx.work_products.export_outside_counsel_packet(MATTER, packet["brief"]["path"],
        reviewed_brief_revision=ctx.document_reviews.revision(packet["brief"]["path"]),
        reviewed_cover_revision=ctx.document_reviews.revision(packet["cover_email"]["path"]),
        attachments=[], output_format="docx", mode="markup", document_exports=ctx.document_exports)
    assert [r["export_state"] for r in result["export_results"]] == ["exported", "failed"]
    assert ctx.vault.read_document(packet["brief"]["path"])["content"] == "Necessary context\n"
    assert ctx.vault.read_document(packet["cover_email"]["path"])["content"] == "Please advise.\n"


def test_mixed_accepted_rejected_and_pending_export(tmp_path):
    vault = VaultService(tmp_path)
    reviews, exports = DocumentReviewService(vault), DocumentExportService(vault)
    vault.write_markdown("memo.md", "Client pays ten dollars.")
    reviews.propose_agent_revision("memo.md", "Client pays twenty dollars.")
    change = reviews.get("memo.md")["changes"][0]["change_id"]
    reviews.apply("memo.md", DocumentReviewAction(action="accept_change", change_id=change))
    reviews.propose_agent_revision("memo.md", "Client pays thirty dollars.")
    change = reviews.get("memo.md")["changes"][0]["change_id"]
    reviews.apply("memo.md", DocumentReviewAction(action="reject_change", change_id=change))
    reviews.propose_agent_revision("memo.md", "Customer pays twenty dollars promptly.")
    path, _ = exports.export("memo.md", "docx", mode="accepted_text")
    text = " ".join(p.text for p in Document(vault.resolve(path)).paragraphs)
    assert "Client pays twenty dollars." in text
    assert "Customer" not in text and "promptly" not in text and "thirty" not in text


def test_reopen_lists_named_artifacts_and_preserved_proposals(app_context):
    ctx = app_context
    memo = draft(ctx, title="Memo")
    clause = draft(ctx, title="Clause")
    with pytest.raises(WorkspaceConflict):
        revise(ctx, memo, "Retained new proposal", local_draft_snapshot="Local draft text")
    reopened = {item["path"]: item for item in ctx.work_products.list_drafts(MATTER)}
    assert memo["vault_path"] in reopened and clause["vault_path"] in reopened
    assert len(reopened[memo["vault_path"]]["proposal_paths"]) == 1
    assert len(reopened[clause["vault_path"]]["proposal_paths"]) == 0


def test_supplied_clause_copy_preserves_original_comments_claims_and_terms(app_context):
    ctx = app_context
    original = f"{ctx.matters.matter_path(MATTER)}/source-documents/clause.extracted.md"
    ctx.vault.write_markdown(original, "Client means Beacon. Client pays ten dollars.", {"claim_ids": ["CLM-1"], "source_id": "SRC-1"})
    ctx.document_reviews.apply(original, DocumentReviewAction(action="add_comment", quote="Client", body="Do not rename this defined term",
        author_name="Counsel", author_id="lawyer", anchor_start=0, anchor_end=6))
    before = ctx.vault.resolve(original).read_bytes()
    source = ctx.vault.read_document(original)
    item = draft(ctx, title="Revised clause", content="Client means Beacon. Client pays twenty dollars.",
        source_document_path=original, source_revision=DocumentReviewService.content_revision(source["content"]),
        source_review_revision=ctx.document_reviews.revision(original))
    assert ctx.vault.resolve(original).read_bytes() == before
    assert item["artifact"]["claim_ids"] == ["CLM-1"]
    assert item["artifact"]["pending_review"]
    review = ctx.document_reviews.get(item["vault_path"])
    assert review["comments"][0]["entries"][0]["body"] == "Do not rename this defined term"
    assert ctx.vault.read_document(item["vault_path"])["metadata"]["supplied_original"]["metadata"]["source_id"] == "SRC-1"


@pytest.mark.parametrize("legacy", [False, True])
def test_unkept_preview_never_becomes_current_by_fallback(app_context, legacy):
    ctx = app_context
    root = ctx.matters.matter_path(MATTER)
    legacy_path = f"{root}/work-product.md"
    if ctx.vault.exists(legacy_path):
        ctx.vault.resolve(legacy_path).unlink()
    if legacy:
        ctx.vault.write_markdown(legacy_path, "Existing legacy advice")
    ctx.vault.update_markdown(f"{root}/matter.md", metadata_updates={"current_work_product_draft_path": None})
    item = draft(ctx, title="Unkept preview", preview=True)
    current = ctx.work_products.current_draft(MATTER)
    assert (current["path"] if current else None) == (legacy_path if legacy else None)
    assert ctx.work_products.current_draft(MATTER, legacy_fallback=False) is None
    kept = ctx.work_products.keep_preview(MATTER, item["vault_path"], expected_revision=item["artifact"]["revision"])
    assert kept["preview"] is False
    assert ctx.work_products.current_draft(MATTER)["path"] == item["vault_path"]


@pytest.mark.parametrize("change", ["insert", "replace", "delete"])
@pytest.mark.parametrize("format", ["docx", "pdf"])
def test_accepted_export_sources_follow_accepted_citations_only(tmp_path, change, format):
    vault = VaultService(tmp_path)
    reviews, exports = DocumentReviewService(vault), DocumentExportService(vault)
    approved = "https://approved.example/rule"
    pending = "https://pending.example/unapproved"
    baseline = "Approved sentence." if change == "insert" else f"Approved sentence. [Authority]({approved})"
    proposal = (baseline + f" [New lead]({pending})" if change == "insert" else
        f"Approved sentence. [New lead]({pending})" if change == "replace" else "Approved sentence.")
    vault.write_markdown("memo.md", baseline)
    reviews.propose_agent_revision("memo.md", proposal)
    paths = {mode: exports.export("memo.md", format, mode=mode)[0] for mode in ("accepted_text", "markup")}
    def extract(path):
        if format == "pdf":
            return " ".join(p.extract_text() for p in PdfReader(vault.resolve(path)).pages)
        with ZipFile(vault.resolve(path)) as archive:
            return archive.read("word/document.xml").decode()
    accepted, markup = extract(paths["accepted_text"]), extract(paths["markup"])
    assert pending not in accepted
    if change != "insert":
        assert approved in accepted
    if change != "delete":
        assert pending in markup
    if change != "insert":
        assert approved in markup


def test_preview_revision_keeps_current_matter_and_recommendation_unchanged(app_context):
    ctx = app_context
    real = draft(ctx, title="Real advice")
    preview = draft(ctx, title="Preview advice", preview=True)
    root = ctx.matters.matter_path(MATTER)
    before = {path: ctx.vault.resolve(path).read_bytes() for path in
        [real["vault_path"], f"{root}/matter.md", f"{root}/recommendations.md", f"{root}/dossier.md"]}
    revise(ctx, preview, "Edited preview", recommendation_content="Must not become working advice")
    for path, data in before.items():
        assert ctx.vault.resolve(path).read_bytes() == data
    assert ctx.work_products.current_draft(MATTER)["path"] == real["vault_path"]
    current = ctx.work_products.reference(MATTER, preview["vault_path"])
    ctx.work_products.keep_preview(MATTER, preview["vault_path"], expected_revision=current["revision"])
    assert ctx.work_products.current_draft(MATTER)["path"] == preview["vault_path"]


@pytest.mark.parametrize("mode", ["accepted_text", "markup"])
@pytest.mark.parametrize("format", ["docx", "pdf"])
def test_reviewed_exports_render_markdown_and_preserve_review_anchors(tmp_path, mode, format):
    from docx.oxml.ns import qn
    vault = VaultService(tmp_path)
    reviews, exports = DocumentReviewService(vault), DocumentExportService(vault)
    baseline = "## Recommendation\n\nUse **conditional approval**.\n\n- Confirm timing.\n\nLawyer note: retain safeguards.\n"
    vault.write_markdown("memo.md", baseline)
    reviews.propose_agent_revision("memo.md", baseline.replace("conditional approval", "limited approval").replace("timing", "custody"))
    change = next(c for c in reviews.get("memo.md")["changes"] if c["new_text"] == "limited")
    reviews.apply("memo.md", DocumentReviewAction(action="accept_change", change_id=change["change_id"]))
    current = vault.read_document("memo.md")["content"]
    start = current.index("retain safeguards")
    reviews.apply("memo.md", DocumentReviewAction(action="add_comment", quote="retain safeguards",
        anchor_start=start, anchor_end=start + len("retain safeguards"), body="Keep the lawyer's note",
        author_name="Counsel", author_id="lawyer"))
    before = vault.resolve("memo.md").read_bytes()
    path, _ = exports.export("memo.md", format, mode=mode)
    assert vault.resolve("memo.md").read_bytes() == before
    if format == "docx":
        doc = Document(vault.resolve(path))
        with ZipFile(vault.resolve(path)) as archive:
            from lxml import etree
            root = etree.fromstring(archive.read("word/document.xml"))
            text = "".join(root.itertext())
            assert b"Keep the lawyer's note" in archive.read("word/comments.xml")
            anchored, in_comment = [], False
            for element in root.iter():
                if element.tag == qn("w:commentRangeStart"):
                    in_comment = True
                elif element.tag == qn("w:commentRangeEnd"):
                    in_comment = False
                elif in_comment and element.tag == qn("w:t"):
                    anchored.append(element.text or "")
            assert "".join(anchored) == "retain safeguards"
            if mode == "markup":
                assert root.find(".//" + qn("w:ins")) is not None
                assert root.find(".//" + qn("w:del")) is not None
        assert doc.paragraphs[0].style.name == "Heading 2"
        bullet = next(p for p in doc.paragraphs if "Confirm" in p.text)
        assert bullet.style.name == "List Bullet"
        assert not bullet.text.startswith("- ")
        assert any(run.bold and "limited" in run.text for p in doc.paragraphs for run in p.runs)
    else:
        reader = PdfReader(vault.resolve(path))
        text = "".join(p.extract_text() for p in reader.pages)
        comments = [a.get_object().get("/Contents", "") for p in reader.pages for a in p.get("/Annots", [])]
        assert any("Keep the lawyer's note" in c for c in comments)
    assert "##" not in text and "**" not in text
    assert "limited" in text and "Lawyer" in text and "safeguards" in text
    assert "timing" in text
    if mode == "accepted_text":
        assert "custody" not in text


@pytest.mark.parametrize("format", ["docx", "pdf"])
def test_accepted_export_keeps_comment_when_its_pending_text_is_omitted(tmp_path, format):
    vault = VaultService(tmp_path)
    reviews, exports = DocumentReviewService(vault), DocumentExportService(vault)
    vault.write_markdown("memo.md", "## Advice\n\nApproved text.")
    reviews.propose_agent_revision("memo.md", "## Advice\n\nApproved text. **Pending words.**")
    text = vault.read_document("memo.md")["content"]
    start = text.index("Pending words.")
    reviews.apply("memo.md", DocumentReviewAction(action="add_comment", quote="Pending words.",
        anchor_start=start, anchor_end=start + len("Pending words."), body="Explain this proposal",
        author_name="Counsel", author_id="lawyer"))
    path, _ = exports.export("memo.md", format, mode="accepted_text")
    if format == "docx":
        with ZipFile(vault.resolve(path)) as archive:
            comments = archive.read("word/comments.xml").decode()
        assert "Pending words." not in " ".join(p.text for p in Document(vault.resolve(path)).paragraphs)
    else:
        reader = PdfReader(vault.resolve(path))
        comments = " ".join(a.get_object().get("/Contents", "") for p in reader.pages for a in p.get("/Annots", []))
        assert "Pending words." not in " ".join(p.extract_text() for p in reader.pages)
    assert "Explain this proposal" in comments
    assert "Selected text is not present in this export" in comments


@pytest.mark.parametrize("mode", ["accepted_text", "markup"])
@pytest.mark.parametrize("format", ["docx", "pdf"])
def test_exports_preserve_literal_identifiers_urls_code_and_nested_emphasis(tmp_path, mode, format):
    vault = VaultService(tmp_path)
    exports = DocumentExportService(vault)
    content = ("## Advice\n\nPreserve customer_account_id and https://example.test/a_b_c.\n\n"
        "***Important condition*** applies. Keep foo__bar__baz.\n\n"
        "Use `**literal_code** and customer_account_id` and [source_label](https://example.test/d_e_f).\n\n"
        "Escaped \\*literal\\* remains. **Bold with _nested italic_**.\n\n"
        "```text\n## literal_header\n**literal_code**\n```\n")
    vault.write_markdown("memo.md", content)
    before = vault.resolve("memo.md").read_bytes()
    path, _ = exports.export("memo.md", format, mode=mode)
    if format == "docx":
        doc = Document(vault.resolve(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        combined = next(r for p in doc.paragraphs for r in p.runs if "Important condition" in r.text)
        assert combined.bold and combined.italic
        code = next(r for p in doc.paragraphs for r in p.runs if "literal_code" in r.text)
        assert code.font.name == "Consolas"
    else:
        text = " ".join("".join(p.extract_text() for p in PdfReader(vault.resolve(path)).pages).split())
    for literal in ("customer_account_id", "https://example.test/a_b_c.", "foo__bar__baz", "**literal_code**",
                    "source_label", "https://example.test/d_e_f", "*literal*", "## literal_header"):
        assert literal in text
    assert "***Important" not in text and "*Important" not in text
    assert "**Bold" not in text and "_nested" not in text
    assert vault.resolve("memo.md").read_bytes() == before
