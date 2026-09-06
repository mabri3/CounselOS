from __future__ import annotations

import io
import hashlib
import json
import re
import unicodedata
from datetime import UTC, datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from app.services.document_review import DocumentReviewService, review_segments
from app.services.dossier import serialized
from app.services.vault import VaultService
from app.services.workspace_evidence import WorkspaceEvidenceService


def markdown_to_plain(markdown: str) -> str:
    keep, _ = _markdown_mask(markdown)
    return "".join(character for offset, character in enumerate(markdown) if keep[offset]).strip() + "\n"


def _markdown_mask(text: str) -> tuple[list[bool], list[int]]:
    """Identify display characters while preserving literal text and review offsets."""
    keep, styles, protected = [True] * len(text), [0] * len(text), [False] * len(text)

    def remove(start: int, end: int) -> None:
        keep[start:end] = [False] * (end - start)

    def protect(start: int, end: int, style: int = 0) -> None:
        protected[start:end] = [True] * (end - start)
        for offset in range(start, end):
            styles[offset] |= style

    # Code contents are literal, including Markdown-looking text and URL punctuation.
    fence, offset = "", 0
    for line in text.splitlines(keepends=True):
        marker = re.match(r"^[ \t]{0,3}(`{3,}|~{3,})", line)
        if marker and (not fence or marker.group(1)[0] == fence[0] and len(marker.group(1)) >= len(fence)):
            fence = "" if fence else marker.group(1)
            protect(offset, offset + len(line))
            remove(offset, offset + len(line.rstrip("\r\n")))
        elif fence:
            protect(offset, offset + len(line), 4)
        offset += len(line)
    for match in re.finditer(r"(?<!`)(`+)(?!`)(.+?)(?<!`)\1(?!`)", text, re.DOTALL):
        if any(protected[match.start():match.end()]):
            continue
        protect(*match.span(), 4)
        remove(match.start(), match.start(2))
        remove(match.end(2), match.end())
    for match in re.finditer(r"\\([!\"#$%&'()*+,\-./:;<=>?@\[\]\\^_`{|}~])", text):
        if not protected[match.start()]:
            remove(match.start(), match.start(1))
            protect(*match.span(1))
    for match in re.finditer(r"!?\[([^\]\n]+)\]\([^\n)]*\)", text):
        if protected[match.start()]:
            continue
        remove(match.start(), match.start(1))
        remove(match.end(1), match.end())
        protect(match.end(1), match.end())
    for match in re.finditer(r"https?://[^\s<>]+", text):
        if not protected[match.start()]:
            end = match.end()
            # A surrounding emphasis delimiter is outside the literal URL.
            prefix = re.search(r"([*_]+)$", text[:match.start()])
            if prefix and match.group().endswith(prefix.group(1)):
                end -= len(prefix.group(1))
            protect(match.start(), end)
    for match in re.finditer(r"^[ \t]{0,3}(?:#{1,6}[ \t]+|>[ \t]?|[-+*][ \t]+|\d+[.)][ \t]+)", text, re.MULTILINE):
        if not any(protected[match.start():match.end()]):
            remove(*match.span())

    def punctuation(character: str) -> bool:
        return unicodedata.category(character).startswith(("P", "S"))

    # Delimiter flanking rules keep intraword underscores literal. Consume paired runs
    # from the inside out, so triple and nested emphasis retain both font styles.
    openers: list[dict[str, Any]] = []
    for match in re.finditer(r"\*+|_+", text):
        if any(protected[match.start():match.end()]) or not all(keep[match.start():match.end()]):
            continue
        previous = text[match.start() - 1] if match.start() else " "
        following = text[match.end()] if match.end() < len(text) else " "
        left = not following.isspace() and (not punctuation(following) or previous.isspace() or punctuation(previous))
        right = not previous.isspace() and (not punctuation(previous) or following.isspace() or punctuation(following))
        marker, length = match.group()[0], len(match.group())
        can_open = left and (marker == "*" or not right or punctuation(previous))
        can_close = right and (marker == "*" or not left or punctuation(following))
        used = 0
        while can_close and used < length:
            selected = next((index for index in range(len(openers) - 1, -1, -1)
                if openers[index]["marker"] == marker and not (
                    (openers[index]["can_close"] or can_open)
                    and (openers[index]["length"] + length) % 3 == 0
                    and (openers[index]["length"] % 3 or length % 3))), None)
            if selected is None:
                break
            opener = openers[selected]
            amount = 2 if min(opener["remaining"], length - used) >= 2 else 1
            left_end = opener["start"] + opener["remaining"]
            right_start = match.start() + used
            remove(left_end - amount, left_end)
            remove(right_start, right_start + amount)
            for position in range(left_end, right_start):
                styles[position] |= 1 if amount == 2 else 2
            opener["remaining"] -= amount
            used += amount
            del openers[selected + 1:]
            if not opener["remaining"]:
                openers.pop()
        if can_open and used < length:
            openers.append({"start": match.start() + used, "remaining": length - used,
                "length": length, "marker": marker, "can_close": can_close})
    return keep, styles


def _plain_review_segments(segments: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[int]]:
    # Parse both sides together: Markdown delimiters often span several tracked runs.
    current = "".join(item["text"] for item in segments if item.get("kind") != "delete")
    original = "".join(item["text"] for item in segments if item.get("kind") != "insert")
    current_keep, current_styles = _markdown_mask(current)
    old_keep, old_styles = _markdown_mask(original)
    current_offset = old_offset = 0
    output = []
    for segment in segments:
        deleted = segment.get("kind") == "delete"
        keep, styles = (old_keep, old_styles) if deleted else (current_keep, current_styles)
        offset = old_offset if deleted else current_offset
        fragment, active_style = "", None
        for index, character in enumerate(segment["text"]):
            if not keep[offset + index]:
                continue
            style = styles[offset + index]
            if active_style is not None and style != active_style:
                output.append({**segment, "text": fragment, "export_style": active_style})
                fragment = ""
            active_style = style
            fragment += character
        if fragment:
            output.append({**segment, "text": fragment, "export_style": active_style or 0})
        if segment.get("kind") != "delete":
            current_offset += len(segment["text"])
        if segment.get("kind") != "insert":
            old_offset += len(segment["text"])
    offsets = [0]
    for retained in current_keep:
        offsets.append(offsets[-1] + int(retained))
    return output, offsets


def _export_comments(comments: list[dict[str, Any]], source: str, chosen: str,
                     offsets: list[int]) -> list[dict[str, Any]]:
    blocks = SequenceMatcher(None, source, chosen, autojunk=False).get_matching_blocks() if source != chosen else []
    result = []
    for comment in comments:
        item = dict(comment)
        start, end = _integer(item.get("anchor_start")), _integer(item.get("anchor_end"))
        quote = str(item.get("quote") or "")
        if start is None or end is None or not (0 <= start < end <= len(source)) or source[start:end] != quote:
            start = source.find(quote) if quote else -1
            end = start + len(quote)
        if source != chosen and start >= 0:
            block = next((block for block in blocks if block.a <= start and end <= block.a + block.size), None)
            start = block.b + start - block.a if block else (chosen.find(quote) if quote and chosen.count(quote) == 1 else -1)
            end = start + len(quote)
        if 0 <= start < end <= len(chosen):
            item.update({"anchor_start": offsets[start], "anchor_end": offsets[end]})
        else:
            item.update({"anchor_start": -1, "anchor_end": -1,
                "export_anchor_note": f"Selected text is not present in this export: {quote}"})
        item["quote"] = markdown_to_plain(quote).strip() if quote else ""
        result.append(item)
    return result


class DocumentExportService:
    def __init__(self, vault: VaultService):
        self.vault = vault

    @serialized
    def export(self, path: str, output_format: str, *, mode: str = "markup",
               expected_revision: str | None = None, expected_review_revision: str | None = None,
               revision_path: str | None = None) -> tuple[str, str]:
        if mode not in {"markup", "accepted_text"}:
            raise ValueError("Export mode must be markup or accepted_text.")
        selected_path = revision_path or path
        document = self.vault.read_document(selected_path)
        if revision_path and (document.get("metadata", {}).get("source_path") != path
                              or not document.get("metadata", {}).get("immutable")):
            raise ValueError("The selected saved version does not belong to this document.")
        revision = DocumentReviewService.content_revision(document["content"])
        if expected_revision is not None and expected_revision != revision:
            raise ValueError("The selected saved version changed. Refresh before exporting.")
        review_revision = (document["metadata"].get("saved_review_revision") if revision_path
                           else DocumentReviewService(self.vault).revision(path))
        if expected_review_revision is not None and review_revision != expected_review_revision:
            raise ValueError("The selected review changed. Refresh before exporting.")
        if document.get("kind") != "markdown":
            raise ValueError("Only Markdown documents can be exported.")
        review = document.get("metadata", {}).get("review", {})
        review = review if isinstance(review, dict) else {}
        current = document["content"]
        stored_segments = review.get("segments")
        if review.get("version") == 2 and isinstance(stored_segments, list):
            segments = [dict(item) for item in stored_segments if isinstance(item, dict)]
        else:
            baseline = str(review.get("baseline") or document["content"])
            segments = review_segments(baseline, current)
        if mode == "accepted_text":
            segments = DocumentReviewService.accepted_segments(segments)
            current = "".join(item["text"] for item in segments)
            # Accepted text has no review boundaries. Join it before export so prior token-level
            # revisions cannot split public URLs into disconnected Word runs or PDF text objects.
            segments = [{"kind": "equal", "text": current}]
        effective_markdown = current if mode == "accepted_text" else document["content"]
        original_markdown = "".join(item["text"] for item in segments if item.get("kind") != "insert")
        segments, offsets = _plain_review_segments(segments)
        current = "".join(item["text"] for item in segments if item.get("kind") != "delete")
        comments = _export_comments(
            [dict(item) for item in review.get("comments", []) if isinstance(item, dict)],
            document["content"], effective_markdown, offsets)
        # Derive sources and styles from the selected text, so a pending citation cannot leak
        # into an accepted-text export through the source appendix.
        source_markdown = effective_markdown + ("\n" + original_markdown if mode == "markup" else "")
        links = list(dict.fromkeys(re.findall(r"\[[^\]]+\]\((https?://[^)]+)\)", source_markdown)))
        visible = "".join(item["text"] for item in segments)
        missing_links = [link for link in links if link not in visible]
        if missing_links:
            segments.append({"kind": "equal", "text": "\n\nSource links\n" + "\n".join(missing_links) + "\n"})
        document_references = self._document_reference_appendix(source_markdown)
        if document_references:
            segments.append({"kind": "equal", "text": document_references})
        claim_appendix = self._claim_support_appendix(document.get("metadata", {}), source_markdown)
        if claim_appendix:
            segments.append({"kind": "equal", "text": claim_appendix})
        stem = Path(path).name.removesuffix(".md").removesuffix(".extracted")
        if output_format == "docx":
            data = self._docx(segments, comments, document["name"], effective_markdown)
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif output_format == "pdf":
            data = self._pdf(segments, comments, current, effective_markdown)
            media_type = "application/pdf"
        else:
            raise ValueError("Export format must be docx or pdf.")
        identity = hashlib.sha256(path.encode("utf-8")).hexdigest()[:16]
        version = hashlib.sha256(json.dumps({"content": document["content"], "review": review,
            "claims": document.get("metadata", {}).get("claims", []),
            "mode": mode}, sort_keys=True, default=str).encode("utf-8")).hexdigest()[:20]
        output_path = f".exports/{identity}/{version}/{stem}-{mode}.{output_format}"
        self.vault.write_bytes(output_path, data)
        return output_path, media_type

    def _document_reference_appendix(self, selected_markdown: str) -> str:
        keep, _ = _markdown_mask(selected_markdown)
        references: list[tuple[str, str]] = []
        for match in re.finditer(r"(?<!!)\[([^\]\n]+)\]\(([^\n)]*)\)", selected_markdown):
            destination = match.group(2).strip()
            # Active Markdown link destinations are removed by the display mask. A match
            # inside inline or fenced code stays visible and is not a document reference.
            if not destination or any(keep[match.start(2):match.end(2)]):
                continue
            path, separator, locator = destination.partition("#")
            if re.match(r"^[A-Za-z][A-Za-z0-9+.-]*:", path) or path.startswith(("/", "\\")):
                continue
            safe_path = self._safe_vault_path(path)
            try:
                available = safe_path is not None and self.vault.resolve(safe_path).is_file()
            except (OSError, ValueError):
                available = False
            if not available:
                continue
            label = markdown_to_plain(match.group(1)).strip()
            reference = safe_path + (separator + locator if separator and locator else "")
            identity = (label, reference)
            if label and identity not in references:
                references.append(identity)
        if not references:
            return ""
        return "\n\nDocument references\n" + "\n".join(
            f"{label} — {destination}" for label, destination in references
        ) + "\n"

    def _claim_support_appendix(self, metadata: dict[str, Any], selected_markdown: str) -> str:
        raw_claims = metadata.get("claims", [])
        if not isinstance(raw_claims, list):
            return ""
        lines: list[str] = []
        seen_evidence: set[tuple[str, str, str, str]] = set()
        for raw in raw_claims:
            if not isinstance(raw, dict):
                continue
            claim_id = str(raw.get("claim_id") or "").strip()
            claim_text = str(raw.get("text") or "").strip()
            if not claim_id or not claim_text or claim_text not in selected_markdown:
                continue
            if not lines:
                lines = ["\n\nClaim support"]
            claim_revision = str(raw.get("claim_revision") or "").strip()
            output_revision = str(raw.get("output_revision") or metadata.get("output_revision") or "").strip()
            revision_text = ", ".join(value for value in (
                f"claim revision {claim_revision}" if claim_revision else "",
                f"output revision {output_revision}" if output_revision else "",
            ) if value)
            lines.append(f"\nClaim {claim_id}" + (f" ({revision_text})" if revision_text else "") + f": {claim_text}")
            applicability = raw.get("applicability")
            if isinstance(applicability, dict):
                actor = str(applicability.get("regulated_actor") or "").strip()
                jurisdiction = str(applicability.get("jurisdiction") or "").strip()
                explanation = str(applicability.get("explanation") or "").strip()
                if actor or jurisdiction or explanation:
                    lines.append("Applicability: " + "; ".join(value for value in (
                        f"actor: {actor}" if actor else "actor: unknown",
                        f"jurisdiction: {jurisdiction}" if jurisdiction else "jurisdiction: unknown",
                        explanation,
                    ) if value))
            evidence = raw.get("evidence", [])
            for item in evidence if isinstance(evidence, list) else []:
                if not isinstance(item, dict):
                    continue
                source_id = str(item.get("source_id") or "").strip()
                locator = str(item.get("locator") or "").strip()
                identity = (claim_id, claim_revision, source_id, locator)
                if not source_id or identity in seen_evidence:
                    continue
                seen_evidence.add(identity)
                state = str(item.get("support_state") or "unknown").replace("_", " ").title()
                label = str(item.get("source_label") or source_id).strip()
                url = WorkspaceEvidenceService._safe_public_url(item.get("url"))
                path = self._safe_vault_path(item.get("path"))
                location = url or path
                source_line = f"Source {source_id} — {label}"
                if locator:
                    source_line += f" — locator: {locator}"
                source_line += f" — status: {state}"
                if location:
                    source_line += f" — {location}"
                lines.append(source_line)
                excerpt = item.get("available_excerpt")
                if isinstance(excerpt, str) and excerpt:
                    lines.append("Exact available excerpt: " + excerpt)
                explanation = str(item.get("explanation") or "").strip()
                if explanation:
                    lines.append("Why it supports this claim: " + explanation)
            support_gap = str(raw.get("support_gap") or "").strip()
            if support_gap:
                lines.append("Support gap: " + support_gap)
        return "\n".join(lines) + ("\n" if lines else "")

    def _safe_vault_path(self, value: Any) -> str | None:
        if not isinstance(value, str) or not value.strip():
            return None
        try:
            return self.vault.relative(self.vault.resolve(value.strip()))
        except (OSError, ValueError):
            return None

    @serialized
    def export_original(self, path: str, *, expected_revision: str) -> tuple[str, str]:
        """Copy the selected original bytes; never export its extraction in its place."""
        import mimetypes
        data = self.vault.resolve(path).read_bytes()
        revision = hashlib.sha256(data).hexdigest()
        if revision != expected_revision:
            raise ValueError("The selected original file changed. Review its current version first.")
        identity = hashlib.sha256(path.encode("utf-8")).hexdigest()[:16]
        output_path = f".exports/{identity}/{revision[:20]}/{Path(path).name}"
        self.vault.write_bytes(output_path, data)
        return output_path, mimetypes.guess_type(path)[0] or "application/octet-stream"

    @staticmethod
    def _docx(
        segments: list[dict[str, Any]],
        comments: list[dict[str, Any]],
        title: str,
        markdown: str,
    ) -> bytes:
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.85)
        section.left_margin = section.right_margin = Inches(0.9)
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(11)
        document.core_properties.title = title
        document.core_properties.author = "Themis.ai"

        paragraph = document.add_paragraph()
        tracked_runs: list[tuple[Any, dict[str, Any]]] = []
        visible_runs: list[tuple[Any, int, int]] = []
        visible_offset = 0
        comment_boundaries = {
            boundary
            for comment in comments
            for boundary in (_integer(comment.get("anchor_start")), _integer(comment.get("anchor_end")))
            if boundary is not None and boundary >= 0
        }
        for segment in segments:
            pieces = segment["text"].split("\n")
            for index, piece in enumerate(pieces):
                if piece:
                    cuts = [0, len(piece)]
                    if segment["kind"] != "delete":
                        cuts.extend(boundary - visible_offset for boundary in comment_boundaries
                                    if visible_offset < boundary < visible_offset + len(piece))
                    cuts = sorted(set(cuts))
                    for left, right in zip(cuts, cuts[1:]):
                        run = paragraph.add_run(piece[left:right])
                        style = int(segment.get("export_style") or 0)
                        run.bold, run.italic = (True if style & 1 else None), (True if style & 2 else None)
                        if style & 4:
                            run.font.name = "Consolas"
                        tracked_runs.append((run, segment))
                        if segment["kind"] != "delete":
                            visible_runs.append((run, visible_offset + left, visible_offset + right))
                    if segment["kind"] != "delete":
                        visible_offset += len(piece)
                if index < len(pieces) - 1:
                    paragraph = document.add_paragraph()
                    if segment["kind"] != "delete":
                        visible_offset += 1

        for paragraph, style in zip(document.paragraphs, _block_styles(markdown), strict=False):
            if style:
                paragraph.style = style

        for comment in comments:
            quote = str(comment.get("quote") or "")
            anchor_start = _integer(comment.get("anchor_start"), -1)
            anchor_end = _integer(comment.get("anchor_end"), -1)
            target_runs = [run for run, start, end in visible_runs
                           if start < anchor_end and end > anchor_start]
            if not target_runs and quote:
                target = next((item for item in document.paragraphs if quote in item.text and item.runs), None)
                target_runs = list(target.runs) if target is not None else []
            if not target_runs:
                target = next((item for item in document.paragraphs if item.runs and item.text.strip()), None)
                target_runs = list(target.runs) if target is not None else []
            if target_runs:
                entries = _comment_entries(comment)
                lines = [comment["export_anchor_note"]] if comment.get("export_anchor_note") else []
                if comment.get("resolved"):
                    lines.append("Resolved")
                for entry in entries:
                    label = str(entry.get("author_name") or "Imported reviewer")
                    created_at = str(entry.get("created_at") or "")
                    lines.append(f"{label}{f' — {created_at}' if created_at else ''}: {entry.get('body', '')}")
                if not lines:
                    continue
                document.add_comment(
                    target_runs,
                    text="\n".join(lines),
                    author=str(entries[0].get("author_name") or "Themis.ai User") if entries else "Themis.ai User",
                )

        next_id = 1
        timestamp = datetime.now(UTC).isoformat().replace("+00:00", "Z")
        for run, segment in tracked_runs:
            kind = str(segment.get("kind") or "")
            if kind not in {"insert", "delete"}:
                continue
            element = run._r
            parent = element.getparent()
            if parent is None:
                continue
            index = parent.index(element)
            parent.remove(element)
            wrapper = OxmlElement("w:ins" if kind == "insert" else "w:del")
            wrapper.set(qn("w:id"), str(next_id))
            wrapper.set(qn("w:author"), str(segment.get("author_name") or segment.get("author_id") or "Themis.ai"))
            wrapper.set(qn("w:date"), str(segment.get("created_at") or timestamp))
            next_id += 1
            if kind == "delete":
                for text_node in element.findall(".//" + qn("w:t")):
                    text_node.tag = qn("w:delText")
            wrapper.append(element)
            parent.insert(index, wrapper)

        settings = document.settings._element
        if settings.find(qn("w:trackRevisions")) is None:
            settings.insert(0, OxmlElement("w:trackRevisions"))
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @classmethod
    def _pdf(
        cls,
        segments: list[dict[str, Any]],
        comments: list[dict[str, Any]],
        current: str,
        markdown: str,
    ) -> bytes:
        source = io.BytesIO()
        pdf = canvas.Canvas(source, pagesize=LETTER)
        width, height = LETTER
        margin = 54.0
        x = margin
        y = height - margin
        font = "Helvetica"
        font_size = 10.5
        leading = 15.0
        block_styles = _block_styles(markdown)
        list_labels = [match.group(1) + "." if (match := re.match(r"^\s*(\d+)[.)]\s+", line)) else "•"
                       for line in markdown.splitlines()]
        current_line = 0
        pdf.setFont(font, font_size)
        marks: list[dict[str, Any]] = []
        visible: list[dict[str, Any]] = []
        current_offset = 0

        def new_line(visible_line: bool) -> None:
            nonlocal x, y, page_number, current_line
            x = margin
            y -= leading
            if visible_line:
                current_line += 1
            if y < margin:
                pdf.showPage()
                pdf.setFont(font, font_size)
                y = height - margin
                page_number += 1

        page_number = 0
        for segment in segments:
            kind = segment["kind"]
            for token in re.findall(r"\n|[^\S\n]+|[^\s]+", segment["text"]):
                if token == "\n":
                    if kind != "delete":
                        current_offset += 1
                    new_line(kind != "delete")
                    continue
                style = block_styles[current_line] if current_line < len(block_styles) else ""
                inline_style = int(segment.get("export_style") or 0)
                active_font = ("Courier" if inline_style & 4 else "Helvetica-BoldOblique" if inline_style & 3 == 3
                    else "Helvetica-Bold" if style.startswith("Heading") or inline_style & 1
                    else "Helvetica-Oblique" if inline_style & 2 else font)
                active_size = 15.0 if style == "Heading 1" else 12.0 if style.startswith("Heading") else font_size
                pdf.setFont(active_font, active_size)
                if x == margin and token.strip() and style in {"List Bullet", "List Number"}:
                    label = list_labels[current_line] if current_line < len(list_labels) else "•"
                    pdf.drawString(x, y, label)
                    x += max(12, stringWidth(label, active_font, active_size) + 5)
                token_width = stringWidth(token, active_font, active_size)
                if token.strip() and x + token_width > width - margin:
                    new_line(False)
                if kind == "delete":
                    pdf.setFillColorRGB(0.72, 0.18, 0.14)
                elif kind == "insert":
                    pdf.setFillColorRGB(0.12, 0.43, 0.26)
                else:
                    pdf.setFillColorRGB(0.12, 0.12, 0.12)
                pdf.drawString(x, y, token)
                rect = (x, y - 2, x + token_width, y + active_size)
                if kind in {"insert", "delete"} and token.strip():
                    line_y = y - 1 if kind == "insert" else y + active_size * 0.35
                    pdf.setStrokeColorRGB(*( (0.12, 0.43, 0.26) if kind == "insert" else (0.72, 0.18, 0.14) ))
                    pdf.line(x, line_y, x + token_width, line_y)
                    marks.append({"page": page_number, "kind": kind, "rect": rect, "text": token})
                if kind != "delete":
                    visible.append({"page": page_number, "start": current_offset, "end": current_offset + len(token), "rect": rect})
                    current_offset += len(token)
                x += token_width
        pdf.save()

        reader = PdfReader(io.BytesIO(source.getvalue()))
        writer = PdfWriter()
        writer.clone_document_from_reader(reader)
        for mark in marks:
            subtype = "/Underline" if mark["kind"] == "insert" else "/StrikeOut"
            color = (0.12, 0.43, 0.26) if mark["kind"] == "insert" else (0.72, 0.18, 0.14)
            writer.add_annotation(mark["page"], cls._markup_annotation(
                subtype, mark["rect"], color,
                f"{'Inserted' if mark['kind'] == 'insert' else 'Deleted'}: {mark['text']}",
                "Themis.ai",
            ))
        for comment in comments:
            quote = str(comment.get("quote") or "")
            start = _integer(comment.get("anchor_start"))
            end = _integer(comment.get("anchor_end"))
            if start is None or end is None or start < 0 or end <= start or end > len(current):
                start = current.find(quote)
                end = start + len(quote)
            hits = (visible[:1] if comment.get("export_anchor_note") else
                    [item for item in visible if item["start"] < end and item["end"] > start])
            entries = _comment_entries(comment)
            contents = "\n".join(
                f"{entry.get('author_name') or 'Themis.ai User'}: {entry.get('body') or ''}" for entry in entries
            )
            if comment.get("export_anchor_note"):
                contents = comment["export_anchor_note"] + "\n" + contents
            for hit in hits:
                writer.add_annotation(hit["page"], cls._markup_annotation(
                    "/Text" if comment.get("export_anchor_note") else "/Highlight", hit["rect"], (0.96, 0.75, 0.2),
                    contents,
                    str(entries[0].get("author_name") or "Themis.ai User") if entries else "Themis.ai User",
                ))
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()

    @staticmethod
    def _markup_annotation(
        subtype: str,
        rect: tuple[float, float, float, float],
        color: tuple[float, float, float],
        contents: str,
        author: str,
    ) -> DictionaryObject:
        x1, y1, x2, y2 = rect
        return DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject(subtype),
            NameObject("/Rect"): ArrayObject([FloatObject(v) for v in rect]),
            NameObject("/QuadPoints"): ArrayObject([FloatObject(v) for v in (x1, y2, x2, y2, x1, y1, x2, y1)]),
            NameObject("/C"): ArrayObject([FloatObject(v) for v in color]),
            NameObject("/Contents"): TextStringObject(contents),
            NameObject("/T"): TextStringObject(author),
            NameObject("/F"): FloatObject(4),
        })


def _block_styles(markdown: str) -> list[str]:
    styles: list[str] = []
    fence = ""
    for line in markdown.splitlines():
        marker = re.match(r"^[ \t]{0,3}(`{3,}|~{3,})", line)
        if marker and (not fence or marker.group(1)[0] == fence[0] and len(marker.group(1)) >= len(fence)):
            fence = "" if fence else marker.group(1)
            styles.append("")
            continue
        if fence:
            styles.append("")
            continue
        heading = re.match(r"^\s{0,3}(#{1,6})\s+", line)
        if heading:
            styles.append(f"Heading {min(len(heading.group(1)), 3)}")
        elif re.match(r"^\s*[-+*]\s+", line):
            styles.append("List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            styles.append("List Number")
        elif re.match(r"^\s*>\s?", line):
            styles.append("Quote")
        else:
            styles.append("")
    return styles


def _comment_entries(comment: dict[str, Any]) -> list[dict[str, Any]]:
    entries = [dict(entry) for entry in comment.get("entries", []) if isinstance(entry, dict)]
    if entries:
        return entries
    body = str(comment.get("comment") or comment.get("body") or "")
    if not body:
        return []
    return [{"author_name": str(comment.get("author") or comment.get("author_name") or "Themis.ai User"),
             "body": body, "created_at": str(comment.get("created_at") or "")}]


def _integer(value: Any, default: int | None = None) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default
