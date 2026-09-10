from __future__ import annotations

import asyncio
import hashlib
import io
import logging
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException
from docx import Document as DocxDocument
from fastapi import UploadFile
from pypdf import PdfReader

from app.services.dossier import serialized
from app.services.index import IndexService
from app.services.matters import MatterService
from app.services.matter_paths import MatterPathPolicy
from app.services.vault import VaultService
from app.utils.paths import safe_filename
from app.utils.time import iso_now


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{WORD_NS}}}"
AUTHOR_PALETTE = ("#2F5597", "#7030A0", "#008272", "#A64B00", "#C0006F", "#5B6573", "#7A3E00", "#006B8F")
UPLOAD_READ_CHUNK_BYTES = 64 * 1024
DOCX_MAX_MEMBERS = 256
DOCX_MAX_XML_PART_BYTES = 8 * 1024 * 1024
DOCX_MAX_EXPANDED_BYTES = 50 * 1024 * 1024
DOCX_READ_CHUNK_BYTES = 64 * 1024

logger = logging.getLogger(__name__)


class DocxSecurityError(ValueError):
    """A DOCX package cannot be processed safely."""


class ExtractedText(str):
    """Text plus a truthful partial-extraction signal; legacy text callers work."""
    def __new__(cls, value: str, warning: str = ""):
        instance = super().__new__(cls, value)
        instance.warning = warning
        return instance


class IngestionService:
    TEXT_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}
    IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic"}
    SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES

    def __init__(
        self,
        vault: VaultService,
        index: IndexService,
        matters: MatterService,
        matter_paths: MatterPathPolicy,
        max_upload_mb: int = 25,
    ):
        self.vault = vault
        self.index = index
        self.matters = matters
        self.matter_paths = matter_paths
        self.max_upload_bytes = max_upload_mb * 1024 * 1024

    async def upload_to_matter(
        self, matter_id: str, upload: UploadFile, *, rebuild: bool = True
    ) -> dict[str, Any]:
        name = safe_filename(upload.filename or "uploaded-file")
        suffix = Path(name).suffix.lower()
        if suffix not in self.SUPPORTED_SUFFIXES:
            supported = ", ".join(sorted(self.SUPPORTED_SUFFIXES))
            raise ValueError(f"Unsupported file type. Use one of: {supported}.")
        data, extracted, review = await self._prepare_upload(name, upload)
        source_folder = self.matter_paths.folder(
            matter_id, "matter_files.source_documents_dir"
        )
        destination = f"{source_folder}/{name}"
        content_hash = hashlib.sha256(data).hexdigest()
        result: dict[str, Any] = {
            "path": destination,
            "name": name,
            "extracted_path": None,
            "source_id": f"SRC-{content_hash[:16].upper()}",
            "version": content_hash,
            "content_hash": content_hash,
            "source_only": suffix in self.IMAGE_SUFFIXES,
        }

        companion: str | None = None
        companion_content = ""
        companion_metadata: dict[str, Any] | None = None
        if extracted.strip() or suffix in {".pdf", ".docx"}:
            companion = f"{source_folder}/{name}.extracted.md"
            extracted_body = extracted.strip() or (
                "No extractable text was found. The source may contain only images or empty pages."
            )
            metadata = {
                "record_type": "extracted_document",
                "matter_id": matter_id,
                "source_path": destination,
                "source_filename": name,
                "source_id": result["source_id"],
                "source_version": content_hash,
                "created_at": iso_now(),
            }
            if review:
                review = dict(review)
                prefix = f"# Extracted text: {name}\n\n"
                review["segments"] = [_review_segment("equal", prefix), *review["segments"], _review_segment("equal", "\n")]
                for thread in review["comments"]:
                    thread["anchor_start"] += len(prefix)
                    thread["anchor_end"] += len(prefix)
                metadata["review"] = review
            companion_content = f"# Extracted text: {name}\n\n{extracted_body}\n"
            companion_metadata = metadata
            result["extracted_path"] = companion
        extraction_warning = getattr(extracted, "warning", "")
        if companion_metadata is not None:
            companion_metadata.update(extraction_state="partial" if extraction_warning else "available" if extracted.strip() else "unavailable",
                                      extraction_warning=extraction_warning)
        destination, companion, reused = self._persist_source(destination, data, companion, companion_content, companion_metadata)
        result.update(path=destination, extracted_path=companion, reused=reused,
                      state="saved" if extracted.strip() or suffix in self.IMAGE_SUFFIXES else "partial",
                      extraction_state="available" if extracted.strip() else "unsupported" if suffix in self.IMAGE_SUFFIXES else "unavailable",
                      support_state="supplied", uploaded_at=iso_now())
        if extraction_warning:
            result.update(state="partial", extraction_state="partial", failure_detail=extraction_warning)
        try:
            if not reused:
                self.matters.append_event(
                    matter_id,
                    "document_uploaded",
                    {"title": f"Uploaded {name}", "path": destination},
                    rebuild=False,
                )
        except Exception as exc:
            logger.error("upload event persistence failed: %s", type(exc).__name__)
            result["state"] = "partial"
            result["failure_detail"] = "Source saved, but the upload event could not be saved."
        if rebuild:
            try:
                await self.index.rebuild_async()
            except Exception:
                result.update(state="partial", failure_detail="Source saved; the file index could not refresh.")
        return result

    async def upload_many_to_matter(
        self,
        matter_id: str,
        uploads: list[UploadFile],
        *,
        intent: str = "",
    ) -> dict[str, Any]:
        if not uploads:
            raise ValueError("Select at least one file to upload.")
        attachments: list[dict[str, Any]] = []
        results: list[dict[str, Any]] = []
        for upload in uploads:
            try:
                item = await self.upload_to_matter(matter_id, upload, rebuild=False)
                attachments.append(item)
                results.append(item)
            except Exception as exc:
                results.append({"name": safe_filename(upload.filename or "uploaded-file"),
                                "state": "failed", "failure_detail": str(exc),
                                "extraction_state": "unavailable"})
        preview = self.batch_preview(attachments, intent=intent)
        preview["results"] = results
        preview["upload_state"] = "saved" if all(r["state"] == "saved" for r in results) else "partial" if attachments else "failed"
        preview["state"] = "preview"
        preview["batch_saved"] = False
        try:
            self._save_batch(matter_id, preview)
            preview["batch_saved"] = True
        except Exception:
            preview.update(upload_state="partial" if attachments else "failed", failure_detail="Source results are shown; the document-set record could not be saved.")
        if attachments:
            try:
                await self.index.rebuild_async()
            except Exception:
                preview.update(upload_state="partial", failure_detail="Sources saved; the file index could not refresh.")
        return preview

    @serialized
    def _save_batch(self, matter_id: str, preview: dict[str, Any]) -> None:
        path = self._batch_path(matter_id, preview["batch_id"])
        previous = self.vault.read_markdown(path)["metadata"] if self.vault.exists(path) else {}
        preview["state"] = previous.get("state") if previous.get("action_id") else "preview"
        self.vault.write_markdown(path, f"# Uploaded document set\n\n{preview['count']} source files.\n", {
            **previous, **preview, "matter_id": matter_id, "record_type": "document_batch",
            "created_at": previous.get("created_at") or iso_now(), "action_id": previous.get("action_id", ""),
            "state": preview["state"], "batch_saved": True,
        })

    def get_batch(self, matter_id: str, batch_id: str) -> dict[str, Any]:
        path = self._batch_path(matter_id, batch_id)
        if not self.vault.exists(path):
            raise KeyError(f"Document set not found: {batch_id}")
        metadata = self.vault.read_markdown(path)["metadata"]
        if metadata.get("matter_id") != matter_id:
            raise KeyError(f"Document set not found: {batch_id}")
        return {**metadata, "path": path}

    @serialized
    def update_batch(self, matter_id: str, batch_id: str, *, state: str, action_id: str) -> dict[str, Any]:
        path = self._batch_path(matter_id, batch_id)
        self.get_batch(matter_id, batch_id)
        self.vault.update_markdown(path, metadata_updates={"state": state, "action_id": action_id, "updated_at": iso_now()})
        return self.get_batch(matter_id, batch_id)

    def _batch_path(self, matter_id: str, batch_id: str) -> str:
        return f"{self.matters.matter_path(matter_id)}/documents/batches/{batch_id}.md"

    def batch_preview(self, attachments: list[dict[str, Any]], *, intent: str = "") -> dict[str, Any]:
        """Return read-only scan data. This method does not promote facts or change matter records."""
        count = len(attachments)
        if count == 1:
            presentation = "single"
        elif count <= 3:
            presentation = "plural"
        else:
            presentation = "batch"
        batch_hash = hashlib.sha256(
            "\n".join(str(item.get("path", "")) + ":" + str(item.get("version", "")) for item in attachments).encode("utf-8")
        ).hexdigest()
        scan = []
        for item in attachments:
            try:
                scan.append(self._quick_scan(item))
            except Exception:
                scan.append({"source_id": item["source_id"], "name": item["name"],
                             "state": "unavailable", "failure_detail": "Source saved; preview unavailable."})
        return {
            "batch_id": f"BATCH-{batch_hash[:16].upper()}",
            "presentation": presentation,
            "count": count,
            "intent": intent.strip(),
            "intent_required": not bool(intent.strip()),
            "attachments": [dict(item) for item in attachments],
            "scan": scan,
            "apply_required": True,
        }

    async def upload_many_to_workspace(self, uploads: list[UploadFile]) -> dict[str, Any]:
        if not uploads:
            raise ValueError("Select at least one file to upload.")
        attachments: list[dict[str, Any]] = []
        results: list[dict[str, Any]] = []
        for upload in uploads:
            try:
                name = safe_filename(upload.filename or "uploaded-file")
                suffix = Path(name).suffix.lower()
                if suffix not in self.SUPPORTED_SUFFIXES:
                    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}.")
                data, extracted, review = await self._prepare_upload(name, upload)
                digest = hashlib.sha256(data).hexdigest()
                path = f"04_Inbox/chat-uploads/{digest[:12]}-{name}"
                companion = f"{path}.extracted.md" if extracted.strip() else None
                companion_content = ""
                companion_metadata: dict[str, Any] | None = None
                if companion:
                    metadata = {
                        "record_type": "extracted_document", "source_path": path,
                        "source_id": f"SRC-{digest[:16].upper()}", "source_version": digest,
                        "created_at": iso_now()}
                    if review:
                        review = dict(review)
                        prefix = f"# Extracted text: {name}\n\n"
                        review["segments"] = [_review_segment("equal", prefix), *review["segments"]]
                        for thread in review["comments"]:
                            thread["anchor_start"] += len(prefix)
                            thread["anchor_end"] += len(prefix)
                        metadata["review"] = review
                    companion_content = f"# Extracted text: {name}\n\n{extracted}"
                    companion_metadata = metadata
                path, companion, reused = self._persist_source(path, data, companion, companion_content, companion_metadata)
                item = {"source_id": f"SRC-{digest[:16].upper()}", "path": path, "name": name, "version": digest,
                        "extracted_path": companion, "state": "saved" if extracted.strip() or suffix in self.IMAGE_SUFFIXES else "partial",
                        "extraction_state": "available" if extracted.strip() else "unsupported" if suffix in self.IMAGE_SUFFIXES else "unavailable",
                        "support_state": "supplied", "uploaded_at": iso_now(), "reused": reused}
                extraction_warning = getattr(extracted, "warning", "")
                if extraction_warning:
                    item.update(state="partial", extraction_state="partial", failure_detail=extraction_warning)
                attachments.append(item)
                results.append(item)
            except Exception as exc:
                results.append({"name": safe_filename(upload.filename or "uploaded-file"), "state": "failed",
                                "extraction_state": "unavailable", "failure_detail": str(exc)})
        if not attachments and len(results) == 1:
            raise ValueError(results[0]["failure_detail"])
        index_warning = None
        if attachments:
            try:
                await self.index.rebuild_async()
            except Exception:
                index_warning = "Sources saved; the file index could not refresh."
        return {"attachments": attachments, "results": results, "failure_detail": index_warning,
                "state": "saved" if not index_warning and all(r["state"] == "saved" for r in results) else "partial" if attachments else "failed"}

    async def _read_upload(self, upload: UploadFile) -> bytes:
        data = bytearray()
        while len(data) <= self.max_upload_bytes:
            remaining = self.max_upload_bytes + 1 - len(data)
            chunk = await upload.read(min(UPLOAD_READ_CHUNK_BYTES, remaining))
            if not chunk:
                return bytes(data)
            data.extend(chunk)
        raise ValueError(f"Upload exceeds {self.max_upload_bytes // (1024 * 1024)} MB limit.")

    async def _prepare_upload(
        self, name: str, upload: UploadFile
    ) -> tuple[bytes, str, dict[str, Any] | None]:
        try:
            data = await self._read_upload(upload)
            extracted, review = await self._extract_async(name, data)
            return data, extracted, review
        except ValueError:
            raise
        except Exception as exc:
            logger.error("upload validation or extraction failed: %s", type(exc).__name__)
            raise

    async def _extract_async(self, name: str, data: bytes) -> tuple[str, dict[str, Any] | None]:
        if Path(name).suffix.lower() in {".pdf", ".docx"}:
            return await asyncio.to_thread(self._extract, name, data)
        return self._extract(name, data)

    @serialized
    def _persist_source(
        self,
        destination: str,
        data: bytes,
        companion: str | None,
        companion_content: str,
        companion_metadata: dict[str, Any] | None,
    ) -> tuple[str, str | None, bool]:
        # Allocate both names under one lock. A prior extracted companion is also
        # preserved, including edits made after the original upload.
        original = destination
        digest = hashlib.sha256(data).hexdigest()[:16]
        number = 0
        while self.vault.exists(destination) or (companion and self.vault.exists(companion)):
            current = self.vault.resolve(destination)
            if current.exists() and current.read_bytes() == data:
                if not companion or self.vault.exists(companion):
                    return destination, companion, True
            number += 1
            path = Path(original)
            tag = digest if number == 1 else f"{digest}-{number}"
            destination = str(path.with_name(f"{path.stem}-{tag}{path.suffix}"))
            companion = f"{destination}.extracted.md" if companion else None
        if companion_metadata is not None:
            companion_metadata = {**companion_metadata, "source_path": destination}
        source_path = self.vault.resolve(destination)
        source_before = source_path.read_bytes() if source_path.exists() else None
        companion_path = self.vault.resolve(companion) if companion else None
        companion_before = companion_path.read_bytes() if companion_path and companion_path.exists() else None
        try:
            self.vault.write_bytes(destination, data)
            if companion:
                self.vault.write_markdown(companion, companion_content, companion_metadata)
        except Exception as exc:
            logger.error("upload persistence failed: %s", type(exc).__name__)
            if source_before is not None:
                self.vault.write_bytes(destination, source_before)
            elif source_path.exists():
                source_path.unlink()
            if companion_path:
                if companion_before is not None:
                    self.vault.write_bytes(companion, companion_before)
                elif companion_path.exists():
                    companion_path.unlink()
            raise
        return destination, companion, False

    def _quick_scan(self, attachment: dict[str, Any]) -> dict[str, Any]:
        text = ""
        extracted_path = attachment.get("extracted_path")
        if extracted_path and self.vault.exists(extracted_path):
            text = self.vault.read_markdown(extracted_path)["content"]
        elif Path(str(attachment["name"])).suffix.lower() in {".md", ".txt"}:
            text = self.vault.read_text(attachment["path"])
        words = text.split()
        return {
            "source_id": attachment["source_id"],
            "name": attachment["name"],
            "kind": Path(str(attachment["name"])).suffix.lower().lstrip(".") or "file",
            "source_only": bool(attachment.get("source_only")),
            "word_count": len(words),
            "excerpt": " ".join(words[:60]),
        }

    @staticmethod
    def _extract_text(name: str, data: bytes) -> str:
        return IngestionService._extract(name, data)[0]

    @staticmethod
    def _extract(name: str, data: bytes) -> tuple[str, dict[str, Any] | None]:
        suffix = Path(name).suffix.lower()
        if suffix in {".md", ".txt"}:
            try:
                return data.decode("utf-8"), None
            except UnicodeDecodeError:
                return ExtractedText(data.decode("utf-8", errors="replace"),
                    "Some bytes are not UTF-8 text and were replaced in the preview. Original bytes are preserved."), None
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            pages, missing = [], []
            for number, page in enumerate(reader.pages, 1):
                try:
                    text = page.extract_text() or ""
                    if not text.strip():
                        missing.append(number)
                    pages.append(text)
                except Exception:
                    missing.append(number)
                    pages.append("")
            warning = f"Text unavailable for PDF page(s): {', '.join(map(str, missing))}. Original preserved; no OCR was performed." if missing else ""
            return ExtractedText("\n\n---\n\n".join(page.strip() for page in pages), warning), None
        if suffix == ".docx":
            parts = _preflight_docx(data)
            try:
                imported = _reviewed_docx(data, parts)
                if imported is not None:
                    return imported
            except DocxSecurityError:
                raise
            except (ET.ParseError, KeyError, ValueError, zipfile.BadZipFile):
                pass
            document = DocxDocument(io.BytesIO(data))
            blocks: list[str] = []
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style = paragraph.style.name.lower() if paragraph.style else ""
                if style.startswith("heading"):
                    level_match = re.search(r"(\d+)$", style)
                    level = min(int(level_match.group(1)), 6) if level_match else 2
                    text = f"{'#' * level} {text}"
                elif "list bullet" in style:
                    text = f"- {text}"
                elif "list number" in style:
                    text = f"1. {text}"
                blocks.append(text)
            for table in document.tables:
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
                if rows:
                    blocks.append("| " + " | ".join(rows[0]) + " |")
                    blocks.append("| " + " | ".join("---" for _ in rows[0]) + " |")
                    blocks.extend("| " + " | ".join(row) + " |" for row in rows[1:])
            return "\n\n".join(blocks), None
        return "", None


def _review_segment(kind: str, text: str, change_id: str = "", author: dict[str, str] | None = None,
                    created_at: str = "") -> dict[str, str]:
    author = author or {}
    return {"kind": kind, "text": text, "change_id": change_id,
            "author_id": author.get("author_id", ""), "author_name": author.get("name", ""),
            "author_color": author.get("color", ""), "created_at": created_at}


def _author_id(name: str, used: set[str]) -> str:
    base = "author-imported-" + (re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "reviewer")
    candidate, suffix = base, 2
    while candidate in used:
        candidate, suffix = f"{base}-{suffix}", suffix + 1
    used.add(candidate)
    return candidate


def _preflight_docx(data: bytes) -> dict[str, bytes]:
    """Read a bounded DOCX package before any XML or python-docx parsing."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            members = package.infolist()
            if len(members) > DOCX_MAX_MEMBERS:
                raise DocxSecurityError("DOCX package has too many members")
            parts: dict[str, bytes] = {}
            expanded = 0
            for member in members:
                if member.flag_bits & 0x1:
                    raise DocxSecurityError("DOCX package contains encrypted content")
                size = 0
                member_name = member.filename.lower()
                is_xml_part = member_name.endswith((".xml", ".rels"))
                content = bytearray() if member_name.endswith(".xml") else None
                with package.open(member) as handle:
                    while chunk := handle.read(DOCX_READ_CHUNK_BYTES):
                        size += len(chunk)
                        expanded += len(chunk)
                        if size > DOCX_MAX_EXPANDED_BYTES or expanded > DOCX_MAX_EXPANDED_BYTES:
                            raise DocxSecurityError("DOCX package expands beyond the allowed size")
                        if is_xml_part and size > DOCX_MAX_XML_PART_BYTES:
                            raise DocxSecurityError("DOCX XML part exceeds the allowed size")
                        if content is not None:
                            content.extend(chunk)
                if content is not None:
                    parts[member.filename] = bytes(content)
    except (OSError, RuntimeError, NotImplementedError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise DocxSecurityError("DOCX package is malformed") from exc
    if "word/document.xml" not in parts:
        raise DocxSecurityError("DOCX document XML is missing")
    return parts


def _reviewed_docx(
    data: bytes, parts: dict[str, bytes] | None = None
) -> tuple[str, dict[str, Any] | None] | None:
    """Read classic Word revisions/comments. Return None when no review data exists."""
    parts = parts or _preflight_docx(data)
    try:
        root = SafeET.fromstring(parts["word/document.xml"])
    except DefusedXmlException as exc:
        raise DocxSecurityError("DOCX XML contains unsafe entities") from exc
    try:
        has_revisions = root.find(f".//{W}ins") is not None or root.find(f".//{W}del") is not None
        has_ranges = root.find(f".//{W}commentRangeStart") is not None
        if not has_revisions and not has_ranges:
            return None
        comments_root = None
        if "word/comments.xml" in parts:
            try:
                comments_root = SafeET.fromstring(parts["word/comments.xml"])
            except DefusedXmlException as exc:
                raise DocxSecurityError("DOCX XML contains unsafe entities") from exc
            except ET.ParseError:
                comments_root = None
    except DefusedXmlException as exc:
        raise DocxSecurityError("DOCX XML contains unsafe entities") from exc

    author_map: dict[str, dict[str, str]] = {}
    used_ids: set[str] = set()

    def author(name: str) -> dict[str, str] | None:
        name = name.strip()
        if not name:
            return None
        if name not in author_map:
            author_map[name] = {"author_id": _author_id(name, used_ids), "name": name,
                                "color": AUTHOR_PALETTE[len(author_map) % len(AUTHOR_PALETTE)]}
        return author_map[name]

    segments: list[dict[str, str]] = []
    visible_parts: list[str] = []
    ranges: dict[str, list[int]] = {}
    valid_revision = False

    def append(kind: str, text: str, node: ET.Element | None = None) -> None:
        nonlocal valid_revision
        if not text:
            return
        item_author = author(node.get(f"{W}author", "")) if node is not None else None
        revision_id = node.get(f"{W}id", "") if node is not None else ""
        change_id = f"CHG-WORD-{revision_id}" if revision_id else ""
        segment = _review_segment(kind, text, change_id, item_author, node.get(f"{W}date", "") if node is not None else "")
        if kind in {"insert", "delete"} and item_author is not None and revision_id:
            valid_revision = True
        if segments and all(segments[-1].get(key) == segment.get(key) for key in
                            ("kind", "change_id", "author_id", "created_at")):
            segments[-1]["text"] += text
        else:
            segments.append(segment)
        if kind != "delete":
            visible_parts.append(text)

    def walk(node: ET.Element, revision: ET.Element | None = None) -> None:
        tag = node.tag
        if tag == f"{W}commentRangeStart":
            ranges.setdefault(node.get(f"{W}id", ""), [len("".join(visible_parts)), -1])[0] = len("".join(visible_parts))
            return
        if tag == f"{W}commentRangeEnd":
            ranges.setdefault(node.get(f"{W}id", ""), [0, -1])[1] = len("".join(visible_parts))
            return
        if tag in {f"{W}ins", f"{W}del"}:
            for child in node:
                walk(child, node)
            return
        if tag in {f"{W}t", f"{W}delText", f"{W}tab", f"{W}br"}:
            text = "\t" if tag == f"{W}tab" else "\n" if tag == f"{W}br" else (node.text or "")
            append("delete" if revision is not None and revision.tag == f"{W}del" else
                   "insert" if revision is not None else "equal", text, revision)
            return
        for child in node:
            walk(child, revision)

    body = root.find(f"{W}body")
    if body is None:
        raise ValueError("DOCX document body is missing")
    blocks = [child for child in body if child.tag in {f"{W}p", f"{W}tbl"}]
    for index, block in enumerate(blocks):
        walk(block)
        if index < len(blocks) - 1:
            append("equal", "\n\n")
    visible = "".join(visible_parts)

    threads: list[dict[str, Any]] = []
    if comments_root is not None:
        for comment in comments_root.findall(f"{W}comment"):
            comment_id = comment.get(f"{W}id", "")
            bounds = ranges.get(comment_id)
            if not bounds or bounds[1] < bounds[0]:
                continue
            name = comment.get(f"{W}author", "").strip()
            item_author = author(name)
            body_text = "".join(node.text or "" for node in comment.findall(f".//{W}t")).strip()
            if not item_author or not body_text:
                continue
            start, end = bounds
            threads.append({"thread_id": f"COM-WORD-{comment_id}", "quote": visible[start:end],
                            "anchor_start": start, "anchor_end": end, "resolved": False,
                            "resolved_at": "", "resolved_by": "", "entries": [{
                                "comment_id": f"MSG-WORD-{comment_id}", "author_id": item_author["author_id"],
                                "author_name": item_author["name"], "body": body_text,
                                "created_at": comment.get(f"{W}date", "")}]})
    if not valid_revision and not threads:
        return visible, None
    _merge_word_replacements(segments)
    review = {"version": 2, "tracking": True, "authors": list(author_map.values()),
              "segments": segments, "comments": threads, "comment_events": []}
    return visible, review


def _merge_word_replacements(segments: list[dict[str, str]]) -> None:
    """Give adjacent same-reviewer delete/insert pairs one replacement ID."""
    for before, after in zip(segments, segments[1:]):
        if before.get("kind") != "delete" or after.get("kind") != "insert":
            continue
        if not before.get("author_id") or before.get("author_id") != after.get("author_id"):
            continue
        before_date, after_date = before.get("created_at", ""), after.get("created_at", "")
        if before_date and after_date and before_date != after_date:
            continue
        change_id = before.get("change_id") or after.get("change_id")
        if change_id:
            before["change_id"] = after["change_id"] = change_id
