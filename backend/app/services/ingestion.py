from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document as DocxDocument
from fastapi import UploadFile
from pypdf import PdfReader

from app.services.index import IndexService
from app.services.matters import MatterService
from app.services.matter_paths import MatterPathPolicy
from app.services.vault import VaultService
from app.utils.paths import safe_filename
from app.utils.time import iso_now


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{WORD_NS}}}"
AUTHOR_PALETTE = ("#2F5597", "#7030A0", "#008272", "#A64B00", "#C0006F", "#5B6573", "#7A3E00", "#006B8F")


class IngestionService:
    TEXT_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}
    IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic"}
    SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES

    def __init__(
        self,
        vault: VaultService,
        index: IndexService,
        matters: MatterService,
        matter_paths: MatterPathPolicy,
        max_upload_mb: int = 25,
    ):
        self.vault = vault
        self.index = index
        self.matters = matters
        self.matter_paths = matter_paths
        self.max_upload_bytes = max_upload_mb * 1024 * 1024

    async def upload_to_matter(self, matter_id: str, upload: UploadFile) -> dict[str, Any]:
        name = safe_filename(upload.filename or "uploaded-file")
        suffix = Path(name).suffix.lower()
        if suffix not in self.SUPPORTED_SUFFIXES:
            supported = ", ".join(sorted(self.SUPPORTED_SUFFIXES))
            raise ValueError(f"Unsupported file type. Use one of: {supported}.")
        data = await upload.read()
        if len(data) > self.max_upload_bytes:
            raise ValueError(f"Upload exceeds {self.max_upload_bytes // (1024 * 1024)} MB limit.")
        source_folder = self.matter_paths.folder(
            matter_id, "matter_files.source_documents_dir"
        )
        destination = f"{source_folder}/{name}"
        self.vault.write_bytes(destination, data)
        content_hash = hashlib.sha256(data).hexdigest()
        result: dict[str, Any] = {
            "path": destination,
            "name": name,
            "extracted_path": None,
            "source_id": f"SRC-{content_hash[:16].upper()}",
            "version": content_hash,
            "content_hash": content_hash,
            "source_only": suffix in self.IMAGE_SUFFIXES,
        }

        extracted, review = self._extract(name, data)
        if extracted.strip() or suffix in {".pdf", ".docx"}:
            companion = f"{source_folder}/{name}.extracted.md"
            extracted_body = extracted.strip() or (
                "No extractable text was found. The source may contain only images or empty pages."
            )
            metadata = {
                "record_type": "extracted_document",
                "matter_id": matter_id,
                "source_path": destination,
                "source_filename": name,
                "source_id": result["source_id"],
                "source_version": content_hash,
                "created_at": iso_now(),
            }
            if review:
                review = dict(review)
                prefix = f"# Extracted text: {name}\n\n"
                review["segments"] = [_review_segment("equal", prefix), *review["segments"], _review_segment("equal", "\n")]
                for thread in review["comments"]:
                    thread["anchor_start"] += len(prefix)
                    thread["anchor_end"] += len(prefix)
                metadata["review"] = review
            self.vault.write_markdown(
                companion,
                f"# Extracted text: {name}\n\n{extracted_body}\n",
                metadata,
            )
            result["extracted_path"] = companion
        self.matters.append_event(
            matter_id,
            "document_uploaded",
            {"title": f"Uploaded {name}", "path": destination},
            rebuild=False,
        )
        self.index.rebuild()
        return result

    async def upload_many_to_matter(
        self,
        matter_id: str,
        uploads: list[UploadFile],
        *,
        intent: str = "",
    ) -> dict[str, Any]:
        if not uploads:
            raise ValueError("Select at least one file to upload.")
        attachments = [await self.upload_to_matter(matter_id, upload) for upload in uploads]
        preview = self.batch_preview(attachments, intent=intent)
        path = self._batch_path(matter_id, preview["batch_id"])
        self.vault.write_markdown(path, f"# Uploaded document set\n\n{preview['count']} source files.\n", {
            **preview, "matter_id": matter_id, "record_type": "document_batch", "state": "preview",
            "created_at": iso_now(), "action_id": "",
        })
        return preview

    def get_batch(self, matter_id: str, batch_id: str) -> dict[str, Any]:
        path = self._batch_path(matter_id, batch_id)
        if not self.vault.exists(path):
            raise KeyError(f"Document set not found: {batch_id}")
        metadata = self.vault.read_markdown(path)["metadata"]
        if metadata.get("matter_id") != matter_id:
            raise KeyError(f"Document set not found: {batch_id}")
        return {**metadata, "path": path}

    def update_batch(self, matter_id: str, batch_id: str, *, state: str, action_id: str) -> dict[str, Any]:
        path = self._batch_path(matter_id, batch_id)
        self.get_batch(matter_id, batch_id)
        self.vault.update_markdown(path, metadata_updates={"state": state, "action_id": action_id, "updated_at": iso_now()})
        return self.get_batch(matter_id, batch_id)

    def _batch_path(self, matter_id: str, batch_id: str) -> str:
        return f"{self.matters.matter_path(matter_id)}/documents/batches/{batch_id}.md"

    def batch_preview(self, attachments: list[dict[str, Any]], *, intent: str = "") -> dict[str, Any]:
        """Return read-only scan data. This method does not promote facts or change matter records."""
        count = len(attachments)
        if count == 1:
            presentation = "single"
        elif count <= 3:
            presentation = "plural"
        else:
            presentation = "batch"
        batch_hash = hashlib.sha256(
            "\n".join(str(item.get("version", "")) for item in attachments).encode("utf-8")
        ).hexdigest()
        scan = [self._quick_scan(item) for item in attachments]
        return {
            "batch_id": f"BATCH-{batch_hash[:16].upper()}",
            "presentation": presentation,
            "count": count,
            "intent": intent.strip(),
            "intent_required": not bool(intent.strip()),
            "attachments": [
                {
                    "source_id": item["source_id"],
                    "path": item["path"],
                    "name": item["name"],
                    "version": item["version"],
                }
                for item in attachments
            ],
            "scan": scan,
            "apply_required": True,
        }

    async def upload_many_to_workspace(self, uploads: list[UploadFile]) -> dict[str, Any]:
        if not uploads:
            raise ValueError("Select at least one file to upload.")
        attachments: list[dict[str, str]] = []
        for upload in uploads:
            name = safe_filename(upload.filename or "uploaded-file")
            suffix = Path(name).suffix.lower()
            if suffix not in self.SUPPORTED_SUFFIXES:
                raise ValueError(f"Unsupported file type: {suffix or 'unknown'}.")
            data = await upload.read()
            if len(data) > self.max_upload_bytes:
                raise ValueError(f"Upload exceeds {self.max_upload_bytes // (1024 * 1024)} MB limit.")
            digest = hashlib.sha256(data).hexdigest()
            path = f"04_Inbox/chat-uploads/{digest[:12]}-{name}"
            self.vault.write_bytes(path, data)
            extracted, review = self._extract(name, data)
            if extracted.strip():
                metadata = {
                    "record_type": "extracted_document", "source_path": path,
                    "source_id": f"SRC-{digest[:16].upper()}", "source_version": digest,
                    "created_at": iso_now()}
                if review:
                    review = dict(review)
                    prefix = f"# Extracted text: {name}\n\n"
                    review["segments"] = [_review_segment("equal", prefix), *review["segments"]]
                    for thread in review["comments"]:
                        thread["anchor_start"] += len(prefix)
                        thread["anchor_end"] += len(prefix)
                    metadata["review"] = review
                self.vault.write_markdown(f"{path}.extracted.md", f"# Extracted text: {name}\n\n{extracted}", metadata)
            attachments.append({"source_id": f"SRC-{digest[:16].upper()}", "path": path, "name": name, "version": digest})
        self.index.rebuild()
        return {"attachments": attachments}

    def _quick_scan(self, attachment: dict[str, Any]) -> dict[str, Any]:
        text = ""
        extracted_path = attachment.get("extracted_path")
        if extracted_path and self.vault.exists(extracted_path):
            text = self.vault.read_markdown(extracted_path)["content"]
        elif Path(str(attachment["name"])).suffix.lower() in {".md", ".txt"}:
            text = self.vault.read_text(attachment["path"])
        words = text.split()
        return {
            "source_id": attachment["source_id"],
            "name": attachment["name"],
            "kind": Path(str(attachment["name"])).suffix.lower().lstrip(".") or "file",
            "source_only": bool(attachment.get("source_only")),
            "word_count": len(words),
            "excerpt": " ".join(words[:60]),
        }

    @staticmethod
    def _extract_text(name: str, data: bytes) -> str:
        return IngestionService._extract(name, data)[0]

    @staticmethod
    def _extract(name: str, data: bytes) -> tuple[str, dict[str, Any] | None]:
        suffix = Path(name).suffix.lower()
        if suffix in {".md", ".txt"}:
            return data.decode("utf-8", errors="replace"), None
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n---\n\n".join(page.strip() for page in pages), None
        if suffix == ".docx":
            try:
                imported = _reviewed_docx(data)
                if imported is not None:
                    return imported
            except (ET.ParseError, KeyError, ValueError, zipfile.BadZipFile):
                pass
            document = DocxDocument(io.BytesIO(data))
            blocks: list[str] = []
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style = paragraph.style.name.lower() if paragraph.style else ""
                if style.startswith("heading"):
                    level_match = re.search(r"(\d+)$", style)
                    level = min(int(level_match.group(1)), 6) if level_match else 2
                    text = f"{'#' * level} {text}"
                elif "list bullet" in style:
                    text = f"- {text}"
                elif "list number" in style:
                    text = f"1. {text}"
                blocks.append(text)
            for table in document.tables:
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
                if rows:
                    blocks.append("| " + " | ".join(rows[0]) + " |")
                    blocks.append("| " + " | ".join("---" for _ in rows[0]) + " |")
                    blocks.extend("| " + " | ".join(row) + " |" for row in rows[1:])
            return "\n\n".join(blocks), None
        return "", None


def _review_segment(kind: str, text: str, change_id: str = "", author: dict[str, str] | None = None,
                    created_at: str = "") -> dict[str, str]:
    author = author or {}
    return {"kind": kind, "text": text, "change_id": change_id,
            "author_id": author.get("author_id", ""), "author_name": author.get("name", ""),
            "author_color": author.get("color", ""), "created_at": created_at}


def _author_id(name: str, used: set[str]) -> str:
    base = "author-imported-" + (re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "reviewer")
    candidate, suffix = base, 2
    while candidate in used:
        candidate, suffix = f"{base}-{suffix}", suffix + 1
    used.add(candidate)
    return candidate


def _reviewed_docx(data: bytes) -> tuple[str, dict[str, Any] | None] | None:
    """Read classic Word revisions/comments. Return None when no review data exists."""
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        root = ET.fromstring(package.read("word/document.xml"))
        has_revisions = root.find(f".//{W}ins") is not None or root.find(f".//{W}del") is not None
        has_ranges = root.find(f".//{W}commentRangeStart") is not None
        if not has_revisions and not has_ranges:
            return None
        comments_root = None
        if "word/comments.xml" in package.namelist():
            try:
                comments_root = ET.fromstring(package.read("word/comments.xml"))
            except ET.ParseError:
                comments_root = None

    author_map: dict[str, dict[str, str]] = {}
    used_ids: set[str] = set()

    def author(name: str) -> dict[str, str] | None:
        name = name.strip()
        if not name:
            return None
        if name not in author_map:
            author_map[name] = {"author_id": _author_id(name, used_ids), "name": name,
                                "color": AUTHOR_PALETTE[len(author_map) % len(AUTHOR_PALETTE)]}
        return author_map[name]

    segments: list[dict[str, str]] = []
    visible_parts: list[str] = []
    ranges: dict[str, list[int]] = {}
    valid_revision = False

    def append(kind: str, text: str, node: ET.Element | None = None) -> None:
        nonlocal valid_revision
        if not text:
            return
        item_author = author(node.get(f"{W}author", "")) if node is not None else None
        revision_id = node.get(f"{W}id", "") if node is not None else ""
        change_id = f"CHG-WORD-{revision_id}" if revision_id else ""
        segment = _review_segment(kind, text, change_id, item_author, node.get(f"{W}date", "") if node is not None else "")
        if kind in {"insert", "delete"} and item_author is not None and revision_id:
            valid_revision = True
        if segments and all(segments[-1].get(key) == segment.get(key) for key in
                            ("kind", "change_id", "author_id", "created_at")):
            segments[-1]["text"] += text
        else:
            segments.append(segment)
        if kind != "delete":
            visible_parts.append(text)

    def walk(node: ET.Element, revision: ET.Element | None = None) -> None:
        tag = node.tag
        if tag == f"{W}commentRangeStart":
            ranges.setdefault(node.get(f"{W}id", ""), [len("".join(visible_parts)), -1])[0] = len("".join(visible_parts))
            return
        if tag == f"{W}commentRangeEnd":
            ranges.setdefault(node.get(f"{W}id", ""), [0, -1])[1] = len("".join(visible_parts))
            return
        if tag in {f"{W}ins", f"{W}del"}:
            for child in node:
                walk(child, node)
            return
        if tag in {f"{W}t", f"{W}delText", f"{W}tab", f"{W}br"}:
            text = "\t" if tag == f"{W}tab" else "\n" if tag == f"{W}br" else (node.text or "")
            append("delete" if revision is not None and revision.tag == f"{W}del" else
                   "insert" if revision is not None else "equal", text, revision)
            return
        for child in node:
            walk(child, revision)

    body = root.find(f"{W}body")
    if body is None:
        raise ValueError("DOCX document body is missing")
    blocks = [child for child in body if child.tag in {f"{W}p", f"{W}tbl"}]
    for index, block in enumerate(blocks):
        walk(block)
        if index < len(blocks) - 1:
            append("equal", "\n\n")
    visible = "".join(visible_parts)

    threads: list[dict[str, Any]] = []
    if comments_root is not None:
        for comment in comments_root.findall(f"{W}comment"):
            comment_id = comment.get(f"{W}id", "")
            bounds = ranges.get(comment_id)
            if not bounds or bounds[1] < bounds[0]:
                continue
            name = comment.get(f"{W}author", "").strip()
            item_author = author(name)
            body_text = "".join(node.text or "" for node in comment.findall(f".//{W}t")).strip()
            if not item_author or not body_text:
                continue
            start, end = bounds
            threads.append({"thread_id": f"COM-WORD-{comment_id}", "quote": visible[start:end],
                            "anchor_start": start, "anchor_end": end, "resolved": False,
                            "resolved_at": "", "resolved_by": "", "entries": [{
                                "comment_id": f"MSG-WORD-{comment_id}", "author_id": item_author["author_id"],
                                "author_name": item_author["name"], "body": body_text,
                                "created_at": comment.get(f"{W}date", "")}]})
    if not valid_revision and not threads:
        return visible, None
    _merge_word_replacements(segments)
    review = {"version": 2, "tracking": True, "authors": list(author_map.values()),
              "segments": segments, "comments": threads, "comment_events": []}
    return visible, review


def _merge_word_replacements(segments: list[dict[str, str]]) -> None:
    """Give adjacent same-reviewer delete/insert pairs one replacement ID."""
    for before, after in zip(segments, segments[1:]):
        if before.get("kind") != "delete" or after.get("kind") != "insert":
            continue
        if not before.get("author_id") or before.get("author_id") != after.get("author_id"):
            continue
        before_date, after_date = before.get("created_at", ""), after.get("created_at", "")
        if before_date and after_date and before_date != after_date:
            continue
        change_id = before.get("change_id") or after.get("change_id")
        if change_id:
            before["change_id"] = after["change_id"] = change_id
