from __future__ import annotations

import io
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from app.services.document_review import review_segments
from app.services.vault import VaultService


INLINE_MARKUP = re.compile(r"(!?\[)([^\]]+)(\]\([^)]*\))|(`+)(.*?)\4|(\*\*|__)(.*?)\6|(?<!\*)\*([^*]+)\*|(?<!_)_([^_]+)_")


def markdown_to_plain(markdown: str) -> str:
    lines: list[str] = []
    for line in markdown.splitlines():
        line = re.sub(r"^\s{0,3}#{1,6}\s+", "", line)
        line = re.sub(r"^\s*>\s?", "", line)
        line = re.sub(r"^\s*[-+*]\s+", "- ", line)
        line = re.sub(r"^\s*\d+[.)]\s+", "", line)
        previous = None
        while previous != line:
            previous = line
            line = INLINE_MARKUP.sub(lambda match: next((item for item in (match.group(2), match.group(5), match.group(7), match.group(8), match.group(9)) if item is not None), ""), line)
        lines.append(line)
    return "\n".join(lines).strip() + "\n"


class DocumentExportService:
    def __init__(self, vault: VaultService):
        self.vault = vault

    def export(self, path: str, output_format: str) -> tuple[str, str]:
        document = self.vault.read_document(path)
        if document.get("kind") != "markdown":
            raise ValueError("Only Markdown documents can be exported.")
        review = document.get("metadata", {}).get("review", {})
        review = review if isinstance(review, dict) else {}
        current = markdown_to_plain(document["content"])
        stored_segments = review.get("segments")
        if review.get("version") == 2 and isinstance(stored_segments, list):
            segments = [dict(item) for item in stored_segments if isinstance(item, dict)]
        else:
            baseline = markdown_to_plain(str(review.get("baseline") or document["content"]))
            segments = review_segments(baseline, current)
        comments = [dict(item) for item in review.get("comments", []) if isinstance(item, dict)]
        stem = Path(path).name.removesuffix(".md").removesuffix(".extracted")
        if output_format == "docx":
            data = self._docx(segments, comments, document["name"], document["content"])
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif output_format == "pdf":
            data = self._pdf(segments, comments, current, document["content"])
            media_type = "application/pdf"
        else:
            raise ValueError("Export format must be docx or pdf.")
        output_path = f".exports/{stem}.{output_format}"
        self.vault.write_bytes(output_path, data)
        return output_path, media_type

    @staticmethod
    def _docx(
        segments: list[dict[str, str]],
        comments: list[dict[str, Any]],
        title: str,
        markdown: str,
    ) -> bytes:
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.85)
        section.left_margin = section.right_margin = Inches(0.9)
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(11)
        document.core_properties.title = title
        document.core_properties.author = "Themis.ai"

        paragraph = document.add_paragraph()
        tracked_runs: list[tuple[Any, dict[str, Any]]] = []
        visible_runs: list[tuple[Any, int, int]] = []
        visible_offset = 0
        comment_boundaries = {
            boundary
            for comment in comments
            for boundary in (_integer(comment.get("anchor_start")), _integer(comment.get("anchor_end")))
            if boundary is not None and boundary >= 0
        }
        for segment in segments:
            pieces = segment["text"].split("\n")
            for index, piece in enumerate(pieces):
                if piece:
                    cuts = [0, len(piece)]
                    if segment["kind"] != "delete":
                        cuts.extend(boundary - visible_offset for boundary in comment_boundaries
                                    if visible_offset < boundary < visible_offset + len(piece))
                    cuts = sorted(set(cuts))
                    for left, right in zip(cuts, cuts[1:]):
                        run = paragraph.add_run(piece[left:right])
                        tracked_runs.append((run, segment))
                        if segment["kind"] != "delete":
                            visible_runs.append((run, visible_offset + left, visible_offset + right))
                    if segment["kind"] != "delete":
                        visible_offset += len(piece)
                if index < len(pieces) - 1:
                    paragraph = document.add_paragraph()
                    if segment["kind"] != "delete":
                        visible_offset += 1

        for paragraph, style in zip(document.paragraphs, _block_styles(markdown), strict=False):
            if style:
                paragraph.style = style

        for comment in comments:
            quote = str(comment.get("quote") or "")
            anchor_start = _integer(comment.get("anchor_start"), -1)
            anchor_end = _integer(comment.get("anchor_end"), -1)
            target_runs = [run for run, start, end in visible_runs
                           if start < anchor_end and end > anchor_start]
            if not target_runs and quote:
                target = next((item for item in document.paragraphs if quote in item.text and item.runs), None)
                target_runs = list(target.runs) if target is not None else []
            if not target_runs:
                target = next((item for item in document.paragraphs if item.runs and item.text.strip()), None)
                target_runs = list(target.runs) if target is not None else []
            if target_runs:
                entries = _comment_entries(comment)
                lines = []
                if comment.get("resolved"):
                    lines.append("Resolved")
                for entry in entries:
                    label = str(entry.get("author_name") or "Imported reviewer")
                    created_at = str(entry.get("created_at") or "")
                    lines.append(f"{label}{f' — {created_at}' if created_at else ''}: {entry.get('body', '')}")
                if not lines:
                    continue
                document.add_comment(
                    target_runs,
                    text="\n".join(lines),
                    author=str(entries[0].get("author_name") or "Themis.ai User") if entries else "Themis.ai User",
                )

        next_id = 1
        timestamp = datetime.now(UTC).isoformat().replace("+00:00", "Z")
        for run, segment in tracked_runs:
            kind = str(segment.get("kind") or "")
            if kind not in {"insert", "delete"}:
                continue
            element = run._r
            parent = element.getparent()
            if parent is None:
                continue
            index = parent.index(element)
            parent.remove(element)
            wrapper = OxmlElement("w:ins" if kind == "insert" else "w:del")
            wrapper.set(qn("w:id"), str(next_id))
            wrapper.set(qn("w:author"), str(segment.get("author_name") or segment.get("author_id") or "Themis.ai"))
            wrapper.set(qn("w:date"), str(segment.get("created_at") or timestamp))
            next_id += 1
            if kind == "delete":
                for text_node in element.findall(".//" + qn("w:t")):
                    text_node.tag = qn("w:delText")
            wrapper.append(element)
            parent.insert(index, wrapper)

        settings = document.settings._element
        if settings.find(qn("w:trackRevisions")) is None:
            settings.insert(0, OxmlElement("w:trackRevisions"))
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @classmethod
    def _pdf(
        cls,
        segments: list[dict[str, str]],
        comments: list[dict[str, Any]],
        current: str,
        markdown: str,
    ) -> bytes:
        source = io.BytesIO()
        pdf = canvas.Canvas(source, pagesize=LETTER)
        width, height = LETTER
        margin = 54.0
        x = margin
        y = height - margin
        font = "Helvetica"
        font_size = 10.5
        leading = 15.0
        block_styles = _block_styles(markdown)
        current_line = 0
        pdf.setFont(font, font_size)
        marks: list[dict[str, Any]] = []
        visible: list[dict[str, Any]] = []
        current_offset = 0

        def new_line(visible_line: bool) -> None:
            nonlocal x, y, page_number, current_line
            x = margin
            y -= leading
            if visible_line:
                current_line += 1
            if y < margin:
                pdf.showPage()
                pdf.setFont(font, font_size)
                y = height - margin
                page_number += 1

        page_number = 0
        for segment in segments:
            kind = segment["kind"]
            for token in re.findall(r"\n|[^\S\n]+|[^\s]+", segment["text"]):
                if token == "\n":
                    if kind != "delete":
                        current_offset += 1
                    new_line(kind != "delete")
                    continue
                style = block_styles[current_line] if current_line < len(block_styles) else ""
                active_font = "Helvetica-Bold" if style.startswith("Heading") else font
                active_size = 15.0 if style == "Heading 1" else 12.0 if style.startswith("Heading") else font_size
                pdf.setFont(active_font, active_size)
                token_width = stringWidth(token, active_font, active_size)
                if token.strip() and x + token_width > width - margin:
                    new_line(False)
                if kind == "delete":
                    pdf.setFillColorRGB(0.72, 0.18, 0.14)
                elif kind == "insert":
                    pdf.setFillColorRGB(0.12, 0.43, 0.26)
                else:
                    pdf.setFillColorRGB(0.12, 0.12, 0.12)
                pdf.drawString(x, y, token)
                rect = (x, y - 2, x + token_width, y + active_size)
                if kind in {"insert", "delete"} and token.strip():
                    line_y = y - 1 if kind == "insert" else y + active_size * 0.35
                    pdf.setStrokeColorRGB(*( (0.12, 0.43, 0.26) if kind == "insert" else (0.72, 0.18, 0.14) ))
                    pdf.line(x, line_y, x + token_width, line_y)
                    marks.append({"page": page_number, "kind": kind, "rect": rect, "text": token})
                if kind != "delete":
                    visible.append({"page": page_number, "start": current_offset, "end": current_offset + len(token), "rect": rect})
                    current_offset += len(token)
                x += token_width
        pdf.save()

        reader = PdfReader(io.BytesIO(source.getvalue()))
        writer = PdfWriter()
        writer.clone_document_from_reader(reader)
        for mark in marks:
            subtype = "/Underline" if mark["kind"] == "insert" else "/StrikeOut"
            color = (0.12, 0.43, 0.26) if mark["kind"] == "insert" else (0.72, 0.18, 0.14)
            writer.add_annotation(mark["page"], cls._markup_annotation(
                subtype, mark["rect"], color,
                f"{'Inserted' if mark['kind'] == 'insert' else 'Deleted'}: {mark['text']}",
                "Themis.ai",
            ))
        for comment in comments:
            quote = str(comment.get("quote") or "")
            start = _integer(comment.get("anchor_start"))
            end = _integer(comment.get("anchor_end"))
            if start is None or end is None or start < 0 or end <= start or end > len(current):
                start = current.find(quote)
                end = start + len(quote)
            if start < 0 or end <= start:
                continue
            hits = [item for item in visible if item["start"] < end and item["end"] > start]
            entries = _comment_entries(comment)
            contents = "\n".join(
                f"{entry.get('author_name') or 'Themis.ai User'}: {entry.get('body') or ''}" for entry in entries
            )
            for hit in hits:
                writer.add_annotation(hit["page"], cls._markup_annotation(
                    "/Highlight", hit["rect"], (0.96, 0.75, 0.2),
                    contents,
                    str(entries[0].get("author_name") or "Themis.ai User") if entries else "Themis.ai User",
                ))
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()

    @staticmethod
    def _markup_annotation(
        subtype: str,
        rect: tuple[float, float, float, float],
        color: tuple[float, float, float],
        contents: str,
        author: str,
    ) -> DictionaryObject:
        x1, y1, x2, y2 = rect
        return DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject(subtype),
            NameObject("/Rect"): ArrayObject([FloatObject(v) for v in rect]),
            NameObject("/QuadPoints"): ArrayObject([FloatObject(v) for v in (x1, y2, x2, y2, x1, y1, x2, y1)]),
            NameObject("/C"): ArrayObject([FloatObject(v) for v in color]),
            NameObject("/Contents"): TextStringObject(contents),
            NameObject("/T"): TextStringObject(author),
            NameObject("/F"): FloatObject(4),
        })


def _block_styles(markdown: str) -> list[str]:
    styles: list[str] = []
    for line in markdown.splitlines():
        heading = re.match(r"^\s{0,3}(#{1,6})\s+", line)
        if heading:
            styles.append(f"Heading {min(len(heading.group(1)), 3)}")
        elif re.match(r"^\s*[-+*]\s+", line):
            styles.append("List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            styles.append("List Number")
        elif re.match(r"^\s*>\s?", line):
            styles.append("Quote")
        else:
            styles.append("")
    return styles


def _comment_entries(comment: dict[str, Any]) -> list[dict[str, Any]]:
    entries = [dict(entry) for entry in comment.get("entries", []) if isinstance(entry, dict)]
    if entries:
        return entries
    body = str(comment.get("comment") or comment.get("body") or "")
    if not body:
        return []
    return [{"author_name": str(comment.get("author") or comment.get("author_name") or "Themis.ai User"),
             "body": body, "created_at": str(comment.get("created_at") or "")}]


def _integer(value: Any, default: int | None = None) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default
