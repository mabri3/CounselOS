import pytest
from docx import Document
from pypdf import PdfReader

from app.services.document_export import DocumentExportService
from app.services.document_review import DocumentReviewService
from app.services.vault import VaultService


def _export_text(vault, path, output_format):
    exported, _ = DocumentExportService(vault).export(path, output_format)
    if output_format == "docx":
        return "\n".join(paragraph.text for paragraph in Document(vault.resolve(exported)).paragraphs)
    return "\n".join(page.extract_text() or "" for page in PdfReader(vault.resolve(exported)).pages)


def test_claim_passages_and_revisions_survive_existing_exports(tmp_path):
    vault = VaultService(tmp_path)
    # Synthetic wording tests transport only. The live demo must retrieve the
    # actual 16 CFR 312.2 passages before calling them retrieved authority.
    text = (
        "The child definition covers an individual under 13.\n\n"
        "A persistent identifier is listed as personal information."
    )
    claims = [
        {
            "claim_id": "CLM-CHILD", "text": "The child definition covers an individual under 13.",
            "claim_revision": "claim-child-r1", "output_revision": "output-r7",
            "applicability": {"regulated_actor": "unknown operator", "jurisdiction": "United States",
                              "explanation": "The definition alone does not establish operator status."},
            "evidence": [{"claim_id": "CLM-CHILD", "source_id": "SRC-COPPA", "locator": "child",
                          "available_excerpt": "Child means an individual under the age of 13.",
                          "support_state": "retrieved", "source_label": "16 CFR 312.2",
                          "claim_revision": "claim-child-r1", "output_revision": "output-r7"}],
            "support_gap": "Whether the learning app is a covered operator remains unresolved.",
        },
        {
            "claim_id": "CLM-ID", "text": "A persistent identifier is listed as personal information.",
            "claim_revision": "claim-id-r1", "output_revision": "output-r7",
            "applicability": {"regulated_actor": "unknown operator", "jurisdiction": "United States"},
            "evidence": [{"claim_id": "CLM-ID", "source_id": "SRC-COPPA",
                          "locator": "personal information — persistent identifier",
                          "available_excerpt": "Personal information includes a persistent identifier.",
                          "support_state": "retrieved", "source_label": "16 CFR 312.2",
                          "claim_revision": "claim-id-r1", "output_revision": "output-r7"}],
        },
    ]
    vault.write_markdown("matter/inquiry.md", text, {
        "record_type": "workspace_inquiry", "output_revision": "output-r7", "claims": claims,
    })

    for output_format in ("docx", "pdf"):
        exported = _export_text(vault, "matter/inquiry.md", output_format)
        flat = " ".join(exported.split())
        assert text.split("\n\n")[0] in flat
        assert "claim revision claim-child-r1" in flat
        assert "output revision output-r7" in flat
        assert "Child means an individual under the age of 13." in flat
        assert "Personal information includes a persistent identifier." in flat
        assert "locator: child" in flat
        assert "locator: personal information — persistent identifier" in flat
        assert "status: Retrieved" in flat
        assert "operator status" in flat


def test_export_keeps_prose_and_omits_malformed_or_unsafe_claim_links(tmp_path):
    vault = VaultService(tmp_path)
    text = "Useful analysis remains available."
    vault.write_markdown("matter/inquiry.md", text, {
        "claims": [
            "malformed",
            {"claim_id": "CLM-SAFE", "text": text, "claim_revision": "r1", "output_revision": "out1",
             "evidence": [{"source_id": "SRC-UNSAFE", "support_state": "unknown",
                           "url": "javascript:alert(1)", "available_excerpt": None}],
             "support_gap": "The exact passage is unavailable."},
        ],
    })

    exported = _export_text(vault, "matter/inquiry.md", "docx")

    assert text in exported
    assert "The exact passage is unavailable." in exported
    assert "javascript:" not in exported


@pytest.mark.parametrize("output_format", ["docx", "pdf"])
def test_export_lists_safe_saved_document_references(tmp_path, output_format):
    vault = VaultService(tmp_path)
    matter = "03_Matters/example-matter"
    product_spec = f"{matter}/documents/product-spec.md"
    definitions = f"{matter}/documents/coppa-definitions.md"
    code_sample = f"{matter}/documents/code-sample.md"
    vault.write_markdown(product_spec, "Product specification")
    vault.write_markdown(definitions, "Definitions")
    vault.write_markdown(code_sample, "Code sample")
    vault.write_markdown(
        f"{matter}/work-product/draft/launch-advice.md",
        "\n".join([
            f"Read the [product specification]({product_spec}) and [definitions]({definitions}#personal-information).",
            "Ignore [unsafe](javascript:alert(1)) and [missing](03_Matters/example-matter/documents/missing.md).",
            "Ignore [outside](../../outside.md).",
            f"`[code reference]({code_sample})`",
        ]),
    )

    exported = _export_text(vault, f"{matter}/work-product/draft/launch-advice.md", output_format)
    flat = " ".join(exported.split())

    assert "Document references" in flat
    assert f"product specification — {product_spec}" in flat
    assert f"definitions — {definitions}#personal-information" in flat
    assert "javascript:" not in flat
    assert "documents/missing.md" not in flat
    assert "outside —" not in flat
    assert f"code reference — {code_sample}" not in flat


@pytest.mark.parametrize("output_format", ["docx", "pdf"])
def test_accepted_export_omits_pending_document_reference(tmp_path, output_format):
    vault = VaultService(tmp_path)
    matter = "03_Matters/example-matter"
    pending_source = f"{matter}/documents/pending-source.md"
    draft = f"{matter}/work-product/draft/advice.md"
    pending_link = f"Read [pending source]({pending_source})."
    vault.write_markdown(pending_source, "Pending source")
    vault.write_markdown(draft, "Current advice.")
    DocumentReviewService(vault).propose_agent_revision(draft, f"Current advice.\n\n{pending_link}")

    markup = _export_text(vault, draft, output_format)
    accepted, _ = DocumentExportService(vault).export(draft, output_format, mode="accepted_text")
    if output_format == "docx":
        accepted_text = "\n".join(
            paragraph.text for paragraph in Document(vault.resolve(accepted)).paragraphs
        )
    else:
        accepted_text = "\n".join(
            page.extract_text() or "" for page in PdfReader(vault.resolve(accepted)).pages
        )
    markup = " ".join(markup.split())
    accepted_text = " ".join(accepted_text.split())

    assert f"pending source — {pending_source}" in markup
    assert "Current advice." in accepted_text
    assert f"pending source — {pending_source}" not in accepted_text
