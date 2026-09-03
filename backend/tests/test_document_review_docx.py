from __future__ import annotations

import io
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi import UploadFile
from pypdf import PdfReader

from app.services import ingestion


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
PATH = "03_Matters/beacon-instant-onboarding/work-product/draft/docx-review.md"


def _reviewed_docx(*, malformed_comments: bool = False) -> bytes:
    source = io.BytesIO()
    document = Document()
    document.add_paragraph("Payment is 10 days.")
    document.save(source)
    source.seek(0)

    with zipfile.ZipFile(source) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    root = ET.fromstring(parts["word/document.xml"])
    paragraph = root.find(".//w:body/w:p", NS)
    assert paragraph is not None
    for child in list(paragraph):
        paragraph.remove(child)

    def run(tag: str, text: str) -> ET.Element:
        wrapper = ET.Element(qn(tag))
        r = ET.SubElement(wrapper, qn("w:r")) if tag in {"w:ins", "w:del"} else wrapper
        node = ET.SubElement(r, qn("w:delText" if tag == "w:del" else "w:t"))
        node.text = text
        return wrapper

    paragraph.append(run("w:r", "Payment is "))
    deletion = run("w:del", "10")
    deletion.set(qn("w:id"), "7")
    deletion.set(qn("w:author"), "Alice Reviewer")
    deletion.set(qn("w:date"), "2026-08-28T10:00:00Z")
    paragraph.append(deletion)
    start = ET.Element(qn("w:commentRangeStart"), {qn("w:id"): "4"})
    paragraph.append(start)
    insertion = run("w:ins", "30")
    insertion.set(qn("w:id"), "8")
    insertion.set(qn("w:author"), "Bob Reviewer")
    insertion.set(qn("w:date"), "2026-08-28T10:01:00Z")
    paragraph.append(insertion)
    paragraph.append(ET.Element(qn("w:commentRangeEnd"), {qn("w:id"): "4"}))
    paragraph.append(run("w:r", " days."))
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    parts["word/comments.xml"] = (b"<not-xml" if malformed_comments else f"""
      <w:comments xmlns:w="{W}">
        <w:comment w:id="4" w:author="Carol Commenter" w:date="2026-08-28T10:02:00Z">
          <w:p><w:r><w:t>Confirm the payment period.</w:t></w:r></w:p>
        </w:comment>
      </w:comments>""".encode())
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return output.getvalue()


def _stale_comment_range_docx() -> bytes:
    source = io.BytesIO()
    document = Document()
    document.add_paragraph("Ordinary visible text.")
    document.save(source)
    source.seek(0)
    with zipfile.ZipFile(source) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    root = ET.fromstring(parts["word/document.xml"])
    paragraph = root.find(".//w:body/w:p", NS)
    assert paragraph is not None
    paragraph.insert(0, ET.Element(qn("w:commentRangeStart"), {qn("w:id"): "99"}))
    paragraph.append(ET.Element(qn("w:commentRangeEnd"), {qn("w:id"): "99"}))
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return output.getvalue()


def _docx_package(parts: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return output.getvalue()


def test_docx_xml_part_limit_is_checked_before_python_docx(monkeypatch):
    monkeypatch.setattr(ingestion, "DOCX_MAX_XML_PART_BYTES", 8, raising=False)
    monkeypatch.setattr(ingestion, "DocxDocument", lambda *_: pytest.fail("fallback must not run"))
    data = _docx_package({"word/document.xml": b"<document>longer-than-eight-bytes</document>"})

    with pytest.raises(ValueError, match="DOCX"):
        ingestion.IngestionService._extract("limited.docx", data)


def test_docx_relationship_part_limit_is_case_insensitive(monkeypatch):
    monkeypatch.setattr(ingestion, "DOCX_MAX_XML_PART_BYTES", 16)
    data = _docx_package({
        "word/document.xml": b"<document/>",
        "word/_rels/document.XML.RELS": b"x" * 17,
    })

    with pytest.raises(ValueError, match="XML part"):
        ingestion._preflight_docx(data)


@pytest.mark.parametrize(
    ("constant", "value", "parts", "message"),
    [
        ("DOCX_MAX_MEMBERS", 1, {"word/document.xml": b"<document/>", "extra.bin": b"x"}, "members"),
        ("DOCX_MAX_EXPANDED_BYTES", 32, {"word/document.xml": b"<document/>", "extra.bin": b"x" * 40}, "expands"),
    ],
)
def test_docx_package_limits_reject_small_in_memory_archives(monkeypatch, constant, value, parts, message):
    monkeypatch.setattr(ingestion, constant, value)

    with pytest.raises(ValueError, match=message):
        ingestion.IngestionService._extract("limited.docx", _docx_package(parts))


def test_docx_encrypted_member_is_rejected_before_parsing():
    data = bytearray(_docx_package({"word/document.xml": b"<document/>"}))
    central_directory = data.index(b"PK\x01\x02")
    data[central_directory + 8] |= 0x01

    with pytest.raises(ValueError, match="encrypted"):
        ingestion.IngestionService._extract("encrypted.docx", bytes(data))


def test_docx_entity_rejection_is_terminal_and_never_uses_python_docx(monkeypatch):
    monkeypatch.setattr(ingestion, "DocxDocument", lambda *_: pytest.fail("fallback must not run"))
    data = _docx_package({
        "word/document.xml": (
            b'<!DOCTYPE root [<!ENTITY injected "not safe">]>'
            b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            b"<w:body><w:p><w:r><w:t>&injected;</w:t></w:r></w:p></w:body></w:document>"
        ),
    })

    with pytest.raises(ValueError, match="DOCX"):
        ingestion.IngestionService._extract("entity.docx", data)


@pytest.mark.asyncio
async def test_imports_authors_revisions_dates_and_comment_ranges(app_context):
    result = await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON", UploadFile(file=io.BytesIO(_reviewed_docx()), filename="reviewed.docx")
    )
    companion = app_context.vault.read_markdown(result["extracted_path"])
    review = companion["metadata"]["review"]

    assert "Payment is 30 days." in companion["content"]
    assert {item["name"] for item in review["authors"]} == {
        "Alice Reviewer", "Bob Reviewer", "Carol Commenter"
    }
    assert [item["color"] for item in review["authors"]] == ["#2F5597", "#7030A0", "#008272"]
    deletion = next(item for item in review["segments"] if item["kind"] == "delete")
    insertion = next(item for item in review["segments"] if item["kind"] == "insert")
    assert (deletion["text"], deletion["author_name"], deletion["created_at"]) == (
        "10", "Alice Reviewer", "2026-08-28T10:00:00Z"
    )
    assert (insertion["text"], insertion["author_name"], insertion["change_id"]) == (
        "30", "Bob Reviewer", "CHG-WORD-8"
    )
    thread = review["comments"][0]
    assert companion["content"][thread["anchor_start"]:thread["anchor_end"]] == "30"
    assert thread["entries"][0]["body"] == "Confirm the payment period."
    assert all(item["name"] != "Themis" for item in review["authors"])


@pytest.mark.asyncio
async def test_malformed_comments_keep_revision_text_without_inventing_comment_author(app_context):
    result = await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON",
        UploadFile(file=io.BytesIO(_reviewed_docx(malformed_comments=True)), filename="malformed.docx"),
    )
    companion = app_context.vault.read_markdown(result["extracted_path"])
    review = companion["metadata"]["review"]
    assert "Payment is 30 days." in companion["content"]
    assert review["comments"] == []
    assert {item["name"] for item in review["authors"]} == {"Alice Reviewer", "Bob Reviewer"}


@pytest.mark.asyncio
async def test_stale_comment_markers_without_comment_part_use_ordinary_extraction(app_context):
    result = await app_context.ingestion.upload_to_matter(
        "MAT-DEMO-BEACON",
        UploadFile(file=io.BytesIO(_stale_comment_range_docx()), filename="stale-comment.docx"),
    )
    companion = app_context.vault.read_markdown(result["extracted_path"])
    assert "Ordinary visible text." in companion["content"]
    assert "review" not in companion["metadata"]


def test_adjacent_same_author_word_replacement_uses_one_change_id():
    data = _reviewed_docx()
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    root = ET.fromstring(parts["word/document.xml"])
    insertion = root.find(".//w:ins", NS)
    deletion = root.find(".//w:del", NS)
    assert insertion is not None and deletion is not None
    insertion.set(qn("w:author"), "Alice Reviewer")
    insertion.set(qn("w:date"), deletion.get(qn("w:date")))
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)

    _, review = app_contextless_extract(output.getvalue())
    changes = [item["change_id"] for item in review["segments"] if item["kind"] in {"delete", "insert"}]
    assert changes[0] == changes[1]


def app_contextless_extract(data: bytes):
    from app.services.ingestion import IngestionService

    extracted = IngestionService._extract("reviewed.docx", data)
    assert extracted[1] is not None
    return extracted


def test_adjacent_different_author_word_changes_keep_separate_ids():
    _, review = app_contextless_extract(_reviewed_docx())
    changes = [item["change_id"] for item in review["segments"] if item["kind"] in {"delete", "insert"}]
    assert changes == ["CHG-WORD-7", "CHG-WORD-8"]


def _stored_review() -> dict:
    return {
        "version": 2,
        "tracking": True,
        "authors": [
            {"author_id": "author-alice", "name": "Alice Reviewer", "color": "#2F5597"},
            {"author_id": "author-bob", "name": "Bob Reviewer", "color": "#7030A0"},
        ],
        "segments": [
            {"kind": "equal", "text": "Payment is ", "change_id": "", "author_id": "", "author_name": "", "author_color": "", "created_at": ""},
            {"kind": "delete", "text": "10", "change_id": "CHG-1", "author_id": "author-alice", "author_name": "Alice Reviewer", "author_color": "#2F5597", "created_at": "2026-08-28T10:00:00Z"},
            {"kind": "insert", "text": "30", "change_id": "CHG-1", "author_id": "author-bob", "author_name": "Bob Reviewer", "author_color": "#7030A0", "created_at": "2026-08-28T10:01:00Z"},
            {"kind": "equal", "text": " days.", "change_id": "", "author_id": "", "author_name": "", "author_color": "", "created_at": ""},
        ],
        "comments": [{
            "thread_id": "COM-1", "quote": "30", "anchor_start": 11, "anchor_end": 13,
            "resolved": True, "resolved_at": "2026-08-28T11:00:00Z", "resolved_by": "Alice Reviewer",
            "entries": [
                {"comment_id": "MSG-1", "author_id": "author-alice", "author_name": "Alice Reviewer", "body": "Confirm period.", "created_at": "2026-08-28T10:02:00Z"},
                {"comment_id": "MSG-2", "author_id": "author-bob", "author_name": "Bob Reviewer", "body": "Confirmed.", "created_at": "2026-08-28T10:03:00Z"},
            ],
        }],
        "comment_events": [{"event_id": "EVT-1", "thread_id": "COM-DELETED", "action": "deleted", "actor": "Alice Reviewer", "created_at": "2026-08-28T12:00:00Z"}],
    }


def test_export_preserves_revision_attribution_flattens_resolved_comment_and_round_trips(app_context):
    app_context.vault.write_markdown(PATH, "Payment is 30 days.\n", {"review": _stored_review()})
    exported_path, _ = app_context.document_exports.export(PATH, "docx")
    data = app_context.vault.resolve(exported_path).read_bytes()

    with zipfile.ZipFile(io.BytesIO(data)) as package:
        document_xml = package.read("word/document.xml")
        comments_xml = package.read("word/comments.xml")
    root = ET.fromstring(document_xml)
    insertion = root.find(".//w:ins", NS)
    deletion = root.find(".//w:del", NS)
    assert insertion is not None and deletion is not None
    assert insertion.get(qn("w:author")) == "Bob Reviewer"
    assert insertion.get(qn("w:date")) == "2026-08-28T10:01:00Z"
    assert deletion.get(qn("w:author")) == "Alice Reviewer"
    assert deletion.get(qn("w:date")) == "2026-08-28T10:00:00Z"
    assert root.find(".//w:ins/w:r/w:rPr/w:color", NS) is None
    assert root.find(".//w:del/w:r/w:rPr/w:color", NS) is None
    assert b"Resolved" in comments_xml
    assert b"Alice Reviewer" in comments_xml and b"Bob Reviewer" in comments_xml
    assert b"Confirm period." in comments_xml and b"Confirmed." in comments_xml
    assert b"COM-DELETED" not in comments_xml

    imported = app_context.ingestion._extract("round-trip.docx", data)
    assert imported[1] is not None
    assert {item["author_name"] for item in imported[1]["segments"] if item["change_id"]} == {
        "Alice Reviewer", "Bob Reviewer"
    }
    assert imported[1]["comments"][0]["quote"] == "30"


def test_docx_comment_range_anchors_only_second_repeated_quote(app_context):
    review = {
        "version": 2, "tracking": True, "authors": [],
        "segments": [{"kind": "equal", "text": "term and term", "change_id": "", "author_id": "",
                      "author_name": "", "author_color": "", "created_at": ""}],
        "comments": [{"thread_id": "COM-2", "quote": "term", "anchor_start": 9, "anchor_end": 13,
                      "resolved": False, "resolved_at": "", "resolved_by": "", "entries": [{
                          "comment_id": "MSG-3", "author_id": "author-alice", "author_name": "Alice Reviewer",
                          "body": "Second term only.", "created_at": "2026-08-28T10:00:00Z"}]}],
        "comment_events": [],
    }
    app_context.vault.write_markdown(PATH, "term and term", {"review": review})
    exported_path, _ = app_context.document_exports.export(PATH, "docx")
    with zipfile.ZipFile(app_context.vault.resolve(exported_path)) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    paragraph = root.find(".//w:body/w:p", NS)
    assert paragraph is not None
    inside = False
    anchored = []
    for child in paragraph:
        if child.tag == qn("w:commentRangeStart"):
            inside = True
        elif child.tag == qn("w:commentRangeEnd"):
            inside = False
        elif inside:
            anchored.extend(node.text or "" for node in child.findall(".//w:t", NS))
    assert "".join(anchored) == "term"
    children = list(paragraph)
    marker_index = next(index for index, child in enumerate(children) if child.tag == qn("w:commentRangeStart"))
    text_before = "".join(node.text or "" for child in children[:marker_index] for node in child.findall(".//w:t", NS))
    assert text_before == "term and "


def test_pdf_comment_annotation_uses_second_repeated_quote_anchor(app_context):
    review = {
        "version": 2, "tracking": True, "authors": [],
        "segments": [{"kind": "equal", "text": "term and term", "change_id": "", "author_id": "",
                      "author_name": "", "author_color": "", "created_at": ""}],
        "comments": [{"thread_id": "COM-3", "quote": "term", "anchor_start": 9, "anchor_end": 13,
                      "resolved": False, "resolved_at": "", "resolved_by": "", "entries": [{
                          "comment_id": "MSG-4", "author_id": "author-alice", "author_name": "Alice Reviewer",
                          "body": "Second term only.", "created_at": "2026-08-28T10:00:00Z"}]}],
        "comment_events": [],
    }
    app_context.vault.write_markdown(PATH, "term and term", {"review": review})
    exported_path, _ = app_context.document_exports.export(PATH, "pdf")
    reader = PdfReader(app_context.vault.resolve(exported_path))
    highlights = [item.get_object() for item in reader.pages[0]["/Annots"]
                  if item.get_object().get("/Subtype") == "/Highlight"]
    assert len(highlights) == 1
    assert float(highlights[0]["/Rect"][0]) > 80
