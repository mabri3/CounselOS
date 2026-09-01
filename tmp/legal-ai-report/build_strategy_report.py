"""Build the strategic Themis/CounselOS legal-AI market report."""

from __future__ import annotations

import json
import shutil
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from build_report import (
    BLUE,
    DARK_BLUE,
    INK,
    LIGHT_GRAY,
    MUTED,
    USABLE_DXA,
    add_body,
    add_bullet,
    add_hyperlink,
    add_number,
    callout,
    heading,
    set_font,
    set_paragraph,
    set_table_geometry,
    shade,
    simple_table,
)


ROOT = Path(__file__).resolve().parents[2]
REFERENCE = ROOT / "output/research/legal-ai-small-team-market-research.docx"
CORPUS = ROOT / "output/research/raw/legal-ai-reddit-source-corpus.jsonl"
OUT = ROOT / "output/research/themis-counselos-legal-ai-competitive-strategy.docx"


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def tighten_table(table) -> None:
    set_repeat_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)


def add_link_line(doc, lead: str, links: list[tuple[str, str]], after: int = 8) -> None:
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=after, line=1.0)
    set_font(p.add_run(lead), 8.5, MUTED, True)
    for index, (label, url) in enumerate(links):
        if index:
            set_font(p.add_run(" | "), 8.5, MUTED)
        add_hyperlink(p, label, url)


def add_definition(doc, label: str, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph(p, after=5, line=1.1)
    set_font(p.add_run(f"{label}: "), 11, DARK_BLUE, True)
    set_font(p.add_run(text), 11, INK)


def page(doc) -> None:
    doc.add_page_break()


def source(records: dict[str, dict], source_id: str) -> tuple[str, str]:
    item = records[source_id]
    return item["title"], item["url"]


def build() -> None:
    records_list = [json.loads(line) for line in CORPUS.read_text().splitlines() if line.strip()]
    records = {item["source_id"]: item for item in records_list}
    subreddits = Counter(item.get("subreddit", "not stated") for item in records_list)
    recent_count = 100
    sized_count = 22

    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUT)
    doc = Document(OUT)
    clear_body(doc)

    # Cover
    p = doc.add_paragraph()
    set_paragraph(p, before=10, after=4, line=1.0)
    set_font(p.add_run("COMPETITIVE STRATEGY BRIEF"), 10, BLUE, True)
    p = doc.add_paragraph()
    set_paragraph(p, after=5, line=1.0)
    set_font(p.add_run("How Themis / CounselOS can win the legal AI work layer"), 22, INK, True)
    p = doc.add_paragraph()
    set_paragraph(p, after=16, line=1.0)
    set_font(
        p.add_run(
            "Why teams choose ChatGPT, Claude, Harvey, Legora, CoCounsel, GC AI, or an internal build - and the product wedge that remains open"
        ),
        13,
        MUTED,
    )

    metadata = doc.add_table(rows=5, cols=2)
    metadata.style = "Table Grid"
    set_table_geometry(metadata, [1500, 7860])
    metadata_values = [
        ("Prepared for", "Themis.ai / CounselOS product strategy"),
        ("Date", "August 31, 2026"),
        ("Corpus", f"{len(records_list)} de-duplicated Reddit threads; {recent_count} added sources from March-August 2026"),
        ("Communities", f"r/legaltech ({subreddits['r/legaltech']}), r/biglaw ({subreddits['r/biglaw']}), r/Lawyertalk ({subreddits['r/Lawyertalk']}), r/LawFirm ({subreddits['r/LawFirm']}), and smaller lawyer-centered groups"),
        ("Name note", "Themis.ai is the current product name. CounselOS is used here for the broader product strategy and research work."),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_values):
        shade(row.cells[0], LIGHT_GRAY)
        p1 = row.cells[0].paragraphs[0]
        set_paragraph(p1, after=0, line=1.0)
        set_font(p1.add_run(label), 9.8, INK, True)
        p2 = row.cells[1].paragraphs[0]
        set_paragraph(p2, after=0, line=1.0)
        set_font(p2.add_run(value), 9.8, INK)
        prevent_row_split(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(5)
    callout(
        doc,
        "Strategic verdict",
        "Themis can win, but not as another legal chatbot. Its best opening is a model-neutral matter workspace for small in-house and boutique legal teams. It should let the lawyer keep Claude or ChatGPT, bring in authoritative research, and add what chat does not preserve: matter memory, proved document coverage, review-ready work product, and a visible next action.",
    )

    # Page 2 - verdict
    page(doc)
    heading(doc, "The answer in one page", 1)
    add_body(
        doc,
        "The strongest conclusion from the full corpus is that legal AI is not one market. Buyers assemble a stack. They use a general model for thinking and drafting, a controlled database for legal authority, a specialist tool for selected document workflows, and a matter or practice system for records. Most buying disappointment comes from expecting one product to do all four jobs.",
    )
    add_body(
        doc,
        "This creates a real opening for Themis. The open position is not 'better answers than Claude.' It is 'turn the lawyer's preferred model into reliable matter work without forcing the lawyer to rebuild context, prove coverage by hand, or maintain another administrative system.'",
    )
    table = simple_table(
        doc,
        ["What the evidence says", "What it means for Themis"],
        [
            ("ChatGPT and Claude are the default for broad drafting, reasoning, and flexible work.", "Support them. Do not ask buyers to abandon the models they already trust."),
            ("Harvey and Legora earn value through deployment, review tables, document scale, sources, and integrations.", "Compete on a narrower small-team workflow, not enterprise feature breadth."),
            ("CoCounsel earns value where Westlaw-linked research and litigation workflow matter.", "Integrate or hand off to authority. Do not claim to replace Westlaw or Lexis."),
            ("Internal builds win when firms need exact DMS, precedent, or practice-area fit.", "Be configurable enough to avoid a custom build, but much easier to maintain."),
            ("Users reject any product that creates a second review job or another stale database.", "Every feature must reduce lawyer review or repeated setup in a visible way."),
        ],
        [4600, 4760],
        9.4,
    )
    tighten_table(table)
    heading(doc, "Three conclusions that should guide the product", 1)
    add_number(doc, "The model is a component, not the product. Model quality changes too quickly to be a stable wedge.")
    add_number(doc, "Proof is more valuable than legal branding. Buyers want to see the passage, file coverage, exception, and review status.")
    add_number(doc, "The durable unit of value is the matter. Context, work product, decisions, and next action must survive every chat and model change.")
    callout(
        doc,
        "Position to own",
        "Themis is the matter operating layer between general AI, authoritative legal research, Word, and the systems where the team already works.",
    )

    # Page 3 - market structure
    page(doc)
    heading(doc, "What buyers are actually buying", 1)
    add_body(
        doc,
        "A feature comparison hides the real decision. Buyers are assigning different jobs to different products. The product that wins a job is not always the product with the best answer in a prompt test.",
    )
    table = simple_table(
        doc,
        ["Buyer job", "Current leader", "Why it wins", "Remaining gap"],
        [
            ("Think, draft, explain", "ChatGPT / Claude", "Strong general reasoning, fast writing, flexible instructions, current models, low entry cost.", "No durable matter state; weak proof of file coverage and authority."),
            ("Find legal authority", "Westlaw / Lexis / controlled sources", "Known corpus, links to authority, citator and research habits.", "Research output is still separate from matter work and follow-through."),
            ("Run repeatable document work", "Harvey / Legora / CoCounsel / point tools", "Review tables, large-document work, Word or DMS connections, legal defaults.", "Price, rigid workflows, uneven task performance, and another system."),
            ("Apply firm-specific methods", "Internal build / configured Claude", "Exact playbooks, precedent, connectors, and control.", "Build cost, maintenance, technical skill, adoption, and ownership risk."),
            ("Remember and move the matter", "No clear winner for small teams", "Practice systems store records but rarely create a legal foothold or next action.", "This is the opening for Themis."),
        ],
        [1800, 1700, 3020, 2840],
        8.85,
    )
    tighten_table(table)
    heading(doc, "The market is split by deployment burden", 1)
    add_definition(doc, "General-model path", "Best for capable users with varied work who will supply context, prompts, and legal verification themselves.")
    add_definition(doc, "Specialist-suite path", "Best for larger teams that will pay for standardization, integrations, approved deployment, and broad adoption.")
    add_definition(doc, "Internal-build path", "Best for technically able firms with a narrow high-value workflow and strong firm-specific data needs.")
    add_definition(doc, "Themis path", "Best for a small legal team that wants the flexibility of general AI but cannot afford repeated context setup, scattered work product, and manual follow-through.")
    add_body(
        doc,
        "This last path is a strategic inference from the evidence and the current Themis product thesis. Reddit does not prove market size or willingness to pay. It does show a repeated unmet workflow pattern that can be tested.",
    )

    # Page 4 - general models
    page(doc)
    heading(doc, "Why teams choose ChatGPT or Claude", 1)
    add_body(
        doc,
        "Small teams are not choosing general models because they do not care about legal workflow. They choose them because the first unit of value is useful work now, at low cost, across many tasks.",
    )
    for title, text in [
        ("1. Better price-to-value", "A low-cost subscription can cover drafting, analysis, client communication, operations, coding, and nonlegal work. A specialist contract must create enough additional saved work to justify a second subscription and a second workflow."),
        ("2. Strong raw drafting and broad reasoning", "Several users preferred Claude or ChatGPT for transactional drafting, mixed legal-commercial analysis, correspondence, or issue spotting. Direct comparisons did not show a consistent reasoning advantage for the legal suites."),
        ("3. Faster access to current model capability", "Experienced users want the newest model, their own instructions, and freedom to change providers. They do not want a legal vendor to decide the model, context window, or workflow."),
        ("4. Familiar, low-friction adoption", "Lawyers can start without a rollout, migration, minimum seat count, or new matter database. This matters most in a solo practice or small team with uneven AI use."),
        ("5. Easy customization", "Technically capable users build skills, agents, connectors, and repeatable prompts around their own precedent. They view a fixed prompt library as weak paid value."),
    ]:
        heading(doc, title, 2)
        add_body(doc, text)
    callout(
        doc,
        "Why they still leave general chat",
        "General models lose trust when sources are invented, a file is skipped, context resets, usage limits interrupt a large matter, or the output must be rebuilt in Word. Confidentiality and approved deployment can also block consumer accounts.",
    )
    add_link_line(
        doc,
        "Representative evidence: ",
        [
            ("3-attorney comparison", records["reddit-110"]["url"]),
            ("transactional drafting with Claude", records["reddit-131"]["url"]),
            ("Claude litigation workflow", records["reddit-057"]["url"]),
            ("Claude plus legal connectors", records["reddit-154"]["url"]),
            ("general AI weekly work", records["reddit-166"]["url"]),
        ],
    )

    # Page 5 - specialist products
    page(doc)
    heading(doc, "Why specialist products win - and why they lose", 1)
    table = simple_table(
        doc,
        ["Option", "Why a buyer chooses it", "Why a buyer rejects it", "Strategic meaning"],
        [
            ("Harvey", "Turnkey legal workflows, review tables, large files, Word/DMS and legal-source integrations, enterprise rollout.", "High or unclear pricing, minimum terms, mixed task quality, and weak value over Claude/ChatGPT for skilled users.", "Wins on deployment and workflow breadth more than proven model superiority."),
            ("Legora", "Drafting, diligence, tabular review, templates, Word workflow, standard team use.", "Price, output rigidity, mixed model-quality reports, and convergence with general models.", "Strongest where standardized transactional review matters across a team."),
            ("CoCounsel", "Westlaw-linked research, litigation tasks, document review, discovery, deposition and brief work.", "Generative limits, task-following failures, price increases, and continued checking burden.", "Defensible through authority and litigation workflow; weaker as a general drafting system."),
            ("GC AI", "Legal defaults, playbooks, drafting, contract context, repository work, fast daily use for some users.", "Sparse and polarized evidence; unclear advantage over configured general models for the price.", "Potentially strong legal UX, but the corpus does not establish a clear small-team moat."),
            ("Internal build", "Exact practice-area, DMS, precedent, cost, and model control.", "Maintenance, technical talent, governance, testing, and adoption risk.", "A major competitor for 20-100 lawyer firms with a capable technology lead."),
        ],
        [1250, 2840, 2780, 2490],
        8.45,
    )
    tighten_table(table)
    heading(doc, "The strongest counter-evidence", 1)
    add_body(
        doc,
        "The specialist products are not mere wrappers in every setting. One corporate legal department selected Harvey for its examples, references, interface, and integrations even though answers were not consistently better than ChatGPT. A user at a roughly 5,000-person firm reported meaningful time saved and good property-diligence results. BigLaw users reported strong Harvey regulatory work, Legora tabular review, and better adoption when a domain-specific product made prompting easier.",
    )
    add_body(
        doc,
        "The correct conclusion is therefore not 'general AI always wins.' It is that specialist tools must earn their premium through operational leverage, approved deployment, or scaled adoption. The value case weakens sharply when a capable small team can recreate the same result with a general model and its existing research tools.",
    )
    add_link_line(
        doc,
        "Counter-evidence: ",
        [
            ("corporate Harvey selection", records["reddit-103"]["url"]),
            ("5,000-person firm Harvey use", records["reddit-116"]["url"]),
            ("BigLaw Harvey experience", records["reddit-087"]["url"]),
            ("Harvey versus Legora", records["reddit-097"]["url"]),
            ("legal research discussion", records["reddit-084"]["url"]),
        ],
    )

    # Page 6 - direct decisions
    page(doc)
    heading(doc, "Direct buying and rejection evidence", 1)
    add_body(
        doc,
        "These are the most decision-relevant cases in the corpus and the earlier research. They name a buyer, a tested workflow, a selection or rejection reason, or an observable result. Team size is shown only when the source identified it.",
    )
    table = simple_table(
        doc,
        ["Buyer / size", "Decision or result", "Reason that mattered"],
        [
            ("GC; 3-attorney legal department", "Preferred ChatGPT over Harvey and Claude for the tested prompts; rejected about $30,000/year Harvey value.", "Reasoning, citations, Word documents, redlines, and inconsistent Harvey issue coverage."),
            ("Law-firm technology lead; about 30 attorneys", "Rejected Harvey, Legora, and CoCounsel after three-week trials; chose an internal build.", "All three missed a material issue that Claude Opus caught; DMS/project friction; still needed Westlaw."),
            ("Technology committee; 100+ lawyer firm", "Selected Claude Enterprise plus a smaller research tool after an earlier Harvey trial.", "Harvey did not justify cost; Claude offered flexibility. Counterpoint: specialist UX can aid firm-wide use."),
            ("Corporate legal department; size not stated", "Selected Harvey despite no consistent answer-quality lead.", "Better examples, references, interface, Word integration, legal databases, and deployment fit."),
            ("Lawyer; roughly 5,000-person firm", "Used Harvey regularly and reported time saved and good diligence work.", "Ease of use, document-scale work, and enterprise access; still called it too expensive for small firms."),
            ("Small California firm; size not stated", "Rejected Harvey.", "Expensive, inflexible pricing and no meaningful small-team trial or customization path."),
            ("Small firm; size not stated", "Considered leaving CoCounsel after heavy use.", "Price increases and paid upgrades weakened value; cheaper AI plus classic research remained an option."),
            ("Transactional lawyers; size not stated", "Reported hours saved with Claude; one mock test favored Claude Business over CoCounsel.", "Defined terms, precedent cleanup, cross-references, first drafts, and customizable agents."),
            ("BigLaw users; size not stated", "Reported Harvey and Legora similar for one research task; mixed product preferences.", "Review-table clarity, drafting quality, and ease of use mattered more than a clear model winner."),
        ],
        [2300, 3400, 3660],
        8.35,
    )
    tighten_table(table)
    add_link_line(
        doc,
        "Primary links: ",
        [
            ("3-attorney", records["reddit-110"]["url"]),
            ("30-attorney", records["reddit-175"]["url"]),
            ("100+ lawyer", records["reddit-168"]["url"]),
            ("corporate Harvey", records["reddit-103"]["url"]),
            ("5,000-person firm", records["reddit-116"]["url"]),
            ("small-firm rejection", records["reddit-107"]["url"]),
            ("CoCounsel price objection", records["reddit-109"]["url"]),
            ("Claude transactional drafting", records["reddit-131"]["url"]),
            ("Harvey vs. Legora", records["reddit-097"]["url"]),
        ],
    )

    # Page 7 - target segments
    page(doc)
    heading(doc, "Who Themis should pursue", 1)
    add_body(
        doc,
        "Themis should not start with the whole legal market. Its current product and the evidence point to one beachhead and two adjacent segments.",
    )
    table = simple_table(
        doc,
        ["Segment", "Current default", "Fit for Themis", "Why"],
        [
            ("Beachhead: solo GC and 1-5 person product/legal teams", "ChatGPT or Claude plus outside counsel and a research database", "Highest", "Frequent ambiguous intake, limited junior support, company-specific history, and no appetite for enterprise rollout."),
            ("Adjacent: 5-30 lawyer boutique or specialty firms", "General AI plus Westlaw/Lexis, Word, and point tools", "High if workflow is narrow", "Need repeatable document work and matter continuity, but reject high seat minimums and broad suites."),
            ("Adjacent later: 30-100 lawyer firms with a technology lead", "Claude Enterprise, specialist pilots, or internal build", "Conditional", "Themis must beat internal build on configuration, DMS handoff, maintenance, and team adoption."),
            ("Not now: 150+ lawyer enterprise / BigLaw", "Harvey, Legora, CoCounsel, internal platforms", "Low for current MVP", "Purchases depend on SSO, permissions, DMS, governance, support, and broad deployment."),
            ("Not as a replacement: research-heavy litigation", "Westlaw/Lexis/CoCounsel", "Complement only", "Themis can organize authority and work product, but should not promise a proprietary legal corpus or citator."),
        ],
        [2050, 2400, 1200, 3710],
        8.5,
    )
    tighten_table(table)
    heading(doc, "Beachhead problem statement", 1)
    callout(
        doc,
        "Target user's job",
        "A product lawyer receives an incomplete request, scattered files, prior decisions, and a deadline. The lawyer needs a reliable orientation, a strong first work product, the exact support behind it, and a clear decision or next action - without rebuilding the matter in every chat.",
    )
    heading(doc, "Why this segment can switch", 1)
    add_bullet(doc, "Themis does not require the team to stop using Claude, ChatGPT, Westlaw, Lexis, or Word.")
    add_bullet(doc, "The buyer can test value on one real matter without a firm-wide migration.")
    add_bullet(doc, "The product can show value before team collaboration, SSO, or enterprise governance is complete.")
    add_bullet(doc, "The workflow matches the current Themis product thesis: ambiguous intake to organized foothold, work product, decision, and next action.")

    # Page 8 - current capability and gap
    page(doc)
    heading(doc, "Can Themis win with the current product?", 1)
    add_body(
        doc,
        "Yes as a focused pilot; no as a broad market claim. The current product already supports the core matter concept, but several missing proof and workflow capabilities determine whether buyers see it as a product or another chat wrapper.",
    )
    table = simple_table(
        doc,
        ["Capability", "Current position", "Competitive meaning", "Required move"],
        [
            ("Matter continuity", "Strong", "Workspace, company and matter context, work items, decisions, events, and next action support the wedge.", "Make orientation automatic and visibly persistent after every chat and model change."),
            ("Model neutrality", "Strong", "Multiple providers and per-agent selection reduce model lock-in.", "Let the buyer use the preferred model without losing matter history or work product."),
            ("Editable work product", "Strong foundation", "Markdown-first editor, revisions, comments, finalization, and generated DOCX/PDF move beyond chat.", "Prove that output reaches a lawyer-ready result with low correction burden."),
            ("Source proof", "Partial", "Internal/external source labels exist, but buyers need fast verification.", "Add exact passage support, proposition links, source status, and lawyer-checked markers."),
            ("Document coverage", "Material gap", "Silent omissions are a recurring trust failure across tools.", "Show every file/page reviewed, extraction failures, missing items, and exceptions."),
            ("Structured batch review", "Material gap", "Review tables are a clear reason buyers pay Harvey or Legora.", "Create reusable answer tables with citations, exceptions, and unresolved questions."),
            ("Word workflow", "Partial", "Themis can regenerate Word/PDF, but does not preserve the source Word layout or support a live Word round trip.", "First prove clean review-ready export; later add source-layout preservation or a live Word path."),
            ("Daily-work intake", "Partial", "Inbox watcher and schedules exist; email/Slack/Office integrations do not.", "Add low-noise capture only after manual intake is a measured bottleneck."),
            ("Team/enterprise controls", "Not targeted", "The local MVP cannot compete for broad enterprise deployment.", "Delay until the single-user matter workflow proves pull."),
        ],
        [1600, 1250, 3320, 3190],
        8.15,
    )
    tighten_table(table)
    add_body(
        doc,
        "This assessment is based on the current PRD, current project state, and verified reliability handoff. It is not a claim that every feature has been validated with paying users.",
    )

    # Page 9 - product bets
    page(doc)
    heading(doc, "The three product bets that create a right to win", 1)
    heading(doc, "1. The matter proof layer", 2)
    add_body(
        doc,
        "For every material statement or extracted answer, show the exact source passage, source type, file and page, review status, and any conflict. For every document job, show a coverage receipt and exception list. This converts 'the AI probably read it' into visible work the lawyer can check quickly.",
    )
    heading(doc, "2. The continuity engine", 2)
    add_body(
        doc,
        "Create and maintain a concise matter dossier from confirmed intake, files, company context, prior decisions, open questions, completed work, and next action. It must update from normal work. It must not ask the lawyer to maintain a second database.",
    )
    heading(doc, "3. The work-product bridge", 2)
    add_body(
        doc,
        "Turn the proof layer and matter dossier into a review-ready artifact: a structured review table, advice memo, issue list, redline, or response. Keep source support and unresolved issues connected through export. Bring lawyer-checked Westlaw/Lexis research into the matter instead of trying to replace it.",
    )
    table = simple_table(
        doc,
        ["If this works", "Why the buyer pays", "Why ChatGPT/Claude alone do not replace it"],
        [
            ("The lawyer verifies an answer in minutes", "Review burden falls, not just drafting time.", "General chat usually does not maintain a claim-to-source review record across the matter."),
            ("The system proves what it reviewed", "The lawyer can trust process coverage and focus on exceptions.", "A long context window is not a coverage receipt."),
            ("The matter survives the model", "The team can change providers without starting again.", "Projects and chats remain provider-bound and often need manual upkeep."),
            ("The output reaches Word and the next action", "The work advances instead of becoming another chat transcript.", "The model answer alone does not own matter state, decision records, or follow-through."),
        ],
        [2950, 3000, 3410],
        9.0,
    )
    tighten_table(table)
    callout(
        doc,
        "Feature rule",
        "A Themis feature should be funded only if it removes repeated setup, reduces review work, proves coverage, or moves the matter to a decision or next action.",
    )

    # Page 10 - demo and product priorities
    page(doc)
    heading(doc, "The product demonstration that should prove the thesis", 1)
    add_body(
        doc,
        "Do not demonstrate a broad menu of legal AI features. Demonstrate one matter from raw request to review-ready foothold.",
    )
    for step_number, (step_label, step_text) in enumerate([
        ("Intake once", "The lawyer pastes an email or request, adds the files, and selects Claude, ChatGPT, or another approved model."),
        ("Orient automatically", "Themis creates the business objective, facts, missing facts, issues, relevant prior decisions, file inventory, and next action."),
        ("Do the first work", "Themis runs targeted research or document review and produces a structured, editable result."),
        ("Prove it", "Every material point links to an exact passage. The system shows file and page coverage, unreadable items, conflicts, and exceptions."),
        ("Hand it off", "The lawyer receives a review-ready Word/PDF artifact, records the recommendation or decision, and sees the next action. Switching models does not reset the matter."),
    ], start=1):
        add_definition(doc, f"{step_number}. {step_label}", step_text)
    heading(doc, "Priority order", 1)
    table = simple_table(
        doc,
        ["Order", "Build or prove", "Success observation"],
        [
            ("1", "Exact source passages plus coverage receipt", "A lawyer can identify what supports the answer and what was not reviewed without reopening every file."),
            ("2", "Automatic matter dossier and next action", "A returning lawyer understands the matter in seconds without rereading chat history."),
            ("3", "Structured multi-document review", "The result includes answers, source support, exceptions, and missing data in an editable table."),
            ("4", "Review-ready Word/PDF handoff", "The lawyer can use the artifact without rebuilding structure or losing the support trail."),
            ("5", "Authoritative research import/handoff", "A Westlaw/Lexis source can be attached to a proposition and marked lawyer-checked."),
            ("6", "Low-noise daily intake", "Confirmed email or pasted intake creates or updates a matter without duplicate data entry."),
        ],
        [700, 3100, 5560],
        9.05,
    )
    tighten_table(table)

    # Page 11 - GTM
    page(doc)
    heading(doc, "How Themis should sell against each alternative", 1)
    table = simple_table(
        doc,
        ["Alternative", "Buyer objection", "Themis response", "Proof required"],
        [
            ("ChatGPT / Claude", "We already have a strong model.", "Keep it. Themis makes the model remember the matter, prove its work, and finish the handoff.", "Same preferred model; less context setup and faster verification on a real matter."),
            ("Harvey / Legora", "We want legal workflow, but the suite is too broad or costly.", "Use a focused matter workflow with no enterprise rollout or forced model choice.", "One-matter trial, clear pricing, structured review, proof layer, and Word-ready output."),
            ("CoCounsel / Westlaw", "We need authoritative research.", "Keep the research system. Themis carries checked authority into the matter, work product, decision, and next action.", "Reliable source import, proposition links, and lawyer-checked status."),
            ("Internal build", "We need exact fit and control.", "Configure models, playbooks, agents, and matter records without owning an AI platform.", "Fast setup, open/editable configuration, stable data export, and lower maintenance burden."),
            ("Practice or matter system", "We already store files and tasks.", "Themis creates orientation and work product from the matter; it is not another passive file cabinet.", "Automatic dossier updates and visible work completed by the agent."),
        ],
        [1500, 2170, 3170, 2520],
        8.35,
    )
    tighten_table(table)
    heading(doc, "Recommended message", 1)
    callout(
        doc,
        "Product promise",
        "Give Themis the matter once. Use the model you prefer. Get review-ready work with traceable sources, proved document coverage, preserved context, and a clear next action.",
    )
    heading(doc, "Packaging principles", 1)
    add_bullet(doc, "Let a buyer prove value on one matter before a team rollout.")
    add_bullet(doc, "Avoid a high seat minimum or long commitment in the beachhead segment.")
    add_bullet(doc, "Make model and external research costs clear. Do not hide them inside an unclear legal-AI premium.")
    add_bullet(doc, "Sell saved setup and review work, not access to prompts or a legal label.")
    add_bullet(doc, "Offer a clean export path so the buyer does not fear product lock-in.")

    # Page 12 - risks and validation
    page(doc)
    heading(doc, "What would cause Themis to lose", 1)
    add_body(
        doc,
        "The strategy has real risk. The strongest objection is that Claude Projects, ChatGPT Projects, or a configured internal workflow may become good enough. Themis therefore needs evidence that the matter layer changes the lawyer's work, not just the screen around the model.",
    )
    table = simple_table(
        doc,
        ["Failure condition", "How it appears", "How to test it"],
        [
            ("It feels like another chat wrapper", "Users compare answers but do not return to the matter workspace.", "Test whether users reopen the same matter and rely on its dossier, sources, and next action."),
            ("Matter state becomes manual administration", "Facts, tasks, and decisions become stale because lawyers must update them twice.", "Measure automatic updates and corrections per matter; watch for duplicate entry."),
            ("Proof does not reduce review", "Lawyers still reopen every file and rebuild the answer themselves.", "Compare verification time and missed exceptions with and without exact passages and coverage."),
            ("Word handoff breaks the workflow", "Users copy-paste output or abandon formatting and tracked changes.", "Use real work product and measure reformatting and correction burden."),
            ("The target requires enterprise controls now", "Security review blocks every pilot before the matter value is seen.", "Recruit the intended solo/small-team beachhead and separate pilot blockers from later enterprise needs."),
            ("The wedge is too broad", "Users like individual features but cannot name the job Themis owns.", "Ask users to describe the product after one matter; the answer should center on matter continuity and review-ready proof."),
        ],
        [2450, 3210, 3700],
        8.55,
    )
    tighten_table(table)
    heading(doc, "Validation plan", 1)
    add_bullet(doc, "Run the same three real matter types with at least five target users: product launch advice, multi-document review, and a recurring follow-up matter.")
    add_bullet(doc, "Measure time to orientation, time to first useful artifact, lawyer correction burden, source-verification time, and whether the next action remains visible after return.")
    add_bullet(doc, "Compare Themis with the user's normal ChatGPT or Claude workflow. Do not compare only with specialist vendor demos.")
    add_bullet(doc, "Ask what the user would remove, what they would pay to keep, and which existing tool they would stop using - if any.")
    callout(
        doc,
        "Go / no-go test",
        "If target users can reproduce the same matter continuity, proof, and work-product handoff in their existing model workflow with little effort, Themis does not yet have a paid wedge. If Themis cuts repeated setup and verification while preserving model choice, it has a credible one.",
    )

    # Page 13 - do not build and evidence limits
    page(doc)
    heading(doc, "What Themis should not build first", 1)
    for item in [
        "A proprietary legal model or a claim of universally superior legal reasoning.",
        "A broad prompt library presented as the main paid value.",
        "A new legal research corpus or citator before authoritative-source handoff proves insufficient.",
        "A full practice-management, billing, contract-lifecycle, or document-management replacement.",
        "Complex multi-agent voting, confidence gates, or legal-answer theater that delays useful work.",
        "Enterprise governance, SSO, permissions, and multi-tenant architecture before the beachhead workflow proves pull.",
        "Broad integrations before manual intake, Word handoff, or research import becomes a measured bottleneck.",
    ]:
        add_bullet(doc, item)
    heading(doc, "Evidence strength and limits", 1)
    add_body(
        doc,
        f"The complete corpus contains {len(records_list)} de-duplicated Reddit threads. Product discussion coverage includes 85 records mentioning ChatGPT, 73 Claude, 60 Harvey, 30 CoCounsel, 29 Legora, and 5 GC AI. These are mention counts, not preference or market-share measures. Only about {sized_count} records identify a numeric team size or a solo practice.",
    )
    add_body(
        doc,
        "Reddit evidence is anonymous, self-selected, and not independently verified. Product versions, prices, model access, and negotiated terms vary. Some posts contain vendor, founder, affiliate, or hearsay claims. The report treats direct task tests and named team contexts as stronger evidence and uses broad sentiment only to identify hypotheses.",
    )
    add_body(
        doc,
        "The strategic recommendations are an inference from the corpus plus the current Themis product scope. They are not proof of market size, willingness to pay, or product-market fit. The raw 203-record file remains the audit trail and should be used to inspect any conclusion back to its source link.",
    )
    callout(
        doc,
        "Bottom line after a clear-eyed review",
        "Themis has a credible path to win a narrow market. It does not yet have permission to claim broad legal-AI superiority. The winning product is the smallest system that makes a preferred model dependable across a matter: oriented, source-linked, coverage-proved, editable, and moving toward a decision.",
    )

    # Curated sources
    page(doc)
    heading(doc, "Curated evidence appendix", 1)
    add_body(
        doc,
        "These links are the primary records used for the competitive conclusions. The complete 203-record corpus is stored separately at output/research/raw/legal-ai-reddit-source-corpus.jsonl. Links may change or be removed by Reddit.",
    )
    groups = [
        ("Direct selection, rejection, and pricing", ["reddit-103", "reddit-107", "reddit-109", "reddit-110", "reddit-116", "reddit-117", "reddit-124", "reddit-168", "reddit-175"]),
        ("ChatGPT and Claude workflows", ["reddit-035", "reddit-042", "reddit-057", "reddit-131", "reddit-132", "reddit-153", "reddit-154", "reddit-164", "reddit-166", "reddit-169"]),
        ("Specialist product value and failure modes", ["reddit-084", "reddit-087", "reddit-097", "reddit-114", "reddit-118", "reddit-163", "reddit-165", "reddit-167", "reddit-174", "reddit-187", "reddit-202"]),
        ("Small-team buying and operational fit", ["reddit-017", "reddit-025", "reddit-050", "reddit-054", "reddit-119", "reddit-129", "reddit-152", "reddit-195"]),
    ]
    for group_title, ids in groups:
        if group_title == "Specialist product value and failure modes":
            page(doc)
        heading(doc, group_title, 2)
        for source_id in ids:
            item = records[source_id]
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph(p, after=3, line=1.03)
            add_hyperlink(p, f"{item['title']} ({source_id})", item["url"])

    core = doc.core_properties
    core.title = "How Themis / CounselOS Can Win the Legal AI Work Layer"
    core.subject = "Competitive landscape and product strategy based on 203 Reddit source records"
    core.author = "CounselOS"
    core.keywords = "Themis.ai, CounselOS, legal AI, ChatGPT, Claude, Harvey, Legora, CoCounsel, GC AI, product strategy"
    doc.save(OUT)
    print(f"Wrote {OUT} with {len(records_list)} source records in the research base")


if __name__ == "__main__":
    build()
