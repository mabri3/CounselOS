"""Build a market-first legal AI customer-voices report."""

from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Pt

from build_report import (
    BLUE,
    INK,
    LIGHT_GRAY,
    MUTED,
    add_body,
    add_bullet,
    add_hyperlink,
    callout,
    heading,
    set_font,
    set_paragraph,
    set_table_geometry,
    shade,
    simple_table,
)
from build_strategy_report import clear_body, page, prevent_row_split, tighten_table


ROOT = Path(__file__).resolve().parents[2]
REFERENCE = ROOT / "output/research/legal-ai-small-team-market-research.docx"
CORPUS = ROOT / "output/research/raw/legal-ai-reddit-source-corpus.jsonl"
OUT = ROOT / "output/research/legal-ai-customer-voices-market-research.docx"


THEME_FIELDS = ("reported_outcome", "evidence_excerpt", "task_tested", "notes")
THEME_RULES = [
    (
        "Drafting and rewriting",
        re.compile(r"first draft|drafting|drafted|redline|clause|rewrite|proofread|correspondence|defined terms|cross-reference", re.I),
        "Most common use-case discussion; includes both praise and warnings.",
    ),
    (
        "Workflow and integration",
        re.compile(r"workflow|integration|\bdms\b|imanage|netdocuments|\bword\b|connector|playbook|template|matter-centric|project management|\bclm\b|vault", re.I),
        "Repeated demand for Word, DMS, playbooks, connectors, and process fit.",
    ),
    (
        "Legal research and authority",
        re.compile(r"legal research|case law|citation|holding|westlaw|lexis|authorit|precedent research|research starting|source-linked", re.I),
        "Research stays central, but users often pair AI with controlled sources.",
    ),
    (
        "Document and batch review",
        re.compile(r"document review|document-scale|large document|mass document|diligence|discovery|timeline|medical record|transcript|review table|data extraction|extract(?:ed|ion)|bundle", re.I),
        "Diligence, discovery, timelines, extraction, and review tables recur.",
    ),
    (
        "Accuracy and verification concerns",
        re.compile(r"manual review|human review|verify|verification|hallucinat|invented|fake citation|wrong|incorrect|inaccurat|unreliable|missed (?:a |the )?(?:material )?(?:issue|concept|fact)|citation checking|check citations", re.I),
        "Wrong, missed, or invented outputs make checking part of the workflow.",
    ),
    (
        "Explicit price and value tension",
        re.compile(r"too expensive|expensive|overpriced|ridiculously priced|not worth|worth paying|justify (?:the )?cost|price increase|pricing|cheaper|lower[- ]priced|cost[- ]effective|budget|\$\d|per seat|seat/month|subscription", re.I),
        "Price, subscriptions, seat terms, and whether the product is worth it.",
    ),
    (
        "Privacy, security, and governance",
        re.compile(r"privacy|confidential|security|hipaa|data retention|legal hold|zero.data|soc ?2|governance|privilege", re.I),
        "Confidentiality, retention, security, legal hold, and governance concerns.",
    ),
]


def coded_text(item: dict) -> str:
    return " ".join(str(item.get(field, "")) for field in THEME_FIELDS)


def has_disclosed_legal_team_size(item: dict) -> bool:
    value = str(item.get("firm_or_team_size", ""))
    if "people worldwide" in value.lower():
        return False
    return bool(
        re.search(
            r"\bsolo\b|\b\d[\d,]*(?:\+)?(?:\s*(?:-|–|to)\s*\d[\d,]*(?:\+)?)?[ -]*(?:attorney|attorneys|lawyer|lawyers|person|people|member|members|seat|seats|staff|employee|employees|partner|partners|associate|associates)",
            value,
            re.I,
        )
    )


def add_links(doc, lead: str, links: list[tuple[str, str]], after: int = 8) -> None:
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=after, line=1.0)
    set_font(p.add_run(lead), 8.5, MUTED, True)
    for index, (label, url) in enumerate(links):
        if index:
            set_font(p.add_run(" | "), 8.5, MUTED)
        add_hyperlink(p, label, url)


def add_voice_table(doc, rows: list[tuple[str, str, str]], widths=(2100, 6000, 1260), font_size=8.8):
    table = simple_table(doc, ["Customer context", "Reported voice or result", "Source"], rows, list(widths), font_size)
    tighten_table(table)
    return table


def add_likes_dislikes(doc, likes: str, dislikes: str, why_chosen: str):
    table = simple_table(
        doc,
        ["What users like", "What users dislike", "Why it is chosen"],
        [(likes, dislikes, why_chosen)],
        [3120, 3120, 3120],
        9.15,
    )
    tighten_table(table)
    return table


def add_source_group(doc, title: str, ids: list[str], records: dict[str, dict], new_page=False):
    if new_page:
        page(doc)
    heading(doc, title, 2)
    for source_id in ids:
        item = records[source_id]
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph(p, after=3, line=1.03)
        add_hyperlink(p, f"{item['title']} ({source_id})", item["url"])


def build() -> None:
    records_list = [json.loads(line) for line in CORPUS.read_text().splitlines() if line.strip()]
    records = {item["source_id"]: item for item in records_list}
    subreddits = Counter(item.get("subreddit", "not stated") for item in records_list)
    coded_records = [item for item in records_list if item.get("reported_outcome")]
    theme_counts = {
        label: sum(bool(pattern.search(coded_text(item))) for item in coded_records)
        for label, pattern, _ in THEME_RULES
    }
    size_disclosures = sum(has_disclosed_legal_team_size(item) for item in records_list)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUT)
    doc = Document(OUT)
    clear_body(doc)

    # Cover
    p = doc.add_paragraph()
    set_paragraph(p, before=10, after=4, line=1.0)
    set_font(p.add_run("MARKET RESEARCH REPORT"), 10, BLUE, True)
    p = doc.add_paragraph()
    set_paragraph(p, after=5, line=1.0)
    set_font(p.add_run("What legal AI customers are actually saying"), 22, INK, True)
    p = doc.add_paragraph()
    set_paragraph(p, after=16, line=1.0)
    set_font(
        p.add_run(
            "Customer voices on ChatGPT, Claude, Harvey, Legora, CoCounsel, and GC AI - how the tools are used, why they are chosen, and where the market still falls short"
        ),
        13,
        MUTED,
    )

    metadata = doc.add_table(rows=5, cols=2)
    metadata.style = "Table Grid"
    set_table_geometry(metadata, [1500, 7860])
    metadata_values = [
        ("Prepared for", "Legal AI market and product research"),
        ("Date", "September 1, 2026"),
        ("Corpus", f"{len(records_list)} de-duplicated Reddit threads; 100 added sources from March-August 2026"),
        ("Communities", f"r/legaltech ({subreddits['r/legaltech']}), r/biglaw ({subreddits['r/biglaw']}), r/Lawyertalk ({subreddits['r/Lawyertalk']}), r/LawFirm ({subreddits['r/LawFirm']}), and smaller lawyer-centered groups"),
        ("Evidence standard", "Direct use and task comparisons receive the most weight. Anonymous, vendor-linked, hearsay, and incomplete claims remain directional."),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_values):
        shade(row.cells[0], LIGHT_GRAY)
        p1 = row.cells[0].paragraphs[0]
        set_paragraph(p1, after=0, line=1.0)
        set_font(p1.add_run(label), 9.8, INK, True)
        p2 = row.cells[1].paragraphs[0]
        set_paragraph(p2, after=0, line=1.0)
        set_font(p2.add_run(value), 9.8, INK)
        prevent_row_split(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(5)
    callout(
        doc,
        "Bottom line",
        "There is no single legal AI winner. Lawyers assemble a stack: ChatGPT or Claude for broad thinking and drafting; Westlaw, Lexis, or another controlled source for authority; and specialist tools for selected document, research, or team workflows. Products win when they remove work after the first answer. They lose when they add cost, verification, or another disconnected system.",
    )

    # Market summary
    page(doc)
    heading(doc, "The market in one page", 1)
    add_body(
        doc,
        "The customer voices show an active but unsettled market. Lawyers are using AI every week, but use is concentrated in work that is easy to review: first drafts, summaries, issue lists, correspondence, clause work, document extraction, and research starting points. Trust falls as the work becomes citation-heavy, fact-sensitive, strategic, or difficult to verify.",
    )
    table = simple_table(
        doc,
        ["Market pattern", "What customers report", "What it means"],
        [
            ("General models are the baseline", "ChatGPT appears in 85 records and Claude in 73. Users value breadth, current model quality, flexibility, and price.", "A specialist product must add workflow value, not only a legal label."),
            ("Specialists win selected jobs", "Harvey, Legora, and CoCounsel are valued for legal sources, review tables, large files, Word/DMS links, and approved deployment.", "Operational leverage is the main reason to pay a premium."),
            ("Research stays controlled", "Users commonly use AI for orientation, then verify in Westlaw, Lexis, Google, law-firm memos, or supplied authorities.", "No product in the corpus removes the lawyer's verification duty."),
            ("Firm size changes the choice", "Small teams resist seat minimums and rollout burden; larger firms pay for standardization, security review, integration, and easier adoption.", "The same product can be good value for BigLaw and poor value for a three-lawyer team."),
            ("Users build stacks", "Examples combine Claude with Westlaw or Lexis, ChatGPT Enterprise with point tools, and specialist review with Word or DMS systems.", "The practical competitor is often a bundle, not one vendor."),
        ],
        [2200, 4020, 3140],
        8.9,
    )
    tighten_table(table)
    heading(doc, "Discussion coverage", 1)
    table = simple_table(
        doc,
        ["Product", "Threads mentioning it", "How to read the count"],
        [
            ("ChatGPT", "85", "Broad baseline across drafting, research, client behavior, policy, and operations."),
            ("Claude", "73", "Strong recent discussion around drafting, document work, projects, skills, and internal workflows."),
            ("Harvey", "60", "Frequent enterprise and legaltech comparisons; substantial positive and negative evidence."),
            ("CoCounsel", "30", "Mostly research, litigation, records, discovery, and Westlaw-related workflow."),
            ("Legora", "29", "Concentrated in BigLaw, review-table, drafting, diligence, research, and vendor-choice discussions."),
            ("GC AI", "5", "Evidence is too sparse for a strong market conclusion."),
        ],
        [1650, 1650, 6060],
        9.2,
    )
    tighten_table(table)
    add_body(doc, "These are discussion counts, not usage, satisfaction, preference, or market-share measures.")

    # Quantitative signal map
    page(doc)
    heading(doc, "Quantitative signal map", 1)
    add_body(
        doc,
        f"The corpus supports directional tallies, not survey results. The unit is a de-duplicated Reddit thread record, not a person. Of {len(records_list)} records, {len(coded_records)} contain detailed evidence fields that can be coded. One record can count in more than one theme.",
    )
    table = simple_table(
        doc,
        ["Repeated theme", "Records", "% of coded records", "What the tally supports"],
        [
            (
                label,
                str(theme_counts[label]),
                f"{theme_counts[label] / len(coded_records):.0%}",
                meaning,
            )
            for label, _, meaning in THEME_RULES
        ],
        [2240, 850, 1250, 5020],
        8.45,
    )
    tighten_table(table)
    heading(doc, "Evidence base", 2)
    table = simple_table(
        doc,
        ["Measure", "Count", "Meaning"],
        [
            ("Total de-duplicated records", str(len(records_list)), "All unique Reddit thread records in the source corpus."),
            ("Detailed records coded", str(len(coded_records)), "Records with outcome, evidence, task, and notes fields used for theme coding."),
            ("Legacy links excluded from theme coding", str(len(records_list) - len(coded_records)), "Links retained for coverage but without enough detail for the tally."),
            ("Solo or numeric legal-team size disclosed", f"{size_disclosures} ({size_disclosures / len(records_list):.0%})", "Records that identify a solo practice or give a numeric legal team or firm size."),
            ("Curated strong direct purchase decisions", "8", "High-signal selections or rejections compared directly later in this report."),
        ],
        [2800, 1100, 5460],
        8.5,
    )
    tighten_table(table)
    callout(
        doc,
        "How to read the counts",
        "These totals show repetition, not market share, agreement, or net sentiment. A drafting record may praise drafting or warn against it. The product pages, price section, and direct-decision table explain the customer reasons behind the totals.",
    )

    # Work patterns
    page(doc)
    heading(doc, "How lawyers are using legal AI", 1)
    add_body(
        doc,
        "Use is strongest where the output can be checked quickly against known facts, supplied documents, or an existing draft. The same users who report major time savings often reject AI as final authority.",
    )
    table = simple_table(
        doc,
        ["Work category", "Observed uses", "Reported value", "Recurring failure"],
        [
            ("Drafting and rewriting", "First drafts, clauses, correspondence, advice, discovery requests, outlines, defined terms, cross-references.", "Faster starting point and less typing; sometimes hours saved.", "Generic language, wrong assumptions, weak strategy, or a second editing job."),
            ("Summaries and orientation", "Cases, transcripts, contracts, records, bundles, call notes, business requests.", "Rapid issue map and easier entry into unfamiliar material.", "Missing material facts, overlong summaries, or false confidence that all files were read."),
            ("Legal research", "Issue framing, source leads, case analysis, citation-linked starting points.", "Useful signposting and faster first pass when sources are constrained.", "Invented citations, wrong holdings, missing currentness, and mandatory manual checks."),
            ("Document-scale review", "Diligence, discovery, timelines, testimony search, review tables, extraction, comparisons.", "High value when the system handles many files and exposes results in a table.", "Skipped concepts, incomplete coverage, timeouts, and unclear exception handling."),
            ("Firm-specific workflow", "Playbooks, precedent cleanup, NDA review, prompt libraries, internal connectors, matter templates.", "Better results after context and examples are supplied.", "Setup burden, maintenance, rigid vendor workflows, or poor fit with existing systems."),
            ("Operational work", "Intake, meeting notes, task preparation, pitches, client communications, simple automation.", "Broad nonlegal value improves subscription economics.", "Fragmented tools and weak handoff into the matter or practice system."),
        ],
        [1750, 3150, 2300, 2160],
        8.45,
    )
    tighten_table(table)
    heading(doc, "The common working stack", 1)
    add_bullet(doc, "General AI for thinking, drafting, summarizing, and mixed legal-commercial work.")
    add_bullet(doc, "Westlaw, Lexis, PACER, supplied cases, or another controlled source for authority and citation checking.")
    add_bullet(doc, "Word, a DMS, a CLM, or a point tool for the actual review and document workflow.")
    add_bullet(doc, "Human review for legal judgment, material facts, citations, and final work product.")
    callout(
        doc,
        "Customer behavior",
        "The market is not moving from 'no AI' to 'one AI platform.' It is moving toward mixed stacks in which each tool is trusted for a different part of the work.",
    )

    # Buying logic
    page(doc)
    heading(doc, "Why customers choose one tool over another", 1)
    add_body(
        doc,
        "The choice changes with the work. General models win when drafting quality, flexibility, and price matter most. Legal products win when sources, Word or DMS workflow, review tables, and firm-wide adoption remove more work than the model alone. The examples below show the contrast in actual customer decisions.",
    )
    table = simple_table(
        doc,
        ["Use case / customer", "Chosen over", "Why the customer chose it", "Customer evidence"],
        [
            ("Reasoning, citations, Word documents, and redlines; 3-attorney legal department", "ChatGPT over Harvey and Claude", "ChatGPT performed best on the same prompts. Harvey's uneven issue coverage did not justify about $30,000 per year.", "Reported result: ChatGPT was preferred for every tested work category. (reddit-110)"),
            ("Transactional drafting and precedent cleanup; lawyer and partner commenters", "Claude over CoCounsel for generation", "Claude was strong at defined terms, cross-references, precedent cleanup, and first drafts when supplied with the deal context.", '"For the dirty work of drafting, it\'s insanely good." (reddit-131)'),
            ("Corporate legal deployment; team size not stated", "Harvey over ChatGPT", "The buyer valued references, nuanced examples, Word integration, legal databases, and a usable interface more than a clear answer-quality lead.", '"We couldn\'t prove that the AI answers are consistently better than ChatGPT." (reddit-103)'),
            ("Firm-wide workflows; larger firm", "Harvey or Legora over Claude Enterprise", "The firm wanted packaged workflows and scheduling that ordinary users could adopt, not flexible one-off skills that still required configuration.", '"The skills aren\'t real workflows, neither is the scheduling." (reddit-147)'),
            ("Discovery, motion, and memo trial; about 30 attorneys", "Internal build with Claude Opus over Harvey, Legora, and CoCounsel", "Claude Opus found a material issue the specialist products missed. The firm also wanted its own DMS and project workflow and still needed Westlaw.", '"The only one that caught [the heirship issue] ... was Claude Opus." (reddit-175)'),
            ("Research authority plus generative drafting; size not stated", "Claude plus CoCounsel/Westlaw connector over either alone", "The buyer used CoCounsel for legal-source access and Claude for stronger generation and customization.", '"Claude with custom skills and legal connectors ... works quite well at $200/mo." (reddit-154)'),
            ("Daily legal reasoning and drafting; size not stated", "GC AI over Claude in one direct-use comment", "The user wanted legal defaults and daily convenience without building a custom general-model workflow. Evidence remains sparse.", 'Reported result: GC AI was called "excellent for daily use" and preferred for legal reasoning. (reddit-167)'),
        ],
        [1800, 1920, 3060, 2580],
        7.95,
    )
    tighten_table(table)
    add_links(doc, "Sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-110", "reddit-131", "reddit-103", "reddit-147", "reddit-175", "reddit-154", "reddit-167"]])
    callout(
        doc,
        "What the contrast shows",
        "ChatGPT and Claude win when the buyer values model quality, breadth, customization, and low cost. Harvey and Legora win when packaged workflow and broad adoption matter more than model control. CoCounsel is strongest as an authority and litigation layer. GC AI may offer useful legal defaults, but the evidence is not yet strong enough to show a durable advantage.",
    )

    # Cost and value
    page(doc)
    heading(doc, "Price decides the sale when the value gap is not obvious", 1)
    add_body(
        doc,
        "Customers do not describe price as a separate purchasing factor. They reject the combination of a high fee and an answer that is not better than ChatGPT or Claude, a product that does not replace the rest of the stack, or a minimum commitment before value is proved. They accept a premium when the product removes measurable work. The next page lists every numeric price signal; this page shows the decisions in customers' own words.",
    )
    table = simple_table(
        doc,
        ["Buyer and reported cost", "What the customer said", "Why the price won or lost"],
        [
            (
                "Large sub-AmLaw firm\nHarvey: $1,200/seat/month; $2,400 with Lexis; later about $399",
                '"At four figures a seat, I expect the demo folks to be able to answer basic questions with depth." It was not consistent. (reddit-117)',
                "Harvey lost pricing power when the product demonstration did not show four-figure-per-seat performance. The same buyer reported about $1,600/seat/month for its CoCounsel stack and said rewrite savings were not consistent.",
            ),
            (
                "Small company; exact size not stated\nHarvey: 5-seat minimum at $425/seat/month for 12 months - a $25,500 floor",
                '"For 5 Harvey seats (minimum), 12 months lock-in, I was told it was USD $425/seat/month." (reddit-163)',
                "The issue was not only the seat price. A small buyer had to commit to five seats and a full year before proving use. Claude was the lower-cost comparison in the same discussion.",
            ),
            (
                "Full-service firm; about 100+ lawyers\nHarvey price not stated; Claude Enterprise plus a smaller research tool selected",
                '"We actually tried Harvey early last year but its usefulness didn\'t justify the cost back then." (reddit-168)',
                "Even an enterprise-size buyer rejected Harvey when usefulness did not clear the premium. Claude won on flexibility; the firm kept a separate legal-research layer.",
            ),
            (
                "Small California firm\nHarvey price not stated",
                '"I was honestly shocked by how expensive and inflexible their pricing model is... no real path to try it." (reddit-107)',
                "Inflexible terms and no meaningful small-team trial stopped the buyer before Harvey could prove workflow value.",
            ),
            (
                "Small firm; exact size not stated\nCoCounsel price increased; useful upgrades required more payment",
                '"I use [CoCounsel] a lot... I\'m considering switching to Westlaw \'Classic\' with NO AI." (reddit-109)',
                "A frequent user was willing to unbundle research and AI: keep lower-cost traditional Westlaw access, then buy a cheaper AI tool for generation.",
            ),
            (
                "Team size not stated\nConfigured Claude with legal connectors: about $200/month",
                '"Claude with custom skills and legal connectors set up plus proper context works quite well at $200/mo." (reddit-154)',
                "Claude was 'good enough' because the buyer could add context and connectors while using CoCounsel mainly as a source layer, not as the generator.",
            ),
            (
                "Non-US firm; about 5,000 people\nHarvey rack-rate estimate near $1,000/seat/month; enterprise discount received",
                '"Saved me a hour or so there." The user said a 20-30% time saving could be "massive." (reddit-116)',
                "This is the premium case: repeated time savings, easy use, and Vault document extraction can justify enterprise spend, especially on fixed-fee or capped work.",
            ),
        ],
        [2280, 3300, 3780],
        8.1,
    )
    tighten_table(table)
    add_links(doc, "Direct sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-117", "reddit-163", "reddit-168", "reddit-107", "reddit-109", "reddit-154", "reddit-116"]])
    callout(
        doc,
        "The pricing message",
        "The specialist premium fails when the buyer still gets similar model output, keeps the existing research and document stack, and adds review or rewrite work. ChatGPT and Claude become the default because they are cheaper and broad enough. A specialist can still win when it produces repeated, measurable time savings or makes a large team easier to support.",
    )

    page(doc)
    heading(doc, "Reported prices and what buyers got for the money", 1)
    add_body(
        doc,
        "These anonymous, self-reported figures may reflect different dates, plans, regions, integrations, negotiations, or bundles. They are buyer evidence, not a current vendor price list. Hearsay is labeled.",
    )
    table = simple_table(
        doc,
        ["Product", "Reported price signal", "Customer response and meaning"],
        [
            ("ChatGPT", "$20 consumer access was cited in one discussion; another small-firm user asked whether a $200 monthly tier was worthwhile.", "The low entry point makes ChatGPT the baseline. Buyers tolerate more setup and verification because the same subscription supports legal and nonlegal work."),
            ("Claude", "A configured Claude setup with legal connectors was reported at about $200/month. Another buyer said $20 tiers were too limited, $200 became usable, and higher-limit seats could reach about $250/month.", "At $200, one buyer called Claude good enough for a solo; another said a configured legal stack worked well. The tradeoff was manual matter setup and no complete legal workflow."),
            ("Harvey", "Reports included about $30,000/year for a small department; a $25,500 annual floor for five $425 seats; $1,200/seat/month, $2,400 with Lexis, then about $399; and a rack-rate estimate near $1,000.", "Small teams rejected the premium when tested output was not better or the trial path was weak. A 5,000-person firm saw a possible 20-30% time saving as material and used Harvey regularly."),
            ("Legora", "No dependable direct price was retained. One poster said a friend reported about $400/user/month; another heard about $2,000/seat with Legora lower. Both are hearsay.", "Unclear pricing weakens trust when buyers also question support, document-scale quality, or differentiation from Claude and Harvey. Repeated tabular review is the clearest value case."),
            ("CoCounsel", "A 2023 report cited about $50-$75/task and a subscription becoming preferable above roughly $500/month. A large firm's bundled stack was reported at about $1,600/seat/month.", "The large-firm buyer said generation did not consistently save rewrite time. A frequent small-firm user considered traditional Westlaw plus a lower-priced AI after increases and paid upgrades."),
            ("GC AI", "The five GC AI records do not provide a usable numeric price benchmark.", "Positive daily-use reports are not enough to determine whether buyers see good value versus ChatGPT or Claude."),
        ],
        [1250, 3450, 4660],
        7.95,
    )
    tighten_table(table)
    add_links(doc, "Price evidence: ", [(sid, records[sid]["url"]) for sid in ["reddit-188", "reddit-010", "reddit-154", "reddit-110", "reddit-163", "reddit-117", "reddit-116", "reddit-122", "reddit-028", "reddit-109"]])
    heading(doc, "What customers consider worth paying for", 1)
    table = simple_table(
        doc,
        ["Worth a premium", "Usually not worth a premium"],
        [
            ("Reliable large-file and batch review; clear review tables; precise redlines; legal-source access; DMS/Word links; approved data controls; firm-wide templates; low prompting burden; measurable repeated time savings.", "A legal label alone; an answer similar to a general model; citations that still require full reconstruction; a chat interface without matter workflow; features used only occasionally; rigid output; extra seats; long lock-in before a real trial."),
        ],
        [4680, 4680],
        9.0,
    )
    tighten_table(table)
    callout(
        doc,
        "Cost conclusion",
        "For small teams, ChatGPT or Claude is often the default because the downside is limited and the subscription has broad use. A specialist must prove that it removes enough verification, document handling, or workflow effort to exceed a large price gap. For large firms, the calculation changes: adoption, security, standardization, and integration can be worth more than the model itself.",
    )

    # ChatGPT
    page(doc)
    heading(doc, "ChatGPT: the broad, familiar baseline", 1)
    add_body(
        doc,
        "ChatGPT is used as a general work tool rather than a complete legal system. Lawyers use it for first drafts, proofreading, brainstorming, issue spotting, clauses, correspondence, commercial advice, operations, and research framing. Its broad value makes the specialist premium difficult to justify for small teams.",
    )
    add_likes_dislikes(
        doc,
        "Fast, familiar, broad across legal and nonlegal work, flexible instructions, strong writing and reasoning for some users, and low entry cost.",
        "Fake cases, unreliable holdings, missed transcript or timeline details, weak final research, confidentiality concerns, and client-generated errors that create lawyer work.",
        "Chosen when the team wants one versatile tool, can supply its own context, and will verify legal authority elsewhere.",
    )
    heading(doc, "Customer voices", 1)
    add_voice_table(doc, [
        ("GC; 3-attorney legal department", "Preferred ChatGPT over Harvey and Claude on the same prompts for reasoning, citations, Word documents, and redlines; did not see about $30,000/year of Harvey value.", "reddit-110"),
        ("Commercial-advice user; size not stated", "Reported that ChatGPT Enterprise completed roughly 80% of a 25-page offshore-power advice task; review took about 10 hours instead of about a week of drafting.", "reddit-167"),
        ("Solo/small-firm lawyer", "Used it for proofreading, brainstorming, and basic clause drafts after attorney editing; did not use it as a source of legal precedent.", "reddit-010"),
        ("Law-firm user; size not stated", "Could not get ChatGPT to summarize an examination-before-trial transcript accurately.", "reddit-009"),
        ("SME legal department", "Found enterprise controls useful for simple drafting and spotting missing items, but called legal research and substantive legal problem-solving failures.", "reddit-031"),
    ])
    add_links(doc, "Sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-110", "reddit-167", "reddit-010", "reddit-009", "reddit-031"]])
    callout(
        doc,
        "Why users choose it over legal AI",
        "ChatGPT often wins when the work is broad, the user is skilled, and authoritative research remains separate. It loses when confidentiality, reliable sourcing, document coverage, or firm-wide controls matter more than flexibility.",
    )

    # Claude
    page(doc)
    heading(doc, "Claude: favored for drafting and configurable workflows", 1)
    add_body(
        doc,
        "Claude is commonly used for transactional drafting, litigation workflows, contract review with playbooks, document synthesis, internal agents, coding, and repeatable projects. Recent customer voices often describe Claude as the strongest general-model alternative to a legal suite.",
    )
    add_likes_dislikes(
        doc,
        "Strong drafting, large-context document work, projects and reusable instructions, custom skills and connectors, model choice, and good results when precedent or playbooks are supplied.",
        "Usage limits, clunky Word-file generation, setup effort, no authoritative legal corpus, verification needs, and consumer-account privacy concerns.",
        "Chosen by capable users and technology teams that want flexibility, lower cost, and control over models, precedent, and workflow.",
    )
    heading(doc, "Customer voices", 1)
    add_voice_table(doc, [
        ("Junior transactional lawyer; size not stated", "Reported hours saved each week on defined terms, precedent cleanup, cross-references, and first drafts.", "reddit-131"),
        ("Partner commenter; size not stated", "Said Claude Business plus agents outperformed a paralegal-and-associate team in a mock $100 million loan workflow and was much better than CoCounsel in that test.", "reddit-131"),
        ("Law-firm technology lead; about 30 attorneys", "Claude Opus caught a material heirship issue that Claude Sonnet, Harvey, Legora, and CoCounsel missed.", "reddit-175"),
        ("Technology committee; 100+ lawyer firm", "Selected Claude Enterprise plus a smaller research tool after an earlier Harvey trial did not justify cost.", "reddit-168"),
        ("Plaintiff-employment lawyer", "Reported a surprisingly decent case/discovery result but reached usage limits quickly.", "reddit-042"),
        ("In-house contract team", "Separate Claude chats with a contract, playbook, and context produced useful work, but the team wanted a repeatable process instead of manual upload and email handoff.", "reddit-132"),
    ], font_size=8.55)
    add_links(doc, "Sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-131", "reddit-175", "reddit-168", "reddit-042", "reddit-132"]])
    callout(
        doc,
        "Why users choose it over a legal suite",
        "Claude wins when the buyer believes the model and custom context do most of the useful work. It loses when the team needs turnkey legal sources, review tables, Word/DMS integration, or low-skill firm-wide adoption.",
    )

    # Harvey
    page(doc)
    heading(doc, "Harvey: enterprise workflow value with contested small-team economics", 1)
    add_body(
        doc,
        "Harvey has the broadest specialist discussion in the corpus. Positive users describe legal sources, review tables, large-document work, Word and DMS integration, Vault workflows, and firm-wide deployment. Negative users question whether those benefits justify the price when ChatGPT or Claude performs the underlying thinking and drafting well enough.",
    )
    add_likes_dislikes(
        doc,
        "Turnkey legal workflows, easier prompting, review tables, document scale, redlining, legal-source links, Word/DMS integration, enterprise deployment, and shared standardization.",
        "High or unclear pricing, minimum seats or commitments, mixed task quality, incomplete issue spotting, wrapper perception, and poor fit for small or uneven teams.",
        "Chosen by larger organizations that value deployment, integration, and adoption across many lawyers more than direct model control.",
    )
    heading(doc, "Customer voices", 1)
    add_voice_table(doc, [
        ("Lawyer; about 5,000-person firm", "Used Harvey regularly, reported about one hour saved on a pitch and good property-diligence results, but called it too expensive for small firms.", "reddit-116"),
        ("Mid-level BigLaw lawyer", "Called Harvey life-changing for fast regulatory and clearance work; other commenters in the same thread reported limited use.", "reddit-087"),
        ("Corporate legal department; size not stated", "Selected Harvey for nuanced examples, references, interface, Word integration, and legal databases despite no consistent answer-quality lead over ChatGPT.", "reddit-103"),
        ("GC; 3-attorney legal department", "Rejected the value at roughly $30,000/year after preferring ChatGPT on the tested work.", "reddit-110"),
        ("Small California firm", "Rejected expensive, inflexible pricing and the lack of a meaningful small-team trial or customization path.", "reddit-107"),
        ("Law-firm technology lead; about 30 attorneys", "Rejected Harvey with Legora and CoCounsel after all three missed a material issue and created DMS/project friction.", "reddit-175"),
    ], font_size=8.5)
    add_links(doc, "Sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-116", "reddit-087", "reddit-103", "reddit-110", "reddit-107", "reddit-175"]])
    callout(
        doc,
        "Why it wins or loses",
        "Harvey wins when operational deployment is the product. It loses when a capable small team sees the legal workflow as an expensive layer around a general model it can already use.",
    )

    # Legora
    page(doc)
    heading(doc, "Legora: strong review and drafting reports, but limited differentiation", 1)
    add_body(
        doc,
        "Legora is discussed mainly in BigLaw and legaltech threads about research, drafting, diligence, tabular review, templates, and Word workflow. Positive users report frequent use and meaningful efficiency. Negative users see similar answers to Harvey or weaker output than enterprise Claude or ChatGPT.",
    )
    add_likes_dislikes(
        doc,
        "Tabular review, diligence, drafting, research starting points, templates, Word workflow, standardized team use, and lower prompting burden.",
        "High price, unclear difference from Harvey, mixed model quality, rigid output, missed issues, and concern that general models are catching up quickly.",
        "Chosen when a firm wants a repeatable transactional or review workflow that many lawyers can use consistently.",
    )
    heading(doc, "Customer voices", 1)
    add_voice_table(doc, [
        ("BigLaw research discussion", "One user said recent Legora use had been highly helpful for research; others still preferred constrained databases or traditional search because narrative explanations can hallucinate.", "reddit-084"),
        ("Large-firm discussion", "One commenter reported using Legora every day for almost every legal task, while another mainly used Westlaw AI despite access to enterprise models.", "reddit-088"),
        ("BigLaw lawyer", "Found Harvey and Legora similar on one research question and slightly preferred Harvey's clearer review-table categories; commenters also reported strong Legora drafting and a claimed 30% efficiency gain.", "reddit-097"),
        ("Law-firm technology lead; about 30 attorneys", "Legora missed the same material issue as Harvey and CoCounsel in the three-week comparative trial.", "reddit-175"),
        ("Technology committee; 100+ lawyer firm", "Had not trialed Legora and expected limited differentiation; this is expectation, not direct product evidence.", "reddit-168"),
        ("UK lawyer discussion", "A user described extensive Legora tabular review for document questions such as identifying which document says a fact.", "reddit-202"),
    ], font_size=8.5)
    add_links(doc, "Sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-084", "reddit-088", "reddit-097", "reddit-175", "reddit-168", "reddit-202"]])
    callout(
        doc,
        "Why it wins or loses",
        "Legora wins where review tables and standardized legal workflow produce visible team efficiency. It loses when the buyer cannot see a durable advantage over Harvey or a configured general model.",
    )

    # CoCounsel
    page(doc)
    heading(doc, "CoCounsel: strongest around authority and litigation workflow", 1)
    add_body(
        doc,
        "CoCounsel is most often discussed as a Westlaw-connected research and litigation tool. Users apply it to discovery, depositions, records, timelines, brief review, and source-linked research. Its position is less secure for broad drafting and generative work.",
    )
    add_likes_dislikes(
        doc,
        "Westlaw connection, source-linked starting points, legal safeguards, large uploads, discovery and deposition workflows, records review, and known legal-research deployment.",
        "Generative limits, poor task following, missed concepts, timeline noise, weak full-document drafting, price increases, and continued citation checking.",
        "Chosen when research authority, litigation workflow, or approved legal data matters more than open-ended drafting flexibility.",
    )
    heading(doc, "Customer voices", 1)
    add_voice_table(doc, [
        ("Firm user; size not stated", "Praised legal-specific safeguards and pay-per-use options, but the firm paused a deposition use case over HIPAA certification concerns.", "reddit-028"),
        ("Small-firm lawyer", "Reported using CoCounsel frequently but considered switching after price increases and paid upgrades weakened the value.", "reddit-109"),
        ("Paralegal and support users", "Called discovery requests acceptable but reported that an amended complaint failed; records review missed requested concepts and created noisy timelines.", "reddit-187"),
        ("Westlaw research user", "Preferred Westlaw Deep Research to the bundled CoCounsel experience and still found it useful even when it was seriously wrong; treated it as a checked starting point.", "reddit-047"),
        ("Claude/connector commenter", "Said Westlaw plus CoCounsel worked as a connector with Claude, but CoCounsel was weak for generative work.", "reddit-154"),
        ("Law-firm technology lead; about 30 attorneys", "Rejected CoCounsel with Harvey and Legora after the products missed a material issue and did not replace Westlaw.", "reddit-175"),
    ], font_size=8.5)
    add_links(doc, "Sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-028", "reddit-109", "reddit-187", "reddit-047", "reddit-154", "reddit-175"]])
    callout(
        doc,
        "Why it wins or loses",
        "CoCounsel wins through the Westlaw relationship and litigation-specific work. It loses when users want stronger generative drafting, exact instruction following, lower cost, or a flexible model-centered workflow.",
    )

    # GC AI
    page(doc)
    heading(doc, "GC AI: promising daily legal use, but the evidence is sparse", 1)
    add_body(
        doc,
        "GC AI appears in only five records. The available customer voices are positive about legal reasoning and drafting, but the sample is too small to establish a stable product position, target segment, or comparative advantage.",
    )
    add_likes_dislikes(
        doc,
        "Legal-specific defaults, daily-use convenience, drafting, legal reasoning, playbooks, contract context, and repository-oriented workflow are reported positively by some users.",
        "Too few independent reports, unclear firm-size fit, limited direct trials, polarized views, and uncertainty about whether the paid layer beats Claude or ChatGPT with good context.",
        "Chosen by users who want legal context and workflow without building their own general-model setup, based on the limited evidence available.",
    )
    heading(doc, "Customer voices", 1)
    add_voice_table(doc, [
        ("Legaltech commenter; size not stated", "Called GC AI excellent for daily use and preferred it to Claude for legal reasoning.", "reddit-167"),
        ("UK lawyer discussion; varied sizes", "A participant reported using GC AI for drafting while others used Claude, Legora, Copilot, or bespoke tools for different jobs.", "reddit-202"),
        ("Frequent legal-AI evaluator; size not stated", "Said most tested tools were not worth it but placed GC AI among the few products worth considering.", "reddit-104"),
        ("Broader market discussion", "Other users argued that general models with web search, skills, and proper context handle much everyday legal work, making the specialist premium hard to prove.", "reddit-167"),
    ])
    add_links(doc, "Sources: ", [(sid, records[sid]["url"]) for sid in ["reddit-167", "reddit-202", "reddit-104"]])
    callout(
        doc,
        "Evidence warning",
        "The available evidence supports a hypothesis that GC AI can be a useful legal daily driver. It does not support a confident conclusion about product-market fit, price-to-value, or superiority over Claude or ChatGPT.",
    )
    heading(doc, "What the six products reveal together", 1)
    add_bullet(doc, "General models own breadth, flexibility, and price-to-value.")
    add_bullet(doc, "Specialist tools own selected workflows, controlled sources, integrations, and deployment.")
    add_bullet(doc, "No product has established a consistent lead in legal reasoning across the direct comparisons.")
    add_bullet(doc, "The buyer's team size, technical ability, existing stack, and verification burden often matter more than the product brand.")

    # Direct decisions
    page(doc)
    heading(doc, "Direct selection and rejection decisions", 1)
    add_body(
        doc,
        "The table below concentrates the strongest buying evidence. It shows why a product that works well for one organization can be rejected by another.",
    )
    table = simple_table(
        doc,
        ["Customer / size", "Products considered", "Decision or outcome", "Reason"],
        [
            ("GC; 3-attorney legal department", "Harvey, ChatGPT, Claude", "Preferred ChatGPT; rejected Harvey value.", "Better tested reasoning, citations, Word and redline work; no reason to pay about $30,000/year."),
            ("Technology lead; about 30-attorney firm", "Harvey, Legora, CoCounsel, Claude Opus", "Rejected all three specialists; chose internal build.", "Material issue missed; DMS and project friction; products did not replace Westlaw."),
            ("Technology committee; 100+ lawyer firm", "Claude Enterprise, Harvey, local research tool", "Selected Claude Enterprise plus smaller research tool.", "Earlier Harvey trial did not justify cost; Claude offered flexibility."),
            ("Lawyer; about 5,000-person firm", "Harvey and alternatives", "Used Harvey regularly and reported meaningful value.", "Easy use, time saved, document-scale diligence, and enterprise availability."),
            ("Corporate legal department; size not stated", "Harvey, ChatGPT", "Selected Harvey despite no consistent answer lead.", "References, examples, interface, Word integration, databases, and deployment."),
            ("Small California firm", "Harvey", "Rejected Harvey.", "Price, inflexibility, no small-team trial, and weak customization path."),
            ("Small firm; size not stated", "CoCounsel and alternatives", "Considered leaving after frequent use.", "Price increases and paid upgrades reduced value."),
            ("BigLaw discussion", "Harvey, Legora, Claude, ChatGPT, Copilot", "Mixed; no product winner.", "Review-table clarity, drafting quality, research controls, and user preference differed by task."),
        ],
        [1900, 2100, 2470, 2890],
        8.15,
    )
    tighten_table(table)
    add_links(doc, "Primary links: ", [(sid, records[sid]["url"]) for sid in ["reddit-110", "reddit-175", "reddit-168", "reddit-116", "reddit-103", "reddit-107", "reddit-109", "reddit-097"]])
    heading(doc, "Firm size changes the value equation", 1)
    add_body(
        doc,
        "Small teams tend to compare a specialist product with one or two inexpensive general-model seats and an existing research subscription. Large firms compare it with the cost of inconsistent work across hundreds of users, security review, training, DMS access, and firm-wide standardization. This is why high specialist pricing can be rational in a large deployment and unacceptable in a small department.",
    )
    callout(
        doc,
        "Important limit",
        f"Only {size_disclosures} corpus records identify a numeric legal-team or firm size or a solo practice. The size pattern is strong enough to guide interviews, but not to estimate demand or market share.",
    )

    # Market gaps
    page(doc)
    heading(doc, "Where the current market still falls short", 1)
    table = simple_table(
        doc,
        ["Market gap", "Customer problem", "Why current options do not fully solve it"],
        [
            ("Fast verification", "Lawyers must check citations, holdings, facts, and source support before relying on the answer.", "General models can invent sources; specialists can still be wrong or incomplete; source links do not prove the proposition was reviewed."),
            ("Proved document coverage", "Users do not know whether every file, page, concept, or exception was reviewed.", "Long context windows and large-upload claims do not produce a clear coverage receipt or exception list."),
            ("Reliable Word handoff", "Copy-paste, broken formatting, weak redlines, and unexplained edits create a second review job.", "General models are clunky; specialist Word tools vary; source-layout preservation remains difficult."),
            ("Persistent context", "Facts, prior decisions, playbooks, open questions, and work product are rebuilt across chats and tools.", "Projects help, but users still report manual setup, separate chats, and disconnected systems."),
            ("Small-team economics", "Solo and small teams need value without enterprise minimums, rollout, or unused seats.", "Many specialist products are designed and priced for larger firms; point tools fragment the workflow."),
            ("Flexible but usable workflow", "Experts want customization; ordinary users need a turnkey workflow.", "General models require skill and setup; specialist products can be rigid or hide model choice."),
            ("Authority plus drafting", "Users want controlled research and strong generative drafting in one flow.", "Research tools are weaker generators; general models lack controlled authority; users bridge them manually."),
            ("Matter follow-through", "AI produces an answer but does not reliably preserve the decision, owner, deadline, and next action.", "Chat and point tools often stop at output rather than moving the legal matter forward."),
        ],
        [2300, 3250, 3810],
        8.45,
    )
    tighten_table(table)
    heading(doc, "The central market tension", 1)
    add_body(
        doc,
        "Customers want the intelligence and flexibility of the newest general models, but they also want the sources, document control, integrations, and repeatability of a legal workflow product. The market has not settled on one product that delivers both without high cost, technical setup, or continued manual verification.",
    )
    callout(
        doc,
        "Unmet job",
        "Turn a strong model answer into dependable legal work: source-linked, complete enough to review, connected to the actual documents, and preserved with the matter's next action.",
    )

    # Comparative summary and Themis short section
    page(doc)
    heading(doc, "Comparative market summary", 1)
    table = simple_table(
        doc,
        ["Product", "Strongest observed position", "Main reason to choose", "Main objection"],
        [
            ("ChatGPT", "Broad general work", "Versatility, familiarity, price, and useful drafting/reasoning.", "Research reliability, confidentiality, and no durable legal workflow."),
            ("Claude", "Drafting and configurable document workflow", "Writing quality, projects, skills, context, and model control.", "No authoritative corpus, usage/setup friction, and weak Word handoff."),
            ("Harvey", "Enterprise legal workflow", "Review tables, document scale, sources, integration, and adoption.", "Price, commitments, and mixed advantage over general models."),
            ("Legora", "Review, diligence, and standardized drafting", "Tabular review, Word workflow, and team consistency.", "Unclear differentiation, price, and mixed model-quality reports."),
            ("CoCounsel", "Research and litigation", "Westlaw, legal sources, discovery, records, and deposition workflow.", "Generative limits, task failures, price, and continued checking."),
            ("GC AI", "Legal daily-use assistant", "Legal reasoning, drafting, playbooks, and convenience for some users.", "Too little evidence to establish a clear moat or buyer segment."),
        ],
        [1350, 2700, 2850, 2460],
        8.75,
    )
    tighten_table(table)
    heading(doc, "Short comparison: Themis / CounselOS", 1)
    add_body(
        doc,
        "Themis is currently closest to the persistent-matter and follow-through gap. Its strengths are model neutrality, matter context, editable work product, decisions, work items, and visible next action. It is not a replacement for Harvey or Legora's enterprise deployment, CoCounsel's controlled legal research, or a live Word/DMS workflow.",
    )
    table = simple_table(
        doc,
        ["Current overlap with market needs", "Material gaps"],
        [
            ("Durable matter workspace; model choice; company and matter context; editable Markdown work product; research packets; decisions; work items; schedules; DOCX/PDF export.", "Exact passage-level proof; document-coverage receipts; structured batch review; source-layout-preserving or live Word workflow; authoritative Westlaw/Lexis handoff; team access; email/Slack/DMS integration; enterprise controls."),
        ],
        [4680, 4680],
        9.1,
    )
    tighten_table(table)
    callout(
        doc,
        "Market-research conclusion",
        "Themis addresses a real gap identified in the customer voices, but the evidence does not yet show that customers will pay for its approach. The next proof should be whether it reduces repeated setup and verification on real matters better than a user's existing ChatGPT or Claude stack.",
    )

    # Methods and sources
    page(doc)
    heading(doc, "Evidence strength and limitations", 1)
    add_body(
        doc,
        "The corpus contains 203 de-duplicated Reddit threads. The latest expansion added 100 unique threads dated March through August 2026. Direct product use, comparative trials, reported selection decisions, observable work, and stated team size receive the most weight. Broad opinions, vendor claims, affiliate comments, hearsay prices, and removed posts receive less weight.",
    )
    add_bullet(doc, "A product mention is not a positive review. Counts show discussion coverage only.")
    add_bullet(doc, "A Reddit report is not an independent product benchmark. Model versions, product releases, prompts, files, and user skill differ.")
    add_bullet(doc, "Self-reported pricing may reflect negotiation, region, plan tier, database bundles, or incomplete terms.")
    add_bullet(doc, "Firm and legal-team size is usually missing. Size conclusions are directional.")
    add_bullet(doc, "The report describes customer voices and market gaps. It does not estimate market size or market share.")
    heading(doc, "Curated source appendix", 1)
    add_body(
        doc,
        "The links below are the primary sources used for the product profiles and cross-product conclusions. The complete machine-readable corpus remains at output/research/raw/legal-ai-reddit-source-corpus.jsonl.",
    )
    add_source_group(doc, "ChatGPT and Claude", ["reddit-009", "reddit-010", "reddit-031", "reddit-035", "reddit-042", "reddit-057", "reddit-110", "reddit-131", "reddit-132", "reddit-154", "reddit-168", "reddit-175"], records)
    add_source_group(doc, "Harvey and Legora", ["reddit-084", "reddit-087", "reddit-088", "reddit-097", "reddit-103", "reddit-107", "reddit-116", "reddit-117", "reddit-124", "reddit-163", "reddit-168", "reddit-175", "reddit-202"], records, new_page=True)
    add_source_group(doc, "CoCounsel and GC AI", ["reddit-028", "reddit-047", "reddit-104", "reddit-109", "reddit-154", "reddit-165", "reddit-167", "reddit-175", "reddit-187", "reddit-202"], records)
    add_source_group(doc, "Small-team and market context", ["reddit-017", "reddit-025", "reddit-050", "reddit-054", "reddit-119", "reddit-129", "reddit-152", "reddit-195"], records)

    core = doc.core_properties
    core.title = "Legal AI Customer Voices Market Research"
    core.subject = "How customers use and assess ChatGPT, Claude, Harvey, Legora, CoCounsel, and GC AI"
    core.author = "CounselOS"
    core.keywords = "legal AI, customer voices, ChatGPT, Claude, Harvey, Legora, CoCounsel, GC AI, market research"
    doc.save(OUT)
    print(f"Wrote {OUT} from a {len(records_list)}-record research corpus")


if __name__ == "__main__":
    build()
