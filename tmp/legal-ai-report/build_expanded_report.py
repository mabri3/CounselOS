"""Build a complete, source-linked Word report from the de-duplicated Reddit corpus."""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_report import (
    BLUE, CALLOUT, DARK_BLUE, INK, LIGHT_BLUE, LIGHT_GRAY, MUTED, USABLE_DXA,
    add_body, add_bottom_border, add_bullet, add_hyperlink, add_number, add_page_field,
    callout, heading, set_font, set_paragraph, set_table_geometry, shade,
    simple_table,
)


CORPUS = Path("output/research/raw/legal-ai-reddit-source-corpus.jsonl")
OUT = Path("output/research/legal-ai-reddit-market-research-expanded.docx")


def product_counts(records):
    counts = Counter()
    for record in records:
        values = record.get("products", [])
        if isinstance(values, str):
            values = re.split(r"[,;/]", values)
        for value in values:
            name = str(value).strip().lower()
            if name:
                counts[name] += 1
    return counts


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph(hp, after=0, line=1.0)
    set_font(hp.add_run("COUNSELOS MARKET RESEARCH"), 8.5, MUTED, True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph(fp, after=0, line=1.0)
    set_font(fp.add_run("Confidential working research | Page "), 8.5, MUTED)
    add_page_field(fp)


def add_metadata(doc, records, source_files):
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1600, 7760])
    values = [
        ("Prepared for", "CounselOS product strategy"),
        ("Date", str(date(2026, 8, 31))),
        ("Research corpus", f"{len(records)} de-duplicated Reddit threads; all original raw records retained."),
        ("Expansion", "100 additional unique threads dated March 1-August 31, 2026."),
        ("Evidence standard", f"Firsthand accounts prioritized. {source_files} raw collection files were normalized into one source corpus; anonymous claims remain directional evidence."),
    ]
    for row, (label, value) in zip(table.rows, values):
        shade(row.cells[0], LIGHT_GRAY)
        p = row.cells[0].paragraphs[0]
        set_paragraph(p, after=0, line=1.0)
        set_font(p.add_run(label), 9.5, INK, True)
        p = row.cells[1].paragraphs[0]
        set_paragraph(p, after=0, line=1.0)
        set_font(p.add_run(value), 9.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_styled_hyperlink(paragraph, text, url, size=8.1):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(round(size * 2)))
    r_pr.append(sz)
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def add_source_card(doc, record):
    """One compact paragraph per complete raw evidence record."""
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=3, line=1.0)
    p.paragraph_format.keep_together = True
    set_font(p.add_run(f"{record['source_id']} | "), 8.1, DARK_BLUE, True)
    add_styled_hyperlink(p, record.get("title", "Untitled Reddit thread"), record["url"])
    fields = [
        ("Source", f"{record.get('subreddit', 'not stated')} | {record.get('date_if_visible', 'not stated')}"),
        ("Role/team", f"{record.get('poster_role', 'not stated')} | {record.get('firm_or_team_size', 'not stated')}"),
        ("Products", ", ".join(record.get('products', [])) if isinstance(record.get('products'), list) else str(record.get('products', 'not stated'))),
        ("Task", str(record.get('task_tested', 'not stated'))),
        ("Reported outcome", str(record.get('reported_outcome', 'not stated'))),
        ("ChatGPT/Claude comparison", str(record.get('chatgpt_or_claude_comparison', 'not stated'))),
        ("Evidence excerpt", str(record.get('evidence_excerpt', 'not stated'))),
        ("Confidence", str(record.get('confidence', 'not stated'))),
        ("Notes", str(record.get('notes', 'not stated'))),
    ]
    for label, value in fields:
        set_font(p.add_run(f" {label}: "), 7.7, DARK_BLUE, True)
        set_font(p.add_run(value), 7.7, INK)


def enable_two_columns(section):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "360")


def make_document():
    records = [json.loads(line) for line in CORPUS.read_text().splitlines() if line.strip()]
    records.sort(key=lambda item: item["source_id"])
    sources_by_id = {record["source_id"]: record for record in records}
    source_files = len({name for record in records for name in record.get("raw_files", [])})
    counts = product_counts(records)
    subreddits = Counter(record.get("subreddit", "not stated") for record in records)
    confidence = Counter(str(record.get("confidence", "not stated")).lower() for record in records)
    size_known = sum(bool(re.search(r"\b\d+[\- ]?(attorney|lawyer|person|people|firm|employee)", str(record.get("firm_or_team_size", "")), re.I)) or "solo" in str(record.get("firm_or_team_size", "")).lower() for record in records)

    doc = Document()
    configure(doc)

    p = doc.add_paragraph()
    set_paragraph(p, before=10, after=4, line=1.0)
    set_font(p.add_run("MARKET RESEARCH REPORT"), 10, BLUE, True)
    p = doc.add_paragraph()
    set_paragraph(p, after=5, line=1.0)
    set_font(p.add_run("Legal AI buying behavior: ChatGPT, Claude, Harvey, Legora, CoCounsel, GC AI, and CounselOS"), 21, INK, True)
    p = doc.add_paragraph()
    set_paragraph(p, after=15, line=1.0)
    set_font(p.add_run("Expanded Reddit research with 203 source records, product comparisons, and a complete raw evidence register"), 12.5, MUTED)
    add_metadata(doc, records, source_files)
    callout(doc, "Bottom line", "The paid legal-AI opportunity is not a better generic chatbot. It is a durable legal-work layer: reliable matter context, source verification, document coverage, repeatable review, and low-friction handoff while the lawyer continues to use the model they prefer.")

    doc.add_page_break()
    heading(doc, "Executive conclusion", 1)
    add_body(doc, "The expanded corpus reinforces the earlier conclusion. ChatGPT and Claude are the low-cost baseline for broad drafting, reasoning, brainstorming, correspondence, and custom workflows. Specialist tools are most defensible when they provide controlled legal sources, review of large document sets, Word or document-management integration, permissions, and a workflow that reduces the lawyer's verification burden.")
    add_body(doc, "The evidence does not establish that any single specialist platform has consistently superior legal reasoning. Direct trials are mixed. The more durable buying distinction is operational: whether a product saves real work after the model produces its first answer.")
    simple_table(doc, ["Finding", "Evidence in this corpus", "CounselOS implication"], [
        ("General models are the baseline", "ChatGPT appears in 85 records and Claude in 73. Users repeatedly cite price, flexibility, and strong general drafting.", "Remain model-neutral; do not compete on raw model claims."),
        ("Paid value is operational", "Harvey, Legora, CoCounsel, and GC AI are most often valued for workflows, sources, integrations, and document scale.", "Build source-grounded work products, coverage receipts, and durable matter state."),
        ("Verification remains mandatory", "Recent posts describe invented or misleading citations, missed issues, and broken AI edits across general and specialist tools.", "Show exact supporting passages, review status, and exceptions."),
        ("Small teams resist heavy contracts", "Price, minimum seats, and the cost of a second workflow recur in direct small-team discussions.", "Offer an immediate, low-maintenance matter workflow that complements existing model subscriptions."),
    ], [2550, 3400, 3410], 8.9)

    heading(doc, "Corpus and limitations", 1)
    add_body(doc, f"The corpus contains {len(records)} distinct Reddit threads. The source mix is {subreddits['r/legaltech']} r/legaltech, {subreddits['r/biglaw']} r/biglaw, {subreddits['r/Lawyertalk']} r/Lawyertalk, and {subreddits['r/LawFirm']} r/LawFirm threads, with smaller counts from other lawyer-centered communities. {size_known} records name either a numerical firm/team size or a solo practice.")
    add_body(doc, f"Agent-assigned evidence labels are {confidence['high']} high, {confidence['medium']} medium, {confidence['low']} low, and {confidence['not stated']} not stated. These labels are a review aid, not a scientific confidence score. Reddit is self-selected and anonymous; treat the report as directional buyer research, not a product benchmark or market-size study.")
    add_bullet(doc, "A product mention is not a positive review. Counts show discussion coverage, not preference or market share.")
    add_bullet(doc, "Pricing is illustrative only. It can reflect a negotiated quote, region, plan tier, or bundled databases.")
    add_bullet(doc, "The newer 100-source expansion is limited to March-August 2026, as requested.")

    heading(doc, "What pushes teams toward ChatGPT or Claude", 1)
    for title, text in [
        ("Price-to-value", "Small and uneven workloads make high per-seat contracts hard to justify when a team already has a general-model subscription. The question is not whether specialist products can help; it is whether their additional workflow value clears the contract cost."),
        ("Raw drafting and broad reasoning", "Users often prefer general models for writing, mixed legal-commercial questions, brainstorming, internal communication, and rapid access to the newest model capability. This is reported preference, not a controlled benchmark."),
        ("Control of the workflow", "Technically capable teams prefer their own instructions, skills, precedent, connectors, and model settings. They are skeptical of a fixed vendor workflow that duplicates a general model without preserving their specific context."),
        ("Low change-management cost", "A tool lawyers already know can start with low training and no new matter database. A paid system has to remove enough work to outweigh another interface, another migration, and another source of truth."),
        ("Separate authoritative research", "General models are commonly used for framing, synthesis, and drafting, while Westlaw, Lexis, or another controlled source remains the authority for current legal research and citation checking."),
    ]:
        heading(doc, title, 2)
        add_body(doc, text)

    heading(doc, "Where specialist legal tools still win", 1)
    simple_table(doc, ["Value area", "Why general chat is insufficient", "What users report valuing"], [
        ("Source-linked research", "A plausible answer or citation still needs verification.", "Controlled legal content, linked passages, and a starting point for research review."),
        ("Large-document and batch review", "Chat sessions can lose file coverage, structure, and exception tracking.", "Review tables, comparison, extraction, discovery, and a visible list of missing items."),
        ("Word and DMS workflow", "Copy-paste or file upload can create a second review task.", "In-document drafting, redlines, formatting preservation, and integrations with firm systems."),
        ("Enterprise controls", "Consumer model accounts may not meet confidentiality, retention, or access requirements.", "Permissions, approved deployment, auditability, and governed data connections."),
        ("Adoption at scale", "Not every lawyer will create custom prompts, skills, or connectors.", "Turnkey legal workflows that shorten training and standardize output."),
    ], [2150, 3550, 3660], 9.0)

    heading(doc, "Product comparison from the discussion", 1)
    simple_table(doc, ["Product", "Reported strengths", "Repeated objections"], [
        ("Harvey", "Review tables, document-scale work, legal-source and DMS integrations, shared workflows, and familiar enterprise deployment.", "High price or commitment, uneven task quality, and limited advantage over a well-configured Claude or ChatGPT workflow for capable users."),
        ("Legora", "Transactional/diligence workflow, template work, Word support, and standardization for teams.", "High per-seat cost, model or output concerns, and weak differentiation when firms can use general models with their own precedent."),
        ("CoCounsel", "Westlaw-connected research and litigation work, discovery, deposition and document-review workflows.", "Inconsistent task following, generative-work limits, review failures, and continued need to check substantive analysis."),
        ("GC AI", "Legal context, playbooks, contract work, drafting, and repository-oriented workflow for some users.", "Sparse and polarized evidence; users question whether the paid layer beats general models plus tailored context."),
    ], [1450, 4050, 3860], 8.8)

    heading(doc, "High-leverage direct examples", 1)
    examples = [
        ("3-attorney legal department", "A GC compared Harvey, ChatGPT, and Claude with the same prompts; ChatGPT was preferred for reasoning, citations, Word documents, and redlines. The user did not see value in approximately $30,000/year for Harvey.", "reddit-053"),
        ("About 30 attorneys", "A three-week trial of Harvey, Legora, and CoCounsel missed a material issue that Claude Opus caught. The firm chose a narrower internal tool.", "reddit-089"),
        ("100+ lawyer full-service firm", "A technology committee chose Claude Enterprise plus a smaller research tool after a Harvey trial did not justify its cost; counter-comments noted that specialist tools can be easier to deploy across a team.", "reddit-167"),
        ("About 5,000 people worldwide", "A non-US firm user reported time saved on a pitch and good property-diligence results with Harvey, while calling it expensive and unsuitable for small firms.", "reddit-059"),
        ("Recent Harvey versus Claude discussion", "Users identified review tables, large-document handling, redlining, integrations, and legal-source access as Harvey's case; others argued Claude plus connectors offers similar day-to-day work at lower cost.", "reddit-162"),
    ]
    simple_table(doc, ["Team", "Reported result", "Source ID"], [(a, b, c) for a, b, c in examples], [2000, 6150, 1210], 8.8)
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=8, line=1.0)
    set_font(p.add_run("Linked records: "), 8.3, MUTED, True)
    for index, (_, _, source_id) in enumerate(examples):
        if index:
            set_font(p.add_run(" | "), 8.3, MUTED)
        record = sources_by_id[source_id]
        add_hyperlink(p, source_id, record["url"])

    heading(doc, "CounselOS product requirements from the evidence", 1)
    priority_rows = [
        ("1", "Exact source passages and review status", "Let lawyers see the support for a proposition and mark it checked or unreviewed."),
        ("2", "Document-coverage receipts", "Prove which files, pages, and extracted issues the system reviewed; surface gaps and exceptions."),
        ("3", "Structured multi-document review", "Produce editable answer tables with source support, missing information, and follow-up questions."),
        ("4", "Durable matter orientation", "Keep intake, facts, decisions, files, work product, and next action together across model changes."),
        ("5", "Model-neutral work", "Make Claude, ChatGPT, and other models replaceable components, not a forced buying choice."),
        ("6", "Authoritative research handoff", "Bring Westlaw/Lexis or another lawyer-checked source into the matter without trying to build a legal corpus first."),
        ("7", "Low-maintenance daily-work capture", "Turn confirmed email, Slack, or intake work into a matter and handoff with minimal duplicate entry."),
        ("8", "Reliable Word workflow", "Keep tracked changes, explanations, formatting, and source support connected to the reviewed work product."),
    ]
    simple_table(doc, ["Priority", "Capability", "Reason"], priority_rows, [700, 3000, 5660], 9.0)
    heading(doc, "What not to build first", 1)
    for item in [
        "A proprietary legal model or a claim of universally superior legal reasoning.",
        "A generic prompt library as the main paid value.",
        "A full practice-management, billing, or accounting replacement.",
        "A legal research corpus built from scratch.",
        "Complex multi-agent checks that delay useful work product.",
        "Enterprise-scale governance before the small-team matter workflow proves adoption.",
    ]:
        add_bullet(doc, item)
    callout(doc, "Recommended product promise", "Give CounselOS the matter once. Use the model you prefer. Get review-ready work with traceable sources, proved document coverage, preserved context, and a clear next action.")

    heading(doc, "Complete raw evidence register", 1)
    add_body(doc, "This appendix contains every de-duplicated agent record used in this report. Each entry preserves its source link, role/team detail, products, task, reported outcome, ChatGPT/Claude comparison, evidence excerpt, confidence label, and notes. This is raw directional evidence, not independently verified product fact.")
    doc.add_section(WD_SECTION.CONTINUOUS)
    enable_two_columns(doc.sections[-1])
    for record in records:
        add_source_card(doc, record)

    core = doc.core_properties
    core.title = "Legal AI Market Research: Expanded Reddit Evidence"
    core.subject = "203 Reddit records on legal AI products, ChatGPT, Claude, and CounselOS product strategy"
    core.author = "CounselOS"
    core.keywords = "legal AI, CounselOS, ChatGPT, Claude, Harvey, Legora, CoCounsel, GC AI, Reddit research"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} with {len(records)} source records")


if __name__ == "__main__":
    make_document()
